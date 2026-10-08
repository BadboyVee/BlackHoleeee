#!/usr/bin/env python3
"""Read the figures the slides use out of the finished Segmentation_Matrix.xlsx.

Writes deck_data.json for build_deck.py: each segment's profile, Google Trends interest and rank,
factor scores, weighted score and rank, as Excel worked them out, so the slides match the
workbook; and the desk research.

Usage: python export_deck_data.py Segmentation_Matrix.xlsx deck_data.json
"""
import json
import sys

from openpyxl import load_workbook

from segment_data import BASIS, BUSINESS, FACTORS, RESEARCH, SEGMENTS, SOURCE

workbook, out = sys.argv[1:3]
wb = load_workbook(workbook, data_only=True)
matrix, scores = wb["Segmentation Matrix"], wb["Attractiveness"]
FIRST = 6
rows = []
for i, (name, *_rest) in enumerate(SEGMENTS):
    r = FIRST + i
    m = [matrix.cell(r, c).value for c in range(1, 7)]
    a = [scores.cell(r, c).value for c in range(1, 9)]
    if m[0] != name or a[0] != name or None in m + a:
        sys.exit(f"row {r} of {workbook} is incomplete: run the workbook through LibreOffice first")
    rows.append({"name": name, "who": m[1], "want": m[2], "term": m[3], "interest": m[4], "interest_rank": m[5],
                 "scores": a[1:6], "total": a[6], "rank": a[7], "basis": BASIS[i]})
with open(out, "w") as fh:
    json.dump({"business": BUSINESS, "source": SOURCE, "factors": [f for f, _ in FACTORS],
               "weights": [w for _, w in FACTORS], "segments": rows,
               "research": RESEARCH}, fh, indent=2)
print(f"wrote {out}")
