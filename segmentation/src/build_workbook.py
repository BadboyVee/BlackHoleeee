#!/usr/bin/env python3
"""Build Segmentation_Matrix.xlsx: the segmentation matrix of SkillUp Academy, a new
digital-skills training business, in Excel 2013's own style.

- Segmentation Matrix: the five customer segments (who they are, what they want), each one's
  Google Trends search term and average search interest (0-100), ranked with RANK, and a bar
  chart of the interest.
- Attractiveness: each segment scored 1 to 5 on five weighted factors. The search-interest
  score is =ROUND(interest/MAX(interest)*5,0); the weighted score is SUMPRODUCT of the weights
  and scores (out of 5), ranked with RANK; INDEX/MATCH names the most attractive segment. The
  reasons for each score are listed under the table, and a bar chart shows the weighted scores.
- Research: desk research that backs up the Google Trends findings, with the segment each
  finding supports, its source and a link.

Office 2013 theme, Calibri 11, dark-blue header rows, and charts in grey with the leading bar in
dark blue. Only functions that exist in Excel 2013 are used.

Usage: python build_workbook.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.page import PageMargins

from office2013 import THEME_XML
from segment_data import BASIS, BUSINESS, FACTORS, RESEARCH, SCORES, SEGMENTS, SOURCE

NAVY = Color(theme=4, tint=-0.499984740745262)      # "Blue, Accent 1, Darker 50%"  1F4E79
WHITE = Color(theme=0)
thin = Side(style="thin", color=Color(theme=0, tint=-0.249977111117893))
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N = len(SEGMENTS)
FIRST, LAST = 6, 5 + N                                    # segment rows on both sheets


def font(size=11, bold=False, italic=False, color="000000"):
    return Font(name="Calibri", size=size, bold=bold, italic=italic, color=color, family=2, scheme="minor")


def put(ws, ref, value, *, f=None, fill=None, fmt=None, h="left", wrap=True, bd=BOX):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fill is not None:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=1 if h == "left" else 0)
    if bd:
        c.border = bd
    return c


def head(ws, ref, text):
    return put(ws, ref, text, f=font(bold=True, color=WHITE), fill=NAVY, h="center")


def titles(ws, last_col, title, subtitle, source):
    for row, text, f in ((1, title, font(14, bold=True, color=NAVY)), (2, subtitle, font(italic=True)),
                         (3, source, font(9, italic=True, color="595959"))):
        ws.merge_cells(f"A{row}:{last_col}{row}")
        put(ws, f"A{row}", text, f=f, h="center", bd=None)
    ws.row_dimensions[1].height = 21


def page(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75, header=0.3, footer=0.3)


def rich(size=900, color="595959", bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def bars(ws, anchor, title, values, cats, fmt, vmax, top):
    """A horizontal bar chart in Excel 2013's style: grey bars, the leading one dark blue."""
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = "bar", "clustered", 60
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=1400, b=False, solidFill="595959")
    ch.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),
                                                        r=[RegularTextRun(rPr=cp, t=title)])])), overlay=False)
    ch.add_data(values, titles_from_data=False)
    ch.set_categories(cats)
    ch.legend = None
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill="7F7F7F", ln=LineProperties(noFill=True))
    s.dPt.append(DataPoint(idx=top, spPr=GraphicalProperties(solidFill="1F4E79", ln=LineProperties(noFill=True))))
    ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                  showPercent=False)
    ch.dataLabels.numFmt = fmt
    ch.dataLabels.position = "outEnd"
    ch.dataLabels.txPr = rich(1000, "404040", bold=True)
    ch.x_axis.delete = False
    ch.x_axis.scaling.orientation = "maxMin"              # first segment at the top, as in the table
    ch.x_axis.txPr = rich(1000)
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
    ch.y_axis.delete = True
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, vmax
    ch.y_axis.majorGridlines = None
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
    ch.width, ch.height = 18, 8.5
    ws.add_chart(ch, anchor)


wb = Workbook()
wb.loaded_theme = THEME_XML.encode()
wb._fonts = IndexedList([font()])
wb._named_styles["Normal"].font = font()
wb.properties.creator = wb.properties.lastModifiedBy = ""
wb.properties.title = "Segmentation Matrix"
wb.properties.subject = f"Segmentation matrix of {BUSINESS}, supported by Google Trends"

# ================================================================ 1. Segmentation Matrix
ws = wb.active
ws.title = "Segmentation Matrix"
page(ws, {"A": 27, "B": 44, "C": 36, "D": 22, "E": 15, "F": 9})
titles(ws, "F", f"SEGMENTATION MATRIX: {BUSINESS.upper()}",
       "A new digital-skills training business. Customers are grouped by the skill they want to learn.",
       f"Source: {SOURCE}.")
for col, text in zip("ABCDEF", ("Segment", "Who they are", "What they want", "Google Trends search term",
                                "Average interest (0-100)", "Rank")):
    head(ws, f"{col}5", text)
ws.row_dimensions[5].height = 32
for i, (name, who, want, term, interest) in enumerate(SEGMENTS):
    r = FIRST + i
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", who)
    put(ws, f"C{r}", want)
    put(ws, f"D{r}", term)
    put(ws, f"E{r}", interest, h="center")
    put(ws, f"F{r}", f"=RANK(E{r},$E${FIRST}:$E${LAST})", h="center")
    ws.row_dimensions[r].height = 32
interest = [s[4] for s in SEGMENTS]
bars(ws, f"A{LAST + 3}", "Google Trends Search Interest by Segment (0-100)",
     Reference(ws, min_col=5, min_row=FIRST, max_row=LAST), Reference(ws, min_col=1, min_row=FIRST, max_row=LAST),
     "0", 50, interest.index(max(interest)))

# ================================================================ 2. Attractiveness
ws = wb.create_sheet("Attractiveness")
page(ws, {"A": 27, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 16, "H": 9})
titles(ws, "H", "SEGMENT ATTRACTIVENESS SCORES",
       "Each factor is scored 1 (poor) to 5 (best); for competition, 5 means little competition. "
       "Weighted score = SUMPRODUCT of the weights and the scores, out of 5.",
       "Search interest score = ROUND(average interest / highest average x 5, 0), from Google Trends.")
ws.row_dimensions[2].height = 30
for col, text in zip("ABCDEFGH", ["Segment"] + [f for f, _ in FACTORS] + ["Weighted score", "Rank"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 32
put(ws, "A5", "Weight", f=font(bold=True))
for col, (_, weight) in zip("BCDEF", FACTORS):
    put(ws, f"{col}5", weight, f=font(bold=True), fmt="0%", h="center")
put(ws, "G5", "=SUM(B5:F5)", f=font(bold=True), fmt="0%", h="center")
put(ws, "H5", None)
matrix = "'Segmentation Matrix'"
for i, scores in enumerate(SCORES):
    r = FIRST + i
    put(ws, f"A{r}", f"={matrix}!A{r}", f=font(bold=True))
    put(ws, f"B{r}", f"=ROUND({matrix}!E{r}/MAX({matrix}!$E${FIRST}:$E${LAST})*5,0)", h="center")
    for col, score in zip("CDEF", scores):
        put(ws, f"{col}{r}", score, h="center")
    put(ws, f"G{r}", f"=SUMPRODUCT($B$5:$F$5,B{r}:F{r})", f=font(bold=True), fmt="0.00", h="center")
    put(ws, f"H{r}", f"=RANK(G{r},$G${FIRST}:$G${LAST})", h="center")
    ws.row_dimensions[r].height = 20
r = LAST + 2
ws.merge_cells(f"A{r}:B{r}")
put(ws, f"A{r}", "Most attractive segment", f=font(bold=True, color=WHITE), fill=NAVY)
ws.merge_cells(f"C{r}:H{r}")
put(ws, f"C{r}", f'=INDEX(A{FIRST}:A{LAST},MATCH(MAX(G{FIRST}:G{LAST}),G{FIRST}:G{LAST},0))&" ("&'
                 f'TEXT(MAX(G{FIRST}:G{LAST}),"0.00")&" of 5)"', f=font(12, bold=True))
for col in "BDEFGH":
    ws[f"{col}{r}"].border = BOX
ws.row_dimensions[r].height = 22
BEST = r

r = BEST + 2
ws.merge_cells(f"A{r}:H{r}")
put(ws, f"A{r}", "Basis for the scores", f=font(12, bold=True, color=NAVY), bd=None)
head(ws, f"A{r + 1}", "Segment")
ws.merge_cells(f"B{r + 1}:C{r + 1}")
head(ws, f"B{r + 1}", "Growth (Google Trends)")
ws.merge_cells(f"D{r + 1}:E{r + 1}")
head(ws, f"D{r + 1}", "Spread across Nigeria (Google Trends)")
ws.merge_cells(f"F{r + 1}:H{r + 1}")
head(ws, f"F{r + 1}", "Ability to pay and competition")
for col in "CEGH":
    ws[f"{col}{r + 1}"].border = BOX
for i, (growth, spread, other) in enumerate(BASIS):
    rr = r + 2 + i
    put(ws, f"A{rr}", f"=A{FIRST + i}", f=font(bold=True))
    for first, last, text in (("B", "C", growth), ("D", "E", spread), ("F", "H", other)):
        ws.merge_cells(f"{first}{rr}:{last}{rr}")
        put(ws, f"{first}{rr}", text)
    for col in "CEGH":
        ws[f"{col}{rr}"].border = BOX
    ws.row_dimensions[rr].height = 46
totals = [round(sum(w * v for (_, w), v in zip(FACTORS, (round(s[4] / max(interest) * 5 + 1e-9),) + sc)), 2)
          for s, sc in zip(SEGMENTS, SCORES)]
bars(ws, f"A{r + 3 + N}", "Weighted Attractiveness Score by Segment (out of 5)",
     Reference(ws, min_col=7, min_row=FIRST, max_row=LAST), Reference(ws, min_col=1, min_row=FIRST, max_row=LAST),
     "0.00", 5, totals.index(max(totals)))
ws.freeze_panes = "B5"

# ================================================================ 3. Research
ws = wb.create_sheet("Research")
page(ws, {"A": 6, "B": 62, "C": 20, "D": 34, "E": 24})
titles(ws, "E", "DESK RESEARCH",
       "Other research that backs up the Google Trends findings and the scores.",
       f"Trends data: {SOURCE}. Web sources found on 8 October 2026; figures as the sources report them. "
       "Click a link to open the source.")
for col, text in zip("ABCDE", ("No.", "Finding", "Segment it supports", "Source", "Link")):
    head(ws, f"{col}5", text)
ws.row_dimensions[5].height = 22
for i, (finding, segment, source, link) in enumerate(RESEARCH):
    r = 6 + i
    put(ws, f"A{r}", i + 1, h="center")
    put(ws, f"B{r}", finding)
    put(ws, f"C{r}", segment, f=font(bold=True))
    put(ws, f"D{r}", source)
    site = link.split("/")[2].removeprefix("www.")              # the site's name; the cell links to the page
    c = put(ws, f"E{r}", site, f=Font(name="Calibri", size=11, underline="single", color=Color(theme=10),
                                      family=2, scheme="minor"))
    c.hyperlink = link
    ws.row_dimensions[r].height = 48
ws.freeze_panes = "A6"

wb.save(sys.argv[1])
print(f"wrote {sys.argv[1]}")
