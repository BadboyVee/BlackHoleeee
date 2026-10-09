#!/usr/bin/env python3
"""Build Capstone_Market_Research.xlsx: the capstone customer research for CyberStart, SkillUp
Academy's beginner cybersecurity course, used to find market gaps and opportunities.

Sheets: Market (definition, key facts, market size estimate), Google Search, Google Trends,
Social Media, Reviews (competitors and course marketplaces), Gaps & Opportunities (scored), and
Sources. Live formulas (SUM, AVERAGE, RANK, ROUND, INDEX/MATCH and others that exist in Excel
2013); assumptions in yellow input cells; teal headings; bar charts in grey with the leading bar
in teal.

Usage: python build_workbook.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.page import PageMargins

from office2013 import THEME_XML
import capstone_data as D

TEAL, GREY, YELLOW = "0F766E", "B7C1CA", "FFF2CC"
thin = Side(style="thin", color="D9DEE3")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def font(size=11, bold=False, italic=False, color="000000", underline=None):
    return Font(name="Calibri", size=size, bold=bold, italic=italic, color=color, underline=underline, family=2,
                scheme="minor")


def put(ws, ref, value, *, f=None, fill=None, fmt=None, h="left", wrap=True, bd=BOX):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=1 if h == "left" else 0)
    if bd:
        c.border = bd
    return c


def heads(ws, row, cols, texts, height=30):
    for col, text in zip(cols, texts):
        put(ws, f"{col}{row}", text, f=font(bold=True, color="FFFFFF"), fill=TEAL, h="center")
    ws.row_dimensions[row].height = height


def section(ws, row, text, last_col):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    put(ws, f"A{row}", text, f=font(12, bold=True, color=TEAL), bd=None)
    ws.row_dimensions[row].height = 20


def titles(ws, last_col, title, subtitle):
    for row, text, f in ((1, title, font(14, bold=True, color=TEAL)), (2, subtitle, font(italic=True, color="404040"))):
        ws.merge_cells(f"A{row}:{last_col}{row}")
        put(ws, f"A{row}", text, f=f, bd=None)
    ws.row_dimensions[1].height = 21


def page(ws, widths, landscape=True):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6, header=0.3, footer=0.3)
    ws.sheet_view.showGridLines = False


def rich(size=1000, color="404040", bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def bars(ws, anchor, title, values, cats, fmt, vmax, top, width=16, height=7.5):
    """A horizontal bar chart: grey bars, the leading one teal, values at the bar ends."""
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = "bar", "clustered", 60
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=1300, b=False, solidFill="404040")
    ch.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),
                                                        r=[RegularTextRun(rPr=cp, t=title)])])), overlay=False)
    ch.add_data(values, titles_from_data=False)
    ch.set_categories(cats)
    ch.legend = None
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=GREY, ln=LineProperties(noFill=True))
    for i in top:
        s.dPt.append(DataPoint(idx=i, spPr=GraphicalProperties(solidFill=TEAL, ln=LineProperties(noFill=True))))
    ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                                  showPercent=False)
    ch.dataLabels.numFmt = fmt
    ch.dataLabels.position = "outEnd"
    ch.dataLabels.txPr = rich(1000, "262626", bold=True)
    ch.x_axis.delete = False
    ch.x_axis.scaling.orientation = "maxMin"              # first row at the top, as in the table
    ch.x_axis.txPr = rich(1000)
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
    ch.y_axis.delete = True
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, vmax
    ch.y_axis.majorGridlines = None
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)


def source(key):
    who, _what, _link = D.SOURCES[key]
    return who


wb = Workbook()
wb.loaded_theme = THEME_XML.encode()
wb._fonts = IndexedList([font()])
wb._named_styles["Normal"].font = font()
wb.properties.creator = wb.properties.lastModifiedBy = ""
wb.properties.title = "Capstone: Market Research"
wb.properties.subject = f"Customer research for {D.PRODUCT} by {D.BUSINESS}: market gaps and opportunities"

# ================================================================ 1. Market
ws = wb.active
ws.title = "Market"
page(ws, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 30})
titles(ws, "E", f"CAPSTONE PROJECT: {D.PRODUCT.upper()} BY {D.BUSINESS.upper()}",
       f"{D.PRODUCT}, {D.PRODUCT_LINE}. Research done on {D.RESEARCH_DATE}.")
r = 4
section(ws, r, "The business and product", "E")
for label, text in D.PRODUCT_POINTS:
    r += 1
    put(ws, f"A{r}", label, f=font(bold=True))
    ws.merge_cells(f"B{r}:E{r}")
    put(ws, f"B{r}", text)
    for col in "CDE":
        ws[f"{col}{r}"].border = BOX
r += 2
section(ws, r, "Defining the market", "E")
for label, text in D.MARKET_DEFINITION:
    r += 1
    put(ws, f"A{r}", label, f=font(bold=True))
    ws.merge_cells(f"B{r}:E{r}")
    put(ws, f"B{r}", text)
    for col in "CDE":
        ws[f"{col}{r}"].border = BOX
r += 2
section(ws, r, "The market in numbers", "E")
r += 1
heads(ws, r, "ABCDE", ("Indicator", "Value", "Unit", "When", "Source"), 22)
FACT = {}
for label, value, unit, when, key in D.MARKET_FACTS:
    r += 1
    fmt = "0.0%" if unit in ("of the population", "growth", "of the market") else ("#,##0.00" if value < 1000 and
                                                                                   value != int(value) else "#,##0")
    put(ws, f"A{r}", label)
    put(ws, f"B{r}", value, fmt=fmt, h="right")
    put(ws, f"C{r}", unit)
    put(ws, f"D{r}", when, h="center")
    put(ws, f"E{r}", source(key))
    FACT[(label, when)] = f"B{r}"
r += 1
put(ws, f"A{r}", "Yearly growth of the market, 2026 to 2031", f=font(bold=True))
growth = f"=({FACT[('Nigeria cybersecurity market (forecast)', '2031')]}/{FACT[('Nigeria cybersecurity market', '2026')]})^(1/5)-1"
put(ws, f"B{r}", growth, f=font(bold=True), fmt="0.0%", h="right")
put(ws, f"C{r}", "a year")
put(ws, f"D{r}", "Formula", h="center")
put(ws, f"E{r}", "Worked out from the rows above")
GROWTH = f"B{r}"
r += 1
put(ws, f"A{r}", "Experts in South Africa for each one in Nigeria", f=font(bold=True))
put(ws, f"B{r}", f"=ROUND({FACT[('Cybersecurity professionals in South Africa', '2023')]}/"
                 f"{FACT[('Cybersecurity professionals in Nigeria', '2023')]},1)", f=font(bold=True), fmt="0.0",
    h="right")
put(ws, f"C{r}", "times")
put(ws, f"D{r}", "Formula", h="center")
put(ws, f"E{r}", "Worked out from the rows above")
r += 2
section(ws, r, "Market size estimate (change the yellow cells to test other cases)", "E")
r += 1
heads(ws, r, "ABC", ("Step", "Figure", "How"), 22)
ws.merge_cells(f"C{r}:E{r}")
for col in "DE":
    ws[f"{col}{r}"].border = BOX
first = r + 1
for i, (label, value, kind, note) in enumerate(D.SIZING):
    r += 1
    prev = lambda k: f"B{first + k}"
    if kind == "input":
        cell_value, fill = value, YELLOW
    else:
        cell_value, fill = {2: f"={prev(0)}/{prev(1)}", 4: f"={prev(2)}*{prev(3)}", 6: f"={prev(4)}*{prev(5)}",
                            8: f"={prev(6)}*{prev(7)}"}[i], None
    fmt = "0.0%" if isinstance(value, float) and value < 1 else "#,##0"
    bold = kind == "formula"
    put(ws, f"A{r}", label, f=font(bold=bold))
    put(ws, f"B{r}", cell_value, f=font(bold=bold), fill=fill, fmt=fmt, h="right")
    ws.merge_cells(f"C{r}:E{r}")
    put(ws, f"C{r}", note)
    for col in "DE":
        ws[f"{col}{r}"].border = BOX
r += 2
size_rows = [FACT[("Nigeria cybersecurity market", "2025")], FACT[("Nigeria cybersecurity market", "2026")],
             FACT[("Nigeria cybersecurity market (forecast)", "2031")]]
# the chart's own small table: year and size
section(ws, r, "Market size, US$ million (Mordor Intelligence)", "E")
r += 1
heads(ws, r, "AB", ("Year", "US$ million"), 22)
chart_first = r + 1
for (year, _value), ref in zip(D.MARKET_SIZE, size_rows):
    r += 1
    put(ws, f"A{r}", year)
    put(ws, f"B{r}", f"={ref}", fmt="#,##0", h="right")
bars(ws, f"A{r + 2}", "Nigeria Cybersecurity Market (US$ million)",
     Reference(ws, min_col=2, min_row=chart_first, max_row=r), Reference(ws, min_col=1, min_row=chart_first, max_row=r),
     "#,##0", 500, [2], height=6)

# ================================================================ 2. Google Search
ws = wb.create_sheet("Google Search")
page(ws, {"A": 34, "B": 52, "C": 40})
titles(ws, "C", "GOOGLE SEARCH: WHAT CUSTOMERS FIND",
       "Five searches a customer would make, run on 9 October 2026. Run them again on Google to add screenshots.")
r = 4
heads(ws, r, "ABC", ("Search", "What came up", "What it tells us"))
for query, came_up, means in D.SEARCHES:
    r += 1
    put(ws, f"A{r}", f"“{query}”", f=font(bold=True))
    put(ws, f"B{r}", came_up)
    put(ws, f"C{r}", means)
    ws.row_dimensions[r].height = 48
r += 2
section(ws, r, "Prices found", "C")
r += 1
heads(ws, r, "ABC", ("Provider", "Offer", "Price and length"), 22)
for provider, offer, price, length, key in D.PRICES:
    r += 1
    put(ws, f"A{r}", provider, f=font(bold=True))
    put(ws, f"B{r}", offer)
    put(ws, f"C{r}", f"{price} ({length.lower()})")
    ws.row_dimensions[r].height = 22
r += 1
put(ws, f"A{r}", f"{D.PRODUCT} (proposed)", f=font(bold=True, color=TEAL))
put(ws, f"B{r}", "12-week beginner course: live classes and weekly labs", f=font(color=TEAL))
put(ws, f"C{r}", f"₦{D.PRICE:,} in {D.INSTALMENTS} parts", f=font(bold=True, color=TEAL))

# ================================================================ 3. Google Trends
ws = wb.create_sheet("Google Trends")
page(ws, {"A": 32, "B": 18, "C": 10, "D": 16, "E": 40})
titles(ws, "E", "GOOGLE TRENDS: SEARCH INTEREST IN DIGITAL SKILLS",
       f"{D.TRENDS_PERIOD}; all categories, web search (Market Research Lab Task 1 screenshots).")
r = 4
heads(ws, r, "ABCD", ("Search term", "Average interest (0-100)", "Rank", "Share of the five"))
t_first, t_last = r + 1, r + len(D.TRENDS)
for name, avg in D.TRENDS:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", avg, h="center")
    put(ws, f"C{r}", f"=RANK(B{r},$B${t_first}:$B${t_last})", h="center")
    put(ws, f"D{r}", f"=B{r}/SUM($B${t_first}:$B${t_last})", fmt="0%", h="center")
r += 2
section(ws, r, "What Google Trends shows", "E")
for label, text in D.TRENDS_POINTS + [("Free courses: ", D.TRENDS_OTHER.split(": ", 1)[1])]:
    r += 1
    put(ws, f"A{r}", label.strip(" :"), f=font(bold=True))
    ws.merge_cells(f"B{r}:E{r}")
    put(ws, f"B{r}", text)
    for col in "CDE":
        ws[f"{col}{r}"].border = BOX
    ws.row_dimensions[r].height = 32
bars(ws, f"A{r + 2}", "Average Search Interest, Nigeria, Past 12 Months (0-100)",
     Reference(ws, min_col=2, min_row=t_first, max_row=t_last), Reference(ws, min_col=1, min_row=t_first, max_row=t_last),
     "0", 50, [0])

# ================================================================ 4. Social Media
ws = wb.create_sheet("Social Media")
page(ws, {"A": 22, "B": 18, "C": 18, "D": 18, "E": 56})
titles(ws, "E", "SOCIAL MEDIA TRENDS",
       "Where the audience is in Nigeria, and how people talk about cybersecurity online.")
r = 4
heads(ws, r, "ABCDE", ("Platform", "Users (millions)", "Share of the population", "When", "Source"))
p_first, p_last = r + 1, r + len(D.PLATFORMS)
for name, users, share, when, key in D.PLATFORMS:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", users, fmt="0.0", h="center")
    put(ws, f"C{r}", share, fmt="0.0%", h="center")
    put(ws, f"D{r}", when, h="center")
    put(ws, f"E{r}", source(key))
r += 1
put(ws, f"A{r}", "Largest platform", f=font(bold=True))
put(ws, f"B{r}", f"=INDEX(A{p_first}:A{p_last},MATCH(MAX(B{p_first}:B{p_last}),B{p_first}:B{p_last},0))",
    f=font(bold=True), h="center")
r += 2
section(ws, r, "Social media trends", "E")
r += 1
heads(ws, r, "AB", ("Platform", "What we found"), 22)
ws.merge_cells(f"B{r}:E{r}")
for col in "CDE":
    ws[f"{col}{r}"].border = BOX
for name, text, key in D.SOCIAL_TRENDS:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    ws.merge_cells(f"B{r}:E{r}")
    put(ws, f"B{r}", f"{text} ({source(key)})")
    for col in "CDE":
        ws[f"{col}{r}"].border = BOX
    ws.row_dimensions[r].height = 32
bars(ws, f"A{r + 2}", "Users in Nigeria by Platform (millions)",
     Reference(ws, min_col=2, min_row=p_first, max_row=p_last), Reference(ws, min_col=1, min_row=p_first, max_row=p_last),
     "0.0", 50, [0], height=6)

# ================================================================ 5. Reviews
ws = wb.create_sheet("Reviews")
page(ws, {"A": 34, "B": 14, "C": 12, "D": 16, "E": 40, "F": 44})
titles(ws, "F", "COMPETITOR AND MARKETPLACE REVIEWS",
       "What learners say about the rivals and about the top cybersecurity courses on Udemy and Coursera.")
r = 4
section(ws, r, "Competitor reviews", "F")
r += 1
heads(ws, r, "ABCDEF", ("Competitor", "Rating (of 5)", "Offer", "Rating from", "What learners like",
                        "What learners complain about"))
c_first = r + 1
for name, offer, rating, where, like, dislike, _keys in D.COMPETITORS:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", rating if rating is not None else "None", fmt="0.0", h="center")
    put(ws, f"C{r}", offer)
    put(ws, f"D{r}", where)
    put(ws, f"E{r}", like)
    put(ws, f"F{r}", dislike)
    ws.row_dimensions[r].height = 46
c_last = r
r += 1
put(ws, f"A{r}", "Average rating of the rated rivals", f=font(bold=True))
put(ws, f"B{r}", f"=AVERAGE(B{c_first}:B{c_last})", f=font(bold=True), fmt="0.0", h="center")
r += 1
put(ws, f"A{r}", "Rivals with no rating found", f=font(bold=True))
put(ws, f"B{r}", f'=COUNTIF(B{c_first}:B{c_last},"None")', f=font(bold=True), h="center")
r += 2
section(ws, r, "Marketplace reviews (Udemy and Coursera)", "F")
r += 1
heads(ws, r, "ABCDEF", ("Course", "Rating (of 5)", "Platform", "Learners", "What learners like",
                        "What learners complain about"))
m_first = r + 1
for name, platform, rating, count, learners, like, dislike, _key in D.MARKETPLACE:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", rating, fmt="0.0", h="center")
    put(ws, f"C{r}", platform, h="center")
    put(ws, f"D{r}", learners, fmt="#,##0", h="right")
    put(ws, f"E{r}", like)
    put(ws, f"F{r}", f"{dislike} ({count:,} ratings)")
    ws.row_dimensions[r].height = 32
m_last = r
r += 1
put(ws, f"A{r}", "Total learners on these courses", f=font(bold=True))
put(ws, f"D{r}", f"=SUM(D{m_first}:D{m_last})", f=font(bold=True), fmt="#,##0", h="right")
r += 1
put(ws, f"A{r}", "Average rating", f=font(bold=True))
put(ws, f"B{r}", f"=AVERAGE(B{m_first}:B{m_last})", f=font(bold=True), fmt="0.0", h="center")
r += 1
ws.merge_cells(f"A{r}:F{r}")
put(ws, f"A{r}", D.MARKETPLACE_NOTE + " Complaints that come up again and again in low course reviews: "
    + "; ".join(c.lower() for c in D.COMMON_COMPLAINTS) + " (Course.careers).", f=font(9, italic=True, color="595959"),
    bd=None)
ws.row_dimensions[r].height = 30
bars(ws, f"A{r + 2}", "Learners on Each Course",
     Reference(ws, min_col=4, min_row=m_first, max_row=m_last), Reference(ws, min_col=1, min_row=m_first, max_row=m_last),
     "#,##0", 1_600_000, [0], height=6.5)

# ================================================================ 6. Gaps & Opportunities
ws = wb.create_sheet("Gaps & Opportunities")
page(ws, {"A": 32, "B": 50, "C": 38, "D": 10, "E": 10, "F": 12, "G": 8})
titles(ws, "G", "MARKET GAPS AND OPPORTUNITIES",
       "Gaps: what customers want but do not get today. Demand and fit are scored 1 (low) to 5 (high); "
       "priority = demand × fit (out of 25).")
r = 4
heads(ws, r, "ABCDEFG", ("Gap in the market", "Evidence from the research", f"Opportunity for {D.PRODUCT}",
                         "Demand (1-5)", "Fit (1-5)", "Priority (of 25)", "Rank"))
g_first, g_last = r + 1, r + len(D.GAPS)
for gap, evidence, opportunity, demand, fit in D.GAPS:
    r += 1
    put(ws, f"A{r}", gap, f=font(bold=True))
    put(ws, f"B{r}", evidence)
    put(ws, f"C{r}", opportunity)
    put(ws, f"D{r}", demand, fill=YELLOW, h="center")
    put(ws, f"E{r}", fit, fill=YELLOW, h="center")
    put(ws, f"F{r}", f"=D{r}*E{r}", f=font(bold=True), h="center")
    put(ws, f"G{r}", f"=RANK(F{r},$F${g_first}:$F${g_last})", h="center")
    ws.row_dimensions[r].height = 46
r += 1
put(ws, f"A{r}", "Biggest gap", f=font(bold=True, color="FFFFFF"), fill=TEAL)
ws.merge_cells(f"B{r}:G{r}")
put(ws, f"B{r}", f'=INDEX(A{g_first}:A{g_last},MATCH(MAX(F{g_first}:F{g_last}),F{g_first}:F{g_last},0))&'
                 f'" (priority "&MAX(F{g_first}:F{g_last})&" of 25)"', f=font(12, bold=True))
for col in "CDEFG":
    ws[f"{col}{r}"].border = BOX
r += 2
section(ws, r, "Opportunities in the market", "G")
r += 1
heads(ws, r, "ABC", ("Opportunity", "Evidence", "Source"), 22)
for name, evidence, key in D.OPPORTUNITIES:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    put(ws, f"B{r}", evidence)
    put(ws, f"C{r}", source(key))
    ws.row_dimensions[r].height = 32
r += 2
section(ws, r, f"How {D.PRODUCT} fills the gaps (marketing mix)", "G")
for name, text in D.MIX:
    r += 1
    put(ws, f"A{r}", name, f=font(bold=True))
    ws.merge_cells(f"B{r}:G{r}")
    put(ws, f"B{r}", text)
    for col in "CDEFG":
        ws[f"{col}{r}"].border = BOX
    ws.row_dimensions[r].height = 30
bars(ws, f"A{r + 2}", "Priority of Each Market Gap (out of 25)",
     Reference(ws, min_col=6, min_row=g_first, max_row=g_last), Reference(ws, min_col=1, min_row=g_first, max_row=g_last),
     "0", 25, [i for i, g in enumerate(D.GAPS) if g[3] * g[4] == max(x[3] * x[4] for x in D.GAPS)], height=7)
ws.freeze_panes = "A5"

# ================================================================ 7. Sources
ws = wb.create_sheet("Sources")
page(ws, {"A": 6, "B": 36, "C": 60, "D": 30})
titles(ws, "D", "SOURCES", f"Web sources found on {D.RESEARCH_DATE}; figures are as the sources report them. "
                           "Click a link to open the source.")
r = 4
heads(ws, r, "ABCD", ("No.", "Who", "What", "Link"), 22)
for i, (who, what, link) in enumerate(D.SOURCES.values(), start=1):
    r += 1
    put(ws, f"A{r}", i, h="center")
    put(ws, f"B{r}", who, f=font(bold=True))
    put(ws, f"C{r}", what)
    c = put(ws, f"D{r}", link.split("/")[2].removeprefix("www."), f=font(color="0F766E", underline="single"))
    c.hyperlink = link
    ws.row_dimensions[r].height = 30

wb.save(sys.argv[1])
print(f"wrote {sys.argv[1]}")
