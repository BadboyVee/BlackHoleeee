#!/usr/bin/env python3
"""Build Customer_Analysis.xlsx: capstone part 2, the customer analysis for CyberStart.

Sheets: World Data (world and Nigerian figures, and what the price means against wages),
Customer Needs (twelve needs scored; the five most important found with formulas), Segments (four
segments on demographic, geographic, psychographic and behavioural bases), Segment Scores
(weighted SUMPRODUCT, RANK, the target segment by INDEX/MATCH), Persona, Journey Map (with the
feeling curve) and Sources. Scores and weights are in yellow input cells. Charcoal headings and
gold leading bars, to match the deck; only Excel 2013 functions.

Usage: python build_customer_workbook.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import Marker
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Font as DrawingFont, Paragraph, ParagraphProperties, RegularTextRun
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList

import customer_data as C
from office2013 import THEME_XML
from xl_kit import BOX, YELLOW, bars, font, heads, page, put, rich, section, set_style, titles

CHARCOAL, GOLD, DARK_GOLD = "202124", "F2A900", "8A6100"
set_style(head=CHARCOAL, title=CHARCOAL, top=GOLD, bar="C8C8C4")


def source(key):
    return C.SOURCES[key][0]


def merged(ws, row, first, last, value, **fmt):
    ws.merge_cells(f"{first}{row}:{last}{row}")
    cell = put(ws, f"{first}{row}", value, **fmt)
    for col in range(ord(first) + 1, ord(last) + 1):
        ws[f"{chr(col)}{row}"].border = BOX
    return cell


wb = Workbook()
wb.loaded_theme = THEME_XML.encode()
wb._fonts = IndexedList([font()])
wb._named_styles["Normal"].font = font()
wb.properties.creator = wb.properties.lastModifiedBy = ""
wb.properties.title = "Capstone Part 2: Customer Analysis"
wb.properties.subject = f"Customer needs, segments, persona and journey map for {C.PRODUCT} by {C.BUSINESS}"

# ================================================================ 1. World Data
ws = wb.active
ws.title = "World Data"
page(ws, {"A": 46, "B": 40, "C": 12, "D": 46})
titles(ws, "D", f"CAPSTONE PART 2: CUSTOMER ANALYSIS FOR {C.PRODUCT.upper()}",
       f"World and Nigerian data behind the customer needs and segments. Found on {C.RESEARCH_DATE}.")
r = 4
heads(ws, r, "ABCD", ("Fact", "Figure", "Where", "Source"), 22)
for fact, figure, _number, where, key in C.WORLD:
    r += 1
    put(ws, f"A{r}", fact, f=font(bold=True))
    put(ws, f"B{r}", figure)
    put(ws, f"C{r}", where, h="center")
    put(ws, f"D{r}", source(key))
    ws.row_dimensions[r].height = 30
r += 2
section(ws, r, "What the price means for customers (change the yellow cells)", "D")
r += 1
heads(ws, r, "AB", ("Item", "Figure"), 22)
rows = [("Course price (₦)", C.PRICE, "#,##0", True), ("Number of instalments", C.INSTALMENTS, "0", True),
        ("National minimum wage a month (₦)", C.MIN_WAGE, "#,##0", True),
        ("Junior analyst pay a month, low end (₦)", C.JUNIOR_PAY_LOW, "#,##0", True)]
first = r + 1
for label, value, fmt, is_input in rows:
    r += 1
    put(ws, f"A{r}", label)
    put(ws, f"B{r}", value, fill=YELLOW, fmt=fmt, h="right")
price, parts, wage, pay = (f"B{first + i}" for i in range(4))
for label, formula, fmt in (("Each instalment (₦)", f"={price}/{parts}", "#,##0"),
                            ("Months of minimum wage to pay for the course", f"=ROUND({price}/{wage},1)", "0.0"),
                            ("One instalment as a share of the minimum wage", f"={price}/{parts}/{wage}", "0%"),
                            ("Months of junior analyst pay to earn the price back", f"=ROUND({price}/{pay},1)", "0.0")):
    r += 1
    put(ws, f"A{r}", label, f=font(bold=True))
    put(ws, f"B{r}", formula, f=font(bold=True), fmt=fmt, h="right")
r += 1
merged(ws, r, "A", "D", "So the course is out of reach on the minimum wage, but a junior analyst job earns the price "
       "back in about a month: the buyers are people with some income who want a better-paid job.",
       f=font(italic=True, color="404040"), bd=None)
ws.row_dimensions[r].height = 32

# ================================================================ 2. Customer Needs
ws = wb.create_sheet("Customer Needs")
page(ws, {"A": 12, "B": 34, "C": 30, "D": 56, "E": 12, "F": 12, "G": 13, "H": 9})
titles(ws, "H", "CUSTOMER NEEDS: FUNCTIONAL, ECONOMIC, EMOTIONAL AND SOCIAL",
       "Frequency = how often the need came up in the research; impact = its effect on the decision to buy "
       "(1 to 5). Importance = frequency × impact (of 25).")
ws.row_dimensions[2].height = 30
r = 4
heads(ws, r, "ABCDEFGH", ("Type", "Need", "What it means", "Evidence from the research", "Frequency (1-5)",
                          "Impact (1-5)", "Importance (of 25)", "Rank"), 32)
n_first, n_last = r + 1, r + len(C.NEEDS)
for kind, need, _short, means, evidence, freq, impact in C.NEEDS:
    r += 1
    put(ws, f"A{r}", kind, h="center")
    put(ws, f"B{r}", need, f=font(bold=True))
    put(ws, f"C{r}", means)
    put(ws, f"D{r}", evidence)
    put(ws, f"E{r}", freq, fill=YELLOW, h="center")
    put(ws, f"F{r}", impact, fill=YELLOW, h="center")
    put(ws, f"G{r}", f"=E{r}*F{r}", f=font(bold=True), h="center")
    # Rank without ties: RANK, plus one for each equal score higher up the list.
    put(ws, f"H{r}", f"=RANK(G{r},$G${n_first}:$G${n_last})+COUNTIF($G${n_first}:G{r - 1},G{r})" if r > n_first
        else f"=RANK(G{r},$G${n_first}:$G${n_last})", h="center")
    ws.row_dimensions[r].height = 34
r += 2
section(ws, r, "The 5 most important needs (found by formula)", "H")
r += 1
heads(ws, r, "ABCD", ("Rank", "Need", "Type", "Importance (of 25)"), 22)
TOP5 = r + 1
for k in range(1, 6):
    r += 1
    match = f"MATCH({k},$H${n_first}:$H${n_last},0)"
    put(ws, f"A{r}", k, h="center")
    put(ws, f"B{r}", f"=INDEX($B${n_first}:$B${n_last},{match})", f=font(bold=True))
    put(ws, f"C{r}", f"=INDEX($A${n_first}:$A${n_last},{match})", h="center")
    put(ws, f"D{r}", f"=INDEX($G${n_first}:$G${n_last},{match})", h="center")
r += 2
section(ws, r, "Needs by type: average importance", "H")
r += 1
heads(ws, r, "AB", ("Type", "Average importance"), 22)
for kind in C.NEED_TYPES:
    r += 1
    put(ws, f"A{r}", kind, f=font(bold=True))
    put(ws, f"B{r}", f'=AVERAGEIF($A${n_first}:$A${n_last},"{kind}",$G${n_first}:$G${n_last})', fmt="0.0", h="center")
# The chart's own labels: the short names.
importance = [f * i for *_x, f, i in C.NEEDS]
order = sorted(range(len(C.NEEDS)), key=lambda i: (-importance[i], i))[:5]
r += 2
section(ws, r, "Importance of each need (chart data)", "H")
r += 1
heads(ws, r, "AB", ("Need", "Importance"), 22)
c_first = r + 1
for i, (_kind, _need, short, *_rest) in enumerate(C.NEEDS):
    r += 1
    put(ws, f"A{r}", short)
    put(ws, f"B{r}", f"=G{n_first + i}", h="center")
bars(ws, f"A{r + 2}", "Importance of Each Need (of 25)", Reference(ws, min_col=2, min_row=c_first, max_row=r),
     Reference(ws, min_col=1, min_row=c_first, max_row=r), "0", 25, order, width=17, height=9)
ws.freeze_panes = "A5"

# ================================================================ 3. Segments
ws = wb.create_sheet("Segments")
page(ws, {"A": 18, "B": 34, "C": 34, "D": 34, "E": 34})
titles(ws, "E", "FOUR MARKET SEGMENTS",
       "Built on four bases: demographic, geographic, psychographic and behavioural.")
r = 4
heads(ws, r, "ABCDE", ["Base"] + [s[0] for s in C.SEGMENTS], 24)
for b, (base, about) in enumerate(C.BASES):
    r += 1
    put(ws, f"A{r}", base, f=font(bold=True))
    for col, seg in zip("BCDE", C.SEGMENTS):
        put(ws, f"{col}{r}", seg[1 + b])
    ws.row_dimensions[r].height = 48
r += 1
put(ws, f"A{r}", "Size evidence", f=font(bold=True))
for col, seg in zip("BCDE", C.SEGMENTS):
    put(ws, f"{col}{r}", seg[5], f=font(italic=True))
ws.row_dimensions[r].height = 34
r += 2
section(ws, r, "The four bases", "E")
for base, about in C.BASES:
    r += 1
    put(ws, f"A{r}", base, f=font(bold=True))
    merged(ws, r, "B", "E", about)

# ================================================================ 4. Segment Scores
ws = wb.create_sheet("Segment Scores")
page(ws, {"A": 26, "B": 16, "C": 16, "D": 17, "E": 16, "F": 16, "G": 9})
titles(ws, "G", "SEGMENT SCORES",
       "Each criterion scored 1 (poor) to 5 (best); for competition, 5 means little competition. "
       "Weighted score = SUMPRODUCT of the weights and the scores (of 5).")
ws.row_dimensions[2].height = 30
r = 4
heads(ws, r, "ABCDEFG", ["Segment"] + [c for c, _ in C.CRITERIA] + ["Weighted score", "Rank"], 30)
r += 1
WEIGHTS = r
put(ws, f"A{r}", "Weight", f=font(bold=True))
for col, (_c, w) in zip("BCDE", C.CRITERIA):
    put(ws, f"{col}{r}", w, f=font(bold=True), fill=YELLOW, fmt="0%", h="center")
put(ws, f"F{r}", f"=SUM(B{r}:E{r})", f=font(bold=True), fmt="0%", h="center")
put(ws, f"G{r}", None)
s_first, s_last = r + 1, r + len(C.SEGMENTS)
for seg, scores in zip(C.SEGMENTS, C.SCORES):
    r += 1
    put(ws, f"A{r}", seg[0], f=font(bold=True))
    for col, v in zip("BCDE", scores):
        put(ws, f"{col}{r}", v, fill=YELLOW, h="center")
    put(ws, f"F{r}", f"=SUMPRODUCT($B${WEIGHTS}:$E${WEIGHTS},B{r}:E{r})", f=font(bold=True), fmt="0.00", h="center")
    put(ws, f"G{r}", f"=RANK(F{r},$F${s_first}:$F${s_last})", h="center")
    ws.row_dimensions[r].height = 22
r += 2
put(ws, f"A{r}", "Target segment", f=font(bold=True, color="FFFFFF"), fill=CHARCOAL)
merged(ws, r, "B", "G", f'=INDEX(A{s_first}:A{s_last},MATCH(MAX(F{s_first}:F{s_last}),F{s_first}:F{s_last},0))'
       f'&" ("&TEXT(MAX(F{s_first}:F{s_last}),"0.00")&" of 5)"', f=font(12, bold=True))
ws.row_dimensions[r].height = 22
TARGET = r
r += 2
section(ws, r, "Why each segment got its scores", "G")
r += 1
heads(ws, r, "ABCDE", ["Segment"] + [c for c, _ in C.CRITERIA], 22)
for seg, basis in zip(C.SEGMENTS, C.SCORE_BASIS):
    r += 1
    put(ws, f"A{r}", seg[0], f=font(bold=True))
    for col, text in zip("BCDE", basis):
        put(ws, f"{col}{r}", text)
    ws.row_dimensions[r].height = 46
totals = [sum(w * v for (_c, w), v in zip(C.CRITERIA, sc)) for sc in C.SCORES]
bars(ws, f"A{r + 2}", "Weighted Score of Each Segment (of 5)", Reference(ws, min_col=6, min_row=s_first, max_row=s_last),
     Reference(ws, min_col=1, min_row=s_first, max_row=s_last), "0.00", 5, [totals.index(max(totals))], height=6.5)

# ================================================================ 5. Persona
ws = wb.create_sheet("Persona")
page(ws, {"A": 22, "B": 70}, landscape=False)
P = C.PERSONA
titles(ws, "B", f"CUSTOMER PERSONA: {P['name'].upper()}",
       f"A fictional customer from the target segment ({P['segment']}), built from the research.")
r = 4
for label, text in P["facts"]:
    put(ws, f"A{r}", label, f=font(bold=True))
    put(ws, f"B{r}", text)
    r += 1
put(ws, f"A{r}", "In their words", f=font(bold=True))
put(ws, f"B{r}", P["quote"], f=font(italic=True))
for title, items in (("Goals", P["goals"]), ("Frustrations", P["frustrations"]), ("Needs", P["needs"]),
                     ("Where to reach them", P["channels"])):
    r += 2
    section(ws, r, title, "B")
    for item in items:
        r += 1
        merged(ws, r, "A", "B", item)

# ================================================================ 6. Journey Map
ws = wb.create_sheet("Journey Map")
page(ws, {"A": 20, **{get_column_letter(2 + i): 30 for i in range(len(C.STAGES))}})
last = get_column_letter(1 + len(C.STAGES))
titles(ws, last, f"CUSTOMER JOURNEY MAP: {P['name'].upper()}",
       "From first hearing about cybersecurity to getting a job. Feeling: 1 (very negative) to 5 (very positive).")
r = 4
cols = [get_column_letter(2 + i) for i in range(len(C.STAGES))]
heads(ws, r, ["A"] + cols, ["Stage"] + C.STAGES, 24)
for label, cells in C.JOURNEY:
    r += 1
    put(ws, f"A{r}", label, f=font(bold=True))
    for col, text in zip(cols, cells):
        put(ws, f"{col}{r}", text)
    ws.row_dimensions[r].height = 48
r += 1
FEEL = r
put(ws, f"A{r}", "Feeling (1-5)", f=font(bold=True))
for col, v in zip(cols, C.FEELING):
    put(ws, f"{col}{r}", v, fill=YELLOW, h="center")
r += 1
put(ws, f"A{r}", "Lowest point", f=font(bold=True))
merged(ws, r, "B", last, f"=INDEX(B4:{last}4,MATCH(MIN(B{FEEL}:{last}{FEEL}),B{FEEL}:{last}{FEEL},0))",
       f=font(bold=True))
r += 1
put(ws, f"A{r}", "Highest point", f=font(bold=True))
merged(ws, r, "B", last, f"=INDEX(B4:{last}4,MATCH(MAX(B{FEEL}:{last}{FEEL}),B{FEEL}:{last}{FEEL},0))",
       f=font(bold=True))
# The feeling curve
ch = LineChart()
cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=1300, b=False, solidFill="404040")
ch.title = Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),
                                                    r=[RegularTextRun(rPr=cp, t="How Tunde Feels at Each Stage")])])),
                 overlay=False)
ch.add_data(Reference(ws, min_col=2, max_col=1 + len(C.STAGES), min_row=FEEL), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(ws, min_col=2, max_col=1 + len(C.STAGES), min_row=4))
ch.legend = None
s = ch.series[0]
s.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill=GOLD, w=31750))
s.marker = Marker(symbol="circle", size=9)
s.marker.graphicalProperties = GraphicalProperties(solidFill=CHARCOAL, ln=LineProperties(solidFill=CHARCOAL))
s.smooth = False
ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False, showLegendKey=False,
                              showPercent=False)
ch.dataLabels.position = "t"
ch.dataLabels.txPr = rich(1000, "262626", bold=True)
ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 5
ch.y_axis.majorUnit = 1
ch.y_axis.delete = True
ch.y_axis.majorGridlines = None
ch.x_axis.delete = False
ch.x_axis.txPr = rich(1000)
ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
ch.width, ch.height = 22, 7
ws.add_chart(ch, f"A{r + 2}")

# ================================================================ 7. Sources
ws = wb.create_sheet("Sources")
page(ws, {"A": 6, "B": 40, "C": 60, "D": 30})
titles(ws, "D", "SOURCES", f"Web sources found on {C.RESEARCH_DATE} and in part 1; figures are as the sources report "
                           "them. Click a link to open the source.")
r = 4
heads(ws, r, "ABCD", ("No.", "Who", "What", "Link"), 22)
for i, (who, what, link) in enumerate(C.SOURCES.values(), start=1):
    r += 1
    put(ws, f"A{r}", i, h="center")
    put(ws, f"B{r}", who, f=font(bold=True))
    put(ws, f"C{r}", what)
    c = put(ws, f"D{r}", link.split("/")[2].removeprefix("www."), f=font(color=DARK_GOLD, underline="single"))
    c.hyperlink = link
    ws.row_dimensions[r].height = 30

wb.save(sys.argv[1])
print(f"wrote {sys.argv[1]}")
