#!/usr/bin/env python3
"""Build the capstone deck: customer research for CyberStart, SkillUp Academy's beginner
cybersecurity course, used to find market gaps and opportunities. Made the way it would be in
PowerPoint 2013, in the minimal teal design of ../../assignment/src/deck_kit.py.

Usage: python build_deck.py deck_data.json template.pptx wood.jpg STYLE out.pptx

Slides: 1. title, 2. the business and product, 3. defining the market, 4. the market in
numbers (bar chart), 5. research methods, 6. Google Search, 7. Google Trends (bar chart),
8. social media trends (bar chart), 9. competitor reviews, 10. marketplace reviews, 11. gaps in
the market, 12. the gaps ranked (bar chart), 13. market opportunities, 14. how CyberStart fills
the gaps, 15. conclusion, 16. thank you. The text comes from capstone_data.py, and the worked-out
figures from deck_data.json (export_deck_data.py reads them out of the finished workbook).

Every slide's objects come in by themselves (After Previous); the plan is saved next to
out.pptx as out.anim.json for finish_deck.py.
"""
import json
import sys

from pptx.util import Emu, Inches

import capstone_data as D
from deck_kit import Deck, placeholder

data_path, template_path, wood_path, style_name, out_path = sys.argv[1:6]
deck = Deck(style_name, template_path, wood_path)
S, prs, W, H, LAYOUT = deck.S, deck.prs, deck.W, deck.H, deck.layout
data = json.load(open(data_path))
deck.footer = f"{D.BUSINESS}  ·  Capstone Project"
MINIMAL = bool(S.get("minimal"))
sizing = dict(data["sizing"])


def content_slide(title, subtitle):
    s = prs.slides.add_slide(LAYOUT["Title and Content"])
    top = deck.content(s, title, subtitle)
    deck.anim_chrome(s)
    return s, top


def table_slide(s, top, rows, widths, heights, size, *, centre_from=1, bold_rows=(), note=None):
    """A table under the subtitle (wiping down), and a note under it (fading in last)."""
    table_w = sum(widths, Emu(0))
    heights = heights if isinstance(heights, list) else [heights] * len(rows)
    frame, _ = deck.table(s, placeholder(s, 1), (W - table_w) // 2, top + Inches(0.05), rows, widths, heights, size,
                          centre_from=centre_from, bold_rows=bold_rows)
    deck.anim(s, frame, 3, "wipe-down")
    if note:
        deck.anim(s, deck.footnote(s, top + Inches(0.05) + sum(heights, Emu(0)) + Inches(0.12), note), 4)


def chart_slide(s, top, labels, values, number_format, maximum, top_bars, points, size=17):
    frame, _ = deck.bar_chart(s, placeholder(s, 1), *deck.chart_area(top), labels, values,
                              number_format=number_format, maximum=maximum, top=top_bars)
    deck.anim(s, frame, 3, "wipe-left")
    return deck.side_note(s, top, [[(a, True), (b, False)] for a, b in points], size=size)


def bold_first(points):
    return [[(a, True), (b, False)] for a, b in points]


# ================================================================== 1. Title
s = prs.slides.add_slide(LAYOUT["Title Slide"])
deck.title_card(s, "Capstone Project", [f"Market gaps and opportunities for {D.PRODUCT}, a beginner cybersecurity "
                                        "course", f"{D.BUSINESS}  ·  Market Research Lab"],
                size=54 if MINIMAL else 44)
deck.anim_chrome(s, title_card=True)
s.notes_slide.notes_text_frame.text = (
    f"My capstone project is customer research for {D.PRODUCT}, a beginner cybersecurity course from {D.BUSINESS}. "
    "I used Google Search, Google Trends, social media trends, competitor reviews and marketplace reviews to find "
    "the gaps and opportunities in the market.")

# ================================================================== 2. The business and product
s, top = content_slide("The Business and Product", "What I plan to launch")
deck.anim(s, deck.boxed_text(s, top, [[(f"{label}: ", True), (text, False)] for label, text in D.PRODUCT_POINTS],
                             size=22), 3)
s.notes_slide.notes_text_frame.text = (
    f"{D.BUSINESS} is a digital-skills training business. Its product, {D.PRODUCT}, is a 12-week course for "
    "beginners, taught online with weekend labs in Lagos and Abuja, that prepares people for a first certificate "
    "and a first job.")

# ================================================================== 3. Defining the market
s, top = content_slide("Defining the Market", "Who the customers are, where they are and what they need")
rows = [("Part", "Our market")] + [([[(a, True)]], b) for a, b in D.MARKET_DEFINITION]
table_slide(s, top, rows, [Inches(2.6), Inches(9.2)], [Inches(0.55)] + [Inches(0.72)] * len(D.MARKET_DEFINITION),
            18, centre_from=9)
s.notes_slide.notes_text_frame.text = (
    "The market is beginner cybersecurity training in Nigeria. The customers are young people who want a "
    "cybersecurity job, and the employers who pay to train their staff.")

# ================================================================== 4. The market in numbers
s, top = content_slide("The Market in Numbers", "Nigeria's cybersecurity market, US$ million (Mordor Intelligence)")
note = chart_slide(s, top, [y for y, _ in D.MARKET_SIZE], [v for _, v in D.MARKET_SIZE], "#,##0", 500, [2], [
    ("Growing: ", f"about {data['growth'] * 100:.0f}% a year to 2031."),
    ("Few experts: ", f"8,352 in Nigeria (2023), {data['experts_ratio']:.0f} times fewer than South Africa."),
    ("More attacks: ", "4,388 a week on each organisation in early 2025, up 47%."),
    ("Online: ", "109 million Nigerians use the internet."),
], size=16)
deck.anim(s, note, 4)
s.notes_slide.notes_text_frame.text = (
    f"Nigeria's cybersecurity market was worth about 230 million US dollars in 2025 and is forecast to reach about "
    f"415 million by 2031, about {data['growth'] * 100:.0f} percent a year. Attacks are rising, but the country has "
    "very few trained experts.")

# ================================================================== 5. Research methods
s, top = content_slide("Customer Research Methods", "Five ways I studied the customers and the competition")
rows = [("Method", "Where", "What I looked for")] + [([[(m, True)]], w, x) for m, w, x in D.METHODS]
table_slide(s, top, rows, [Inches(2.7), Inches(4.7), Inches(4.4)], [Inches(0.5)] + [Inches(0.78)] * len(D.METHODS),
            15, centre_from=9, note=f"Web research done on {D.RESEARCH_DATE}; Google Trends screenshots of "
                                    "28 September 2026.")
s.notes_slide.notes_text_frame.text = (
    "I used five methods: Google Search, Google Trends, social media trends, competitor reviews and marketplace "
    "reviews on Udemy and Coursera.")

# ================================================================== 6. Google Search
s, top = content_slide("Google Search", "What a customer finds when searching for cybersecurity courses")
deck.anim(s, deck.boxed_text(s, top, bold_first(D.SEARCH_POINTS), size=20), 3)
s.notes_slide.notes_text_frame.text = (
    "I searched for course fees, free training, how to start a career, and jobs. Naira prices are hard to find and "
    "vary a lot, free programmes fill fast, and every guide tells beginners to practise in labs and get a first "
    "certificate.")

# ================================================================== 7. Google Trends
s, top = content_slide("Google Trends", "Average search interest in Nigeria over the past 12 months (0 to 100)")
interest = [v for _, v in D.TRENDS]
deck.anim(s, chart_slide(s, top, [n for n, _ in D.TRENDS], interest, "0", 50, [0], D.TRENDS_POINTS, size=16), 4)
s.notes_slide.notes_text_frame.text = (
    "Cybersecurity is the most searched digital skill in Nigeria, with an average of about 43. It reached 100 in "
    "April 2026 when the CAC confirmed a cyber breach, and it led in most states.")

# ================================================================== 8. Social media
s, top = content_slide("Social Media Trends", "Users in Nigeria by platform, millions (NapoleonCat, 2026)")
deck.anim(s, chart_slide(s, top, [p[0] for p in D.PLATFORMS], [round(p[1], 1) for p in D.PLATFORMS], "0.0", 50, [0],
                         D.SOCIAL_POINTS, size=16), 4)
s.notes_slide.notes_text_frame.text = (
    "Facebook is the biggest platform in Nigeria, LinkedIn reaches professionals and employers, and cybersecurity "
    "content is popular on Instagram and TikTok. Learners also complain in public on X, so service matters.")

# ================================================================== 9. Competitor reviews
s, top = content_slide("Competitor Reviews", "What learners say about the rivals (Trustpilot and the news)")
rows = [("Competitor", "Rating", "Learners like", "Learners complain about")]
rows += [([[(name, True)]], f"{rating:.1f}" if rating else "None", like, dislike)
         for name, _offer, rating, _where, like, dislike, _keys in D.COMPETITORS]
table_slide(s, top, rows, [Inches(2.9), Inches(1.1), Inches(3.7), Inches(4.2)],
            [Inches(0.45)] + [Inches(0.78)] * len(D.COMPETITORS), 13, centre_from=1,
            note="Ratings out of 5 on Trustpilot: ALX Africa from 544 reviews, TryHackMe from 867.")
s.notes_slide.notes_text_frame.text = (
    "The rated rivals score well, but learners complain about slow support, billing problems, missing lectures "
    "and weak grading. Free programmes are very hard to get into, and classroom providers have almost no reviews "
    "online.")

# ================================================================== 10. Marketplace reviews
s, top = content_slide("Marketplace Reviews", "The top cybersecurity courses on Coursera and Udemy")
rows = [("Course", "Rating", "Learners", "Learners like", "Learners complain about")]
rows += [([[(name, True)], [(platform, False)]], f"{rating:.1f}", f"{learners:,}", like, dislike)
         for name, platform, rating, _count, learners, like, dislike, _key in D.MARKETPLACE]
table_slide(s, top, rows, [Inches(4.1), Inches(1.0), Inches(1.5), Inches(2.5), Inches(2.8)],
            [Inches(0.45)] + [Inches(0.85)] * len(D.MARKETPLACE), 13, centre_from=1,
            note="Common complaints in low reviews: " + "; ".join(c.lower() for c in D.COMMON_COMPLAINTS) + ".")
s.notes_slide.notes_text_frame.text = (
    f"The top online courses average {data['course_rating']:.1f} stars and have over "
    f"{data['course_learners'] / 1e6:.1f} million learners between them, so demand is real. Their low reviews "
    "complain about outdated content, too much theory and weak support.")

# ================================================================== 11. Gaps in the market
s, top = content_slide("Gaps in the Market", "What customers want but do not get today")
rows = [("Gap", "Evidence", f"Opportunity for {D.PRODUCT}")]
rows += [([[(gap, True)]], evidence, opportunity) for gap, evidence, opportunity, _d, _f in D.GAPS]
table_slide(s, top, rows, [Inches(2.8), Inches(5.6), Inches(3.5)], [Inches(0.42)] + [Inches(0.6)] * len(D.GAPS), 12,
            centre_from=9)
s.notes_slide.notes_text_frame.text = (
    "Putting the research together, I found seven gaps. The biggest are affordable courses with live support and "
    "hands-on practice, then a route from course to job and reliable support.")

# ================================================================== 12. The gaps ranked
s, top = content_slide("The Gaps, Ranked", "Priority = demand × fit, each scored 1 to 5 (out of 25)")
gaps = data["gaps"]
best = max(g["priority"] for g in gaps)
top_gaps = [g["gap"] for g in gaps if g["priority"] == best]
deck.anim(s, chart_slide(s, top, [g["gap"] for g in gaps], [g["priority"] for g in gaps], "0", 25,
                         [i for i, g in enumerate(gaps) if g["priority"] == best], [
    ("Biggest gaps: ", f"{top_gaps[0].lower()} and {top_gaps[1].lower()} ({best} of 25)." if len(top_gaps) == 2
     else f"{top_gaps[0].lower()} ({best} of 25)."),
    ("Next: ", "a route from course to job, and reliable support."),
    (f"{D.PRODUCT} ", "is built around these gaps."),
], size=17), 4)
s.notes_slide.notes_text_frame.text = (
    "I scored each gap for how much customers want it and how well it fits SkillUp Academy. Affordable courses "
    "with live support and hands-on practice score highest.")

# ================================================================== 13. Market opportunities
s, top = content_slide("Market Opportunities", "Why now is a good time to launch")
rows = [("Opportunity", "Evidence")] + [([[(name, True)]], evidence) for name, evidence, _key in D.OPPORTUNITIES]
table_slide(s, top, rows, [Inches(3.2), Inches(8.6)], [Inches(0.45)] + [Inches(0.66)] * len(D.OPPORTUNITIES), 16,
            centre_from=9)
s.notes_slide.notes_text_frame.text = (
    "The market is growing, experts are scarce, attacks are rising, employers in banking pay for training, interest "
    "is nationwide, and the free ISC2 exam offer has ended, so learners need exam preparation.")

# ================================================================== 14. How CyberStart fills the gaps
s, top = content_slide(f"How {D.PRODUCT} Fills the Gaps", "The marketing mix")
rows = [("The 4 Ps", f"{D.PRODUCT}")] + [([[(name, True)]], text) for name, text in D.MIX]
table_slide(s, top, rows, [Inches(2.4), Inches(9.4)], [Inches(0.5)] + [Inches(0.85)] * len(D.MIX), 18, centre_from=9,
            note=f"Year 1 estimate: {sizing['Learners in year 1']:,.0f} learners and "
                 f"₦{sizing['Revenue in year 1 (₦)'] / 1e6:,.2f} million (see the Market sheet in Excel).")
s.notes_slide.notes_text_frame.text = (
    f"{D.PRODUCT} answers each gap: weekly labs, a fair price in parts, published online, fast support and job help. "
    f"In year one I estimate {sizing['Learners in year 1']:,.0f} learners.")

# ================================================================== 15. Conclusion
s, top = content_slide("Conclusion", "What the research shows")
deck.anim(s, deck.boxed_text(s, top, bold_first(D.CONCLUSION), size=20), 3)
s.notes_slide.notes_text_frame.text = (
    "In conclusion, the market is growing and customers are searching, but today's options leave clear gaps. "
    f"{D.PRODUCT} can win by filling them.")

# ================================================================== 16. Thank you
s = prs.slides.add_slide(LAYOUT["Title Slide"])
deck.title_card(s, "Thank You", ["Any questions?", f"{D.BUSINESS}  ·  Capstone Project"])
deck.anim_chrome(s, title_card=True)
s.notes_slide.notes_text_frame.text = "Thank you for listening."

deck.save(out_path, title=f"Capstone Project: {D.PRODUCT} by {D.BUSINESS}",
          subject="Customer research to find market gaps and opportunities",
          keywords="market research, Google Trends, reviews, market gaps, opportunities")
