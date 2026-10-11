#!/usr/bin/env python3
"""Build Presenter_Guide.docx: a plain Word guide to presenting the capstone part 2 deck.

It has a checklist for making the work your own before you hand it in, a slide-by-slide guide
(what each slide shows, what to say, and a question your teacher may ask, with an answer), and a
practice quiz with answers at the end. The figures come from the finished Customer_Analysis.xlsx,
so the guide matches the deck. The guide is written as simple HTML and saved as a Word document
by LibreOffice.

Usage: python build_guide.py Customer_Analysis.xlsx out.docx
"""
import html
import os
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

import capstone_data as D
import customer_data as C

workbook, out = sys.argv[1:3]
wb = load_workbook(workbook, data_only=True)
needs = {r[1].value: (r[6].value, r[7].value) for r in wb["3 Customer Needs"].iter_rows(min_row=5, max_row=16)}
seg_ws = wb["5-6 Segment Scores"]
totals = {}
for row in seg_ws.iter_rows():                          # the score table (the first time each name appears)
    if row[0].value in [s[0] for s in C.SEGMENTS] and row[0].value not in totals:
        totals[row[0].value] = row[5].value
if None in [v for v, _ in needs.values()] or None in totals.values() or len(totals) != len(C.SEGMENTS):
    sys.exit("run the workbook through LibreOffice first")
top5 = [n for n, _ in sorted(needs.items(), key=lambda kv: kv[1][1])[:5]]
target = max(totals, key=totals.get)
ranked = sorted(totals, key=totals.get, reverse=True)
weights = {c: w for c, w in C.CRITERIA}
t_scores = dict(zip([c for c, _ in C.CRITERIA], C.SCORES[[s[0] for s in C.SEGMENTS].index(target)]))
working = " + ".join(f"{t_scores[c]} × {weights[c]:.2f}" for c in t_scores)
low = C.STAGES[C.FEELING.index(min(C.FEELING))]

e = html.escape
parts = []


def h(level, text):
    parts.append(f"<h{level}>{e(text)}</h{level}>")


def p(text, bold_lead=None):
    lead = f"<b>{e(bold_lead)}</b> " if bold_lead else ""
    parts.append(f"<p>{lead}{e(text)}</p>")


def items(rows, ordered=False):
    tag = "ol" if ordered else "ul"
    parts.append(f"<{tag}>" + "".join(f"<li>{e(r)}</li>" for r in rows) + f"</{tag}>")


# ------------------------------------------------------------------ the guide
h(1, "Presenter Guide: Customer Analysis")
p(f"Capstone Project, Part 2. {C.PRODUCT}, a beginner cybersecurity course by {C.BUSINESS}.")
p("Read this before you present. Practise each slide out loud until you can say it in your own words without "
  "reading.")

h(2, "Before you hand it in: make it yours")
items([
    "Check your course's rule on using AI tools. The Sources slide says: \"I used an AI assistant to help with "
    "research and slide design.\" Keep that line if AI help is allowed. If you are not sure, ask your teacher.",
    "Rewrite the words in your own style: the bullet points and the speaker notes (in PowerPoint, View > Notes).",
    "Check the scores. They are judgements, so change any you disagree with. In the Excel file they are the yellow "
    "cells: sheet 3 Customer Needs (columns E and F) and sheet 5-6 Segment Scores (the weights and the scores). The "
    "results in Excel update by themselves; then change the same numbers on the slides.",
    "Add one example of your own, for example someone you know who wanted to learn a tech skill.",
    "Know where each number comes from. Every source, with a link, is in the Sources sheet of the Excel file.",
], ordered=True)

h(2, "A tip for presenting")
p("Each slide moves on by itself after 10 seconds. When you present live, you may want more time: in PowerPoint, "
  "go to Slide Show > Set Up Slide Show and choose Advance slides: Manually. Then a click moves you on.")

h(2, "Slide by slide")
slides = [
    ("Customer Analysis (title)", "The topic and the product.",
     f"Good morning. My capstone part 2 is a customer analysis for {C.PRODUCT}, a beginner cybersecurity course.",
     "What is the product?", f"{C.PRODUCT}, a 12-week beginner cybersecurity course with weekend labs, for "
                            f"₦{D.PRICE:,} paid in {D.INSTALMENTS} parts."),
    ("Contents", "The six tasks in order.", "I will go through the six tasks in the order they were set.",
     "What are the six tasks?", "Persona, journey map, the 5 most important needs, four segments, segment scores, "
                                "and the target segment."),
    ("Key Facts", "World and Nigerian data behind the analysis.",
     "Cybersecurity skills are in demand worldwide, and Nigeria has a young population.",
     "Where does the 95% come from?", "ISC2's 2025 Cybersecurity Workforce Study, a survey of 16,029 cybersecurity "
                                      "professionals: 95% of teams have at least one skills gap."),
    ("1. Customer Persona", "Tunde, a made-up customer from the target segment.",
     "Tunde is 27, works in customer service in Lagos and wants a cybersecurity job within a year.",
     "Is Tunde a real person?", "No. A persona is a made-up customer, based on research, that stands for the "
                                "target segment."),
    ("1. Customer Persona (continued)", "What makes Tunde buy.",
     "News of breaches, a friend's new job or a free first class push Tunde to act; Tunde pays in three parts.",
     "What is Tunde's biggest worry?", "Paying for a course that turns out to be a scam."),
    ("2. Customer Journey Map", "The five stages: awareness, consideration, purchase, learning, after the course.",
     "At each stage the table shows what Tunde does, where, and what Tunde thinks.",
     "What is a touchpoint?", "A place where the customer meets the business, such as an advert, the website or a "
                              "WhatsApp chat."),
    ("2. Customer Journey Map (continued)", "The pain point at each stage and what the business can do.",
     "Every pain point is a chance to do better.",
     "Give one pain point and its fix.", "Hidden prices and few reviews when comparing courses; fix it by "
                                         "publishing prices and reviews."),
    ("2. How Tunde Feels", "A bar chart of Tunde's feeling at each stage (1 to 5).",
     f"Tunde feels worst at {low.lower()} and best after the course.",
     f"Why is {low.lower()} the lowest point?", "Tunde is comparing courses, can't find clear prices and is "
                                                "afraid of scams."),
    ("3. Customer Needs", "Three functional, economic, emotional and social needs.",
     "Functional needs are about what the course does; economic about money; emotional about feelings; social "
     "about other people.",
     "What is the difference between a functional and an emotional need?",
     "Functional is what the product does, like hands-on labs; emotional is how it makes the customer feel, like "
     "trust or feeling supported."),
    ("3. The 5 Most Important Needs", "The top five needs and their scores.",
     "I scored each need for how often it came up in my research and how much it affects buying.",
     "How did you score the needs?", "Two scores from 1 to 5, how often it came up and how much it affects "
                                     "buying, multiplied together, so the most is 25."),
    ("3. All Twelve Needs", "A bar chart of all twelve needs; the top five are dark blue.",
     "Skills and money matter most; social needs score lowest.",
     "Which needs scored lowest?", "The social needs, such as the status of a tech job."),
    ("4. Four Market Segments", "Four groups of customers, each described on the four bases.",
     "The segments are " + ", ".join(s[0] for s in C.SEGMENTS[:-1]) + f" and {C.SEGMENTS[-1][0]}.",
     "What are the four bases of segmentation?", "Demographic (age, job, income), geographic (where they live), "
                                                 "psychographic (values and lifestyle) and behavioural (how they "
                                                 "search and buy)."),
    ("5. Segment Scores", "Each segment scored on market size, accessibility, buying potential and competition.",
     f"{target} have the highest weighted score, {totals[target]:.2f} out of 5.",
     f"How is {target}' score worked out?", f"Each score times its weight, added up: {working} = "
                                            f"{totals[target]:.2f}. In Excel it is SUMPRODUCT."),
    ("5. Segment Scores Compared", "A bar chart of the four weighted scores.",
     f"{target} come first, then {ranked[1]} with {totals[ranked[1]]:.2f}.",
     f"Why not {ranked[1]}?", "They are a big group, but they have little money and many free courses already aim "
                              "at them."),
    (f"6. Target Segment: {target}", "Why this segment should come first.",
     "They can pay in parts, have a strong reason to buy, are easy to reach, and few courses suit them.",
     "Why does buying potential have the biggest weight (30%)?", "A new business needs customers who can actually "
                                                                 "pay."),
    ("Conclusion", "The main findings.",
     f"Customers need hands-on skills at a fair price, and {target} are the best place to start.",
     "What would the business do first?", "Launch an evening and weekend course for career switchers, with a clear "
                                          "price, a free first class and payment in parts."),
    ("Sources", "Where the data came from, and the AI help.",
     "The full list of sources, with links, is in the Excel file.",
     "How do you know the figures are right?", "They come from published sources: ISC2, the World Economic Forum, "
                                               "the NUC, the minimum wage law, and my own research in part 1."),
    ("Thank You", "The end.", "Thank you for listening. Are there any questions?", None, None),
]
for n, (title, shows, say, question, answer) in enumerate(slides, start=1):
    h(3, f"Slide {n}: {title}")
    p(shows, "What it shows:")
    p(say, "What to say:")
    if question:
        p(question, "Your teacher may ask:")
        p(answer, "Answer:")

h(2, "Practice quiz")
p("Cover the answers and test yourself.")
quiz = [
    ("What is a customer persona?", "A made-up customer, based on research, that stands for a target segment."),
    ("Name the five stages of the journey map.", ", ".join(C.STAGES) + "."),
    ("Name the four kinds of customer needs.", "Functional, economic, emotional and social."),
    ("What are the 5 most important needs?", "; ".join(top5) + "."),
    ("Name the four bases of segmentation.", "Demographic, geographic, psychographic and behavioural."),
    ("Name the four segments.", ", ".join(s[0] for s in C.SEGMENTS) + "."),
    ("On which four criteria were the segments scored?", ", ".join(c for c, _ in C.CRITERIA) + "."),
    ("Which segment is the target, and what was its score?", f"{target}, {totals[target]:.2f} out of 5."),
    ("What does SUMPRODUCT do?", "It multiplies each score by its weight and adds the results."),
    ("What is the lowest point of Tunde's journey, and how can it be fixed?",
     f"{low}: fix it with clear prices, real reviews and a free first class."),
]
items([q for q, _ in quiz], ordered=True)
h(3, "Answers")
items([a for _, a in quiz], ordered=True)

css = ("body{font-family:Calibri;font-size:11pt} h1,h2,h3{font-family:Calibri;color:#1F3864} h1{font-size:20pt} "
       "h2{font-size:15pt} h3{font-size:12pt} p{margin:4pt 0}")
page = f"<html><head><meta charset='utf-8'><title>Presenter Guide</title><style>{css}</style></head><body>" \
       + "\n".join(parts) + "</body></html>"
with tempfile.TemporaryDirectory() as tmp:
    src = os.path.join(tmp, "Presenter_Guide.html")
    with open(src, "w", encoding="utf-8") as fh:
        fh.write(page)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo", "--headless", "--convert-to",
                    "docx:MS Word 2007 XML", "--outdir", tmp, src], check=True, capture_output=True,
                   env={**os.environ, "SAL_USE_VCLPLUGIN": "svp"})
    os.replace(os.path.join(tmp, "Presenter_Guide.docx"), out)
print(f"wrote {out}")
