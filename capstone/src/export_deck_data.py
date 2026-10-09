#!/usr/bin/env python3
"""Read the worked-out figures the slides use out of the finished Capstone_Market_Research.xlsx
(the market's yearly growth, the size estimate, the gap priorities and the review averages), so
the slides match the workbook.

Usage: python export_deck_data.py Capstone_Market_Research.xlsx deck_data.json
"""
import json
import sys

from openpyxl import load_workbook

import capstone_data as D

workbook, out = sys.argv[1:3]
wb = load_workbook(workbook, data_only=True)


def by_label(sheet, label, col=2):
    for row in wb[sheet].iter_rows():
        if row[0].value == label:
            value = row[col - 1].value
            if value is None:
                sys.exit(f"{sheet}: {label!r} has no value; run the workbook through LibreOffice first")
            return value
    sys.exit(f"{sheet}: no row {label!r}")


data = {
    "growth": by_label("Market", "Yearly growth of the market, 2026 to 2031"),
    "experts_ratio": by_label("Market", "Experts in South Africa for each one in Nigeria"),
    "sizing": [[label, by_label("Market", label)] for label, *_ in D.SIZING],
    "gaps": [{"gap": gap, "priority": by_label("Gaps & Opportunities", gap, 6),
              "rank": by_label("Gaps & Opportunities", gap, 7)} for gap, *_ in D.GAPS],
    "rival_rating": by_label("Reviews", "Average rating of the rated rivals"),
    "course_rating": by_label("Reviews", "Average rating"),
    "course_learners": by_label("Reviews", "Total learners on these courses", 4),
}
with open(out, "w") as fh:
    json.dump(data, fh, indent=2)
print(f"wrote {out}")
