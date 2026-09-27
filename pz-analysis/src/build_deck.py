#!/usr/bin/env python3
"""Build the PZ Nigeria PESTLE, SWOT and industry analysis deck in PowerPoint 2013's own style.

Usage: python build_deck.py slide_data.json template.pptx out.pptx

Every rating, score and conclusion comes from slide_data.json, which export_slide_data.py
reads out of the finished workbook, so the slides always match the Excel file.

The deck is built the way it would be in PowerPoint 2013 itself: on the widescreen Office
Theme (office2013.py), using the standard layouts (Title Slide, Title and Content, Two
Content, Title Only), with tables and charts placed in the content placeholders, as the
placeholder's Insert Table / Insert Chart icons do. Text uses the theme fonts (Calibri Light
titles, Calibri body).

Two colours, both from the theme's colour palette, and black text on white otherwise:
- dark blue ("Blue, Accent 1, Darker 50%") for slide titles (set once on the slide master),
  table header rows, chart bars, labels, box outlines and arrows;
- orange ("Orange, Accent 2, Darker 25%") only to pick things out: the most important bar in
  a chart, the harmful side of the SWOT, "High" priorities, the strongest of the five forces,
  and the accent line on the title slides.

Next to the deck it writes out.anim.json, the entrance animations that finish_deck.py adds:
step 1 appears after the slide transition, and each later step after the one before it.
"""
import io
import json
import re
import sys
import zipfile
from datetime import datetime, timezone

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from office2013 import THEME_XML

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
TABLE_GRID = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # PowerPoint's "No Style, Table Grid"

# The two colours, as theme colours with PowerPoint's Lighter/Darker setting.
NAVY = (MSO_THEME_COLOR.ACCENT_1, -0.5)       # "Blue, Accent 1, Darker 50%"    1F4E79
ORANGE = (MSO_THEME_COLOR.ACCENT_2, -0.25)    # "Orange, Accent 2, Darker 25%"  C55A11
BLACK = (MSO_THEME_COLOR.TEXT_1, 0)
WHITE = (MSO_THEME_COLOR.BACKGROUND_1, 0)

data_path, template_path, out_path = sys.argv[1:4]
data = json.load(open(data_path))
prs = Presentation(template_path)
LAYOUT = {layout.name: layout for layout in prs.slide_layouts}
plan = {}


# ------------------------------------------------------------------ small helpers
def one(x):
    return f"{x:.1f}"


def two(x):
    return f"{x:.2f}"


def lc(text):
    """Lower-case the first letter for use mid-sentence, but leave acronyms (FX, NAFDAC) alone."""
    return text[0].lower() + text[1:] if re.match(r"[A-Z][a-z]", text) else text


def paint(color, colour):
    """Set a python-pptx colour to a theme colour (plus its Lighter/Darker setting)."""
    color.theme_color = colour[0]
    if colour[1]:
        color.brightness = colour[1]


def anim(slide, shape, step, effect="fade"):
    plan.setdefault(str(len(prs.slides)), []).append({"name": shape.name, "step": step, "effect": effect})


def colour_slide_titles():
    """Dark blue titles on every slide, set once on the slide master (View › Slide Master)."""
    style = prs.slide_master.element.find(qn("p:txStyles")).find(qn("p:titleStyle"))
    rpr = style.find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    for old in rpr.findall(qn("a:solidFill")):
        rpr.remove(old)
    rpr.insert(0, etree.fromstring(f'<a:solidFill xmlns:a="{A_NS}"><a:schemeClr val="accent1">'
                                   '<a:lumMod val="50000"/></a:schemeClr></a:solidFill>'))


def new_slide(layout, title=None):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if title is not None:
        slide.shapes.title.text = title
        anim(slide, slide.shapes.title, 1)
    return slide


def placeholder(slide, idx):
    return next(p for p in slide.placeholders if p.placeholder_format.idx == idx)


def write(frame, items, sizes=(24, 20)):
    """items: [(level, [(text, bold[, colour]), ...]), ...]; one paragraph each, sized by level."""
    frame.clear()
    for i, (level, runs) in enumerate(items):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        para.level = level
        for text, bold, *colour in runs:
            run = para.add_run()
            run.text = text
            run.font.size = Pt(sizes[min(level, len(sizes) - 1)])
            if bold:
                run.font.bold = True
            if colour:
                paint(run.font.color, colour[0])


def into_placeholder(slide, ph, frame):
    """Put a table or chart where a content placeholder was, as PowerPoint's Insert icons do:
    the new frame takes over the placeholder's id, name and slot, and the empty placeholder
    goes."""
    nv = frame._element.find(qn("p:nvGraphicFramePr"))
    c_nv = nv.find(qn("p:cNvPr"))
    c_nv.set("id", str(ph.shape_id))
    c_nv.set("name", ph.name)
    nv_pr = nv.find(qn("p:nvPr"))
    nv_pr.insert(0, etree.fromstring(
        f'<p:ph xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" idx="{ph.placeholder_format.idx}"/>'))
    ph._element.addprevious(frame._element)
    ph._element.getparent().remove(ph._element)
    return frame


def cell_border(cell, width=Pt(1)):
    """Thin light-grey lines on all four sides, written out so every viewer shows them."""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
    lines = [etree.fromstring(
        f'<a:{side} xmlns:a="{A_NS}" w="{int(width)}" cap="flat" cmpd="sng" algn="ctr">'
        '<a:solidFill><a:schemeClr val="bg1"><a:lumMod val="75000"/></a:schemeClr></a:solidFill>'
        '<a:prstDash val="solid"/><a:round/><a:headEnd type="none" w="med" len="med"/>'
        f'<a:tailEnd type="none" w="med" len="med"/></a:{side}>')
        for side in ("lnL", "lnR", "lnT", "lnB")]
    for line in reversed(lines):
        tc_pr.insert(0, line)


def bullet(para):
    """The body text bullet (Arial •) for a paragraph inside a table cell."""
    p_pr = para._p.get_or_add_pPr()
    p_pr.set("marL", "228600")
    p_pr.set("indent", "-228600")
    for tag in ("a:buNone", "a:buFont", "a:buChar"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    p_pr.append(etree.fromstring(f'<a:buFont xmlns:a="{A_NS}" typeface="Arial" '
                                 'panose="020B0604020202020204" pitchFamily="34" charset="0"/>'))
    p_pr.append(etree.fromstring(f'<a:buChar xmlns:a="{A_NS}" char="•"/>'))


def table(slide, idx, rows, widths, heights, size, header_rows=(0,), header_colour=lambda r, c: NAVY):
    """A Table Grid table in content placeholder idx, its header rows filled with a colour and
    lettered in white. rows: lists of cell values, where a value is text, a tuple of
    (text, bold[, colour]) runs, a list of bullet lines, or None (the covered part of a merged
    cell, filled in afterwards)."""
    ph = placeholder(slide, idx)
    frame = slide.shapes.add_table(len(rows), len(widths), ph.left, ph.top, sum(widths, Emu(0)), sum(heights, Emu(0)))
    tbl = frame.table
    tbl.first_row = True
    tbl.horz_banding = False
    frame._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = TABLE_GRID
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    for r, h in enumerate(heights):
        tbl.rows[r].height = h
    for r, values in enumerate(rows):
        header = r in header_rows
        for c, value in enumerate(values):
            cell = tbl.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if header:
                cell.fill.solid()
                paint(cell.fill.fore_color, header_colour(r, c))
            else:
                cell.fill.background()
            cell_border(cell)
            if value is None:
                continue
            lines = value if isinstance(value, list) else [value]
            for k, line in enumerate(lines):
                para = cell.text_frame.paragraphs[0] if k == 0 else cell.text_frame.add_paragraph()
                if isinstance(value, list):
                    bullet(para)
                parts = line if isinstance(line, tuple) else ((line, header),)
                for text, bold, *colour in parts:
                    run = para.add_run()
                    run.text = text
                    run.font.size = Pt(size)
                    run.font.bold = bold
                    paint(run.font.color, WHITE if header else (colour[0] if colour else BLACK))
    return into_placeholder(slide, ph, frame), tbl


def chart(slide, idx, labels, values, *, horizontal, fmt, maximum, major, highlight=()):
    """A dark-blue bar or column chart in content placeholder idx, in Office 2013's chart
    style; the bars at the positions in highlight are orange."""
    ph = placeholder(slide, idx)
    chart_data = CategoryChartData(number_format=fmt)
    chart_data.categories = labels
    chart_data.add_series("Score", values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    frame = slide.shapes.add_chart(kind, ph.left, ph.top, ph.width, ph.height, chart_data)
    ch = frame.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.size = Pt(14)
    paint(ch.font.color, (MSO_THEME_COLOR.TEXT_1, 0.35))      # Office 2013 chart text: 595959
    plot = ch.plots[0]
    plot.gap_width = 80
    plot.vary_by_categories = False
    series = plot.series[0]
    series.format.fill.solid()
    paint(series.format.fill.fore_color, NAVY)
    series.format.line.fill.background()
    for i in highlight:
        point = series.points[i].format
        point.fill.solid()
        paint(point.fill.fore_color, ORANGE)
        point.line.fill.background()
    plot.has_data_labels = True
    labels_ = plot.data_labels
    labels_.number_format = fmt
    labels_.number_format_is_linked = False
    labels_.position = XL_LABEL_POSITION.OUTSIDE_END
    labels_.font.size = Pt(14)
    labels_.font.bold = True
    paint(labels_.font.color, (MSO_THEME_COLOR.TEXT_1, 0.25))
    cat, val = ch.category_axis, ch.value_axis
    cat.major_tick_mark = XL_TICK_MARK.NONE
    paint(cat.format.line.color, (MSO_THEME_COLOR.TEXT_1, 0.85))   # D9D9D9, the 2013 axis line
    cat.has_major_gridlines = False
    val.minimum_scale = 0
    val.maximum_scale = maximum
    val.major_unit = major
    val.major_tick_mark = XL_TICK_MARK.NONE
    val.format.line.fill.background()
    val.tick_labels.number_format = fmt
    val.tick_labels.number_format_is_linked = False
    if horizontal:
        cat.reverse_order = True                         # first item at the top, as in the table
        val.visible = False
        val.has_major_gridlines = False
    else:
        val.has_major_gridlines = True
        gridlines = val.major_gridlines.format.line
        paint(gridlines.color, (MSO_THEME_COLOR.TEXT_1, 0.85))
        gridlines.width = Pt(0.75)
    return into_placeholder(slide, ph, frame)


def box(slide, x, y, w, h, lines, *, filled=False):
    """A rectangle with centred text: white with a dark-blue outline, or (filled) solid orange
    with white text. lines: [(text, size, bold, colour), ...]."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    paint(shape.fill.fore_color, ORANGE if filled else WHITE)
    paint(shape.line.color, ORANGE if filled else NAVY)
    shape.line.width = Pt(1.5)
    frame = shape.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for k, (text, size, bold, colour) in enumerate(lines):
        para = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        paint(run.font.color, WHITE if filled else colour)
    return shape


def name_line(line, kind):
    """Name a line as PowerPoint does ("Straight Connector 3"); python-pptx says "Connector 3"."""
    c_nv = line._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvPr"))
    c_nv.set("name", f"{kind} {int(c_nv.get('id')) - 1}")


def arrow(slide, x1, y1, x2, y2):
    """A dark-blue straight arrow (Insert › Shapes › Arrow) pointing from (x1, y1) to (x2, y2)."""
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    paint(line.line.color, NAVY)
    line.line.width = Pt(1.5)
    ln = line.line._get_or_add_ln()
    ln.append(etree.fromstring(f'<a:tailEnd xmlns:a="{A_NS}" type="triangle"/>'))
    name_line(line, "Straight Arrow Connector")
    return line


def accent_line(slide):
    """A short orange line between the title and subtitle of a Title Slide. The layout leaves
    no gap there, so the title moves up 0.3 in to make room (as dragging it up would)."""
    title = slide.shapes.title
    left, width, height = title.left, title.width, title.height      # inherited from the layout
    title.left, title.top, title.width, title.height = left, title.top - Inches(0.3), width, height
    y = title.top + height + Inches(0.16)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.17), y, Inches(8.17), y)
    paint(line.line.color, ORANGE)
    line.line.width = Pt(3)
    name_line(line, "Straight Connector")
    return line


# ------------------------------------------------------------------ figures from Excel
summary = {f["factor"]: f for f in data["pestleSummary"]}
issues_of = {f: [p for p in data["pestle"] if p["factor"] == f] for f in summary}
top = max(data["pestleSummary"], key=lambda f: f["average"])
low = min(data["pestleSummary"], key=lambda f: f["average"])
top_issue = max(data["pestle"], key=lambda p: p["score"])
high = sum(p["priority"] == "High" for p in data["pestle"])
threats = sum(p["type"] == "Threat" for p in data["pestle"])
opps = sum(p["type"] == "Opportunity" for p in data["pestle"])
pos = data["position"]
company = data["company"]
force = {f["force"]: f for f in data["forces"]}

colour_slide_titles()

# ================================================================== 1. Title
s = new_slide("Title Slide", "PZ Nigeria Limited")
anim(s, accent_line(s), 1, "wipe-left")
sub = placeholder(s, 1)
write(sub.text_frame, [(0, [("PESTLE, SWOT and Industry Analysis", False)]),
                       (0, [("Marketing Department, Ilupeju, Lagos", False)])], sizes=(24,))
anim(s, sub, 2)
s.notes_slide.notes_text_frame.text = (
    "Good day. This presentation analyses PZ Nigeria Limited, the company behind Imperial Leather, Premier Cool, "
    "Joy and Morning Fresh. I will look at its business environment with a PESTLE analysis, at the company itself "
    "with a SWOT analysis, and at its industry with Porter's Five Forces. All the scores were calculated in "
    "Microsoft Excel, in the workbook PZ_Nigeria_Analysis.xlsx.")

# ================================================================== 2. Company overview
s = new_slide("Two Content", "Company Overview")
left, right = placeholder(s, 1), placeholder(s, 2)
write(left.text_frame, [
    (0, [("Family Care (consumer)", True, NAVY)]),
    (1, [(company["Consumer business (Family Care)"], False)]),
    (0, [("Electrical", True, NAVY)]),
    (1, [(company["Electrical business"], False)]),
    (0, [("Owner", True, NAVY)]),
    (1, [("PZ Cussons group (United Kingdom)", False)]),
], sizes=(24, 20))
write(right.text_frame, [
    (0, [(company["Managing Director"], True, NAVY), (": Managing Director", False)]),
    (0, [(company["Chief Financial Officer"], True, NAVY), (": Chief Financial Officer", False)]),
    (0, [("Structure", True, NAVY)]),
    (1, [("Managing Director (with the CFO and HR Director)", False)]),
    (1, [("Category Managers", False)]),
    (1, [("Brand Managers: one for each brand", False)]),
    (1, [("Assistant Brand Managers", False)]),
], sizes=(24, 20))
anim(s, left, 2)
anim(s, right, 3)
s.notes_slide.notes_text_frame.text = (
    "PZ Nigeria has two businesses. Family Care makes Imperial Leather, Premier Cool, Joy and Morning Fresh; the "
    "Electrical business sells air conditioners, fridges, freezers and washing machines. The company is led by "
    "Mr Oghale Ogueni, the Managing Director, and Mr Oladare Oresukan, the Chief Financial Officer, and it is part "
    "of the UK-based PZ Cussons group. It uses a product and brand structure: the Managing Director, category "
    "managers, then a brand manager in charge of each brand, supported by assistant brand managers and shared "
    "departments.")

# ================================================================== 3-5. PESTLE tables
for pair in (("Political", "Economic"), ("Social", "Technological"), ("Legal", "Environmental")):
    s = new_slide("Title and Content", f"PESTLE: {pair[0]} and {pair[1]}")
    rows = [["Factor", "Issue", "Effect on PZ Nigeria", "Score (of 25)", "Priority"]]
    merged = []
    for factor in pair:
        for k, p in enumerate(issues_of[factor]):
            if k == 0:
                merged.append((len(rows), factor))
            priority = ((p["priority"], True, ORANGE),) if p["priority"] == "High" else p["priority"]
            rows.append([None, p["issue"], p["effect"], str(p["score"]), priority])
    widths = [Inches(w) for w in (1.75, 2.95, 4.4, 1.2, 1.2)]
    heights = [Inches(0.5)] + [Inches(0.68)] * (len(rows) - 1)
    frame, tbl = table(s, 1, rows, widths, heights, 16)
    for r, factor in merged:                              # factor names, merged down their rows
        cell = tbl.cell(r, 0)
        cell.merge(tbl.cell(r + len(issues_of[factor]) - 1, 0))
        para = cell.text_frame.paragraphs[0]
        run = para.add_run()
        run.text = factor
        run.font.size = Pt(16)
        run.font.bold = True
        paint(run.font.color, NAVY)
        run = cell.text_frame.add_paragraph().add_run()
        run.text = f"Average {one(summary[factor]['average'])}"
        run.font.size = Pt(14)
        paint(run.font.color, BLACK)
    for r in range(len(rows)):
        for c in (3, 4):
            for para in tbl.cell(r, c).text_frame.paragraphs:
                para.alignment = PP_ALIGN.CENTER
    anim(s, frame, 2, "wipe-up")
    lines = []
    for factor in pair:
        best = max(issues_of[factor], key=lambda p: p["score"])
        lines.append(f"{factor} factors average {one(summary[factor]['average'])} out of 25; the biggest issue is "
                     f"{lc(best['issue'])}, which scores {best['score']}: {lc(best['effect'])}.")
    s.notes_slide.notes_text_frame.text = (
        "Each issue is rated for impact and likelihood from 1 to 5; the score is impact times likelihood, so the "
        "highest possible score is 25, and 16 or more is high priority. " + " ".join(lines))

# ================================================================== 6. PESTLE scores
s = new_slide("Two Content", "PESTLE Scores")
factors = [f["factor"] for f in data["pestleSummary"]]
frame = chart(s, 1, factors, [round(f["average"], 1) for f in data["pestleSummary"]],
              horizontal=True, fmt="0.0", maximum=25, major=5, highlight=[factors.index(top["factor"])])
right = placeholder(s, 2)
write(right.text_frame, [
    (0, [("Average score of each factor, out of 25 (Excel: AVERAGEIF)", False)]),
    (0, [("Most important: ", False), (f"{top['factor']} ({one(top['average'])})", True, ORANGE)]),
    (0, [("Biggest issue: ", False), (f"{top_issue['issue']} ({top_issue['score']} of 25)", True, NAVY)]),
    (0, [("High priority: ", False), (f"{high} of {len(data['pestle'])} issues", True, NAVY)]),
    (0, [(f"{threats} threats and {opps} opportunities", True, NAVY)]),
], sizes=(22,))
anim(s, frame, 2, "wipe-left")
anim(s, right, 3)
s.notes_slide.notes_text_frame.text = (
    f"The chart shows each factor's average score, calculated in Excel with AVERAGEIF. {top['factor']} factors "
    f"matter most, with an average of {one(top['average'])} out of 25, and the single biggest issue is "
    f"{lc(top_issue['issue'])}, which scores the maximum of {top_issue['score']}. {low['factor']} factors matter "
    f"least, at {one(low['average'])}. Overall {high} of the {len(data['pestle'])} issues are high priority, and "
    f"there are {threats} threats against {opps} opportunities, so the environment is difficult.")

# ================================================================== 7. SWOT
s = new_slide("Title and Content", "SWOT Analysis")
sw = data["swot"]
rows = [["Strengths (internal)", "Weaknesses (internal)"],
        [sw["strengths"], sw["weaknesses"]],
        ["Opportunities (external)", "Threats (external)"],
        [sw["opportunities"], sw["threats"]]]
# Helpful points (left) in blue, harmful ones (right) in orange.
frame, tbl = table(s, 1, rows, [Inches(5.75)] * 2,
                   [Inches(0.45), Inches(1.9), Inches(0.45), Inches(1.9)], 15, header_rows=(0, 2),
                   header_colour=lambda r, c: NAVY if c == 0 else ORANGE)
anim(s, frame, 2, "wipe-up")
s.notes_slide.notes_text_frame.text = (
    "PZ's main strengths are its well-known brands, more than a century in Nigeria and a wide distribution network. "
    "Its weaknesses are slow decisions, because approvals can wait for the UK owner and there are many management "
    "levels, and its dependence on imported raw materials when the naira is weak. The opportunities come from "
    "Nigeria's large young population, digital selling, smaller packs and the customers left behind by Unilever "
    "and P&G. The threats are the weak naira, high inflation, strong competition, fake products, high energy costs "
    "and policy changes.")

# ================================================================== 8. SWOT scoring
s = new_slide("Two Content", "SWOT Scoring (IFE and EFE)")
frame = chart(s, 1, ["Strengths", "Weaknesses", "Opportunities", "Threats"],
              [round(pos[k], 2) for k in ("strengths", "weaknesses", "opportunities", "threats")],
              horizontal=False, fmt="0.00", maximum=2.5, major=0.5, highlight=[1, 3])
right = placeholder(s, 2)
write(right.text_frame, [
    (0, [("IFE (internal): ", False), (two(data["ifeTotal"]), True, NAVY)]),
    (1, [(data["ifeResult"], False)]),
    (0, [("EFE (external): ", False), (two(data["efeTotal"]), True, NAVY)]),
    (1, [(data["efeResult"], False)]),
    (0, [("IE matrix: ", False), (f"cell {pos['ieCell']}, {lc(pos['strategy'])}", True, NAVY)]),
    (1, [(pos["suggested"], False)]),
    (0, [("Weight × rating in Excel; 2.50 is average", False)]),
], sizes=(22, 18))
anim(s, frame, 2, "wipe-up")
anim(s, right, 3)
s.notes_slide.notes_text_frame.text = (
    f"In Excel each SWOT point has a weight for how important it is and a rating from 1 to 4; weight times rating "
    f"gives the weighted score. The internal total, IFE, is {two(data['ifeTotal'])}: above the average of 2.50, "
    f"because strengths ({two(pos['strengths'])}) outweigh weaknesses ({two(pos['weaknesses'])}). The external "
    f"total, EFE, is {two(data['efeTotal'])}: below average, so PZ is not yet making the most of its opportunities "
    f"or defending well enough against its threats. Together they place PZ in cell {pos['ieCell']} of the "
    f"internal-external matrix, which means {pos['strategy'].lower()}: {pos['suggested'].lower()}.")

# ================================================================== 9. Industry overview
s = new_slide("Title and Content", "Industry Overview")
body = placeholder(s, 1)
o = data["overview"]
write(body.text_frame, [(0, [(f"{k}: ", True, NAVY), (o[k], False)])
                        for k in ("Industry", "Market", "Customers", "Key trends", "Key success factors")],
      sizes=(24,))
anim(s, body, 2)
s.notes_slide.notes_text_frame.text = (
    "PZ competes in fast-moving consumer goods, mainly soaps and dishwashing liquid, and in consumer electricals. "
    "Nigeria has over 200 million people, most of them young, so it is one of Africa's largest consumer markets. "
    "But prices are rising with inflation, shoppers are moving to smaller packs and cheaper brands, and some "
    "multinationals have cut local production. To succeed, a company needs strong brands, low costs, wide "
    "distribution, local raw materials and affordable pack sizes.")

# ================================================================== 10. Competitors
s = new_slide("Title and Content", "Main Competitors")
rows = [["Competitor", "Key brands", "Competes with PZ in", "Position"]]
rows += [[((c["name"], True, NAVY),), c["brands"], c["competes"], c["position"]] for c in data["competitors"]]
frame, tbl = table(s, 1, rows, [Inches(w) for w in (2.6, 3.0, 2.4, 3.5)],
                   [Inches(0.5)] + [Inches(0.66)] * (len(rows) - 1), 15)
anim(s, frame, 2, "wipe-up")
s.notes_slide.notes_text_frame.text = (
    "These are PZ's main competitors. Unilever left home care and soaps in 2023, and Procter and Gamble stopped "
    "making products in Nigeria, so both now leave space PZ can fill. Reckitt is strong in hygiene with Dettol and "
    "Harpic, Hayat Kimya makes low-priced products locally, and many local brands win on price in open markets. "
    "In electricals, PZ faces LG, Samsung, Hisense and Scanfrost, plus many cheaper imports.")

# ================================================================== 11. Five Forces diagram
s = new_slide("Title Only", "Porter's Five Forces")
W = prs.slide_width
places = {
    "Competitive rivalry": (Inches(4.72), Inches(3.55), Inches(3.9), Inches(1.55)),
    "Threat of new entrants": (Inches(4.72), Inches(1.95), Inches(3.9), Inches(1.3)),
    "Bargaining power of suppliers": (Inches(0.92), Inches(3.3), Inches(3.4), Inches(2.05)),
    "Bargaining power of buyers": (W - Inches(0.92) - Inches(3.4), Inches(3.3), Inches(3.4), Inches(2.05)),
    "Threat of substitutes": (Inches(4.72), Inches(5.45), Inches(3.9), Inches(1.3)),
}
mid_x, mid_y = Inches(4.72 + 3.9 / 2), Inches(3.55 + 1.55 / 2)
for step, name in enumerate(places, start=2):
    f = force[name]
    x, y, w, h = places[name]
    shape = box(s, x, y, w, h, [(name, 18, True, NAVY), (f"{f['level']}: {f['rating']} of 5", 16, True, BLACK),
                                (f["why"], 14, False, BLACK)],
                filled=name == data["strongestForce"])
    anim(s, shape, step)
    tip = {  # arrows point at the rivalry box in the middle
        "Threat of new entrants": (mid_x, y + h, mid_x, Inches(3.55)),
        "Threat of substitutes": (mid_x, y, mid_x, Inches(3.55 + 1.55)),
        "Bargaining power of suppliers": (x + w, mid_y, Inches(4.72), mid_y),
        "Bargaining power of buyers": (x, mid_y, Inches(4.72 + 3.9), mid_y),
    }.get(name)
    if tip:
        anim(s, arrow(s, *tip), step)
s.notes_slide.notes_text_frame.text = " ".join(
    f"{f['force']} is {f['level'].lower()}, rated {f['rating']} out of 5: {lc(f['why'])}." for f in data["forces"])

# ================================================================== 12. Five Forces ratings
s = new_slide("Two Content", "Five Forces Ratings")
names = [f["force"] for f in data["forces"]]
frame = chart(s, 1, names, [f["rating"] for f in data["forces"]],
              horizontal=True, fmt="0", maximum=5, major=1, highlight=[names.index(data["strongestForce"])])
right = placeholder(s, 2)
write(right.text_frame, [
    (0, [("Strength of each force, 1 (weak) to 5 (strong)", False)]),
    (0, [("Average: ", False), (f"{one(data['forcesAverage'])} of 5 ({data['forcesLevel']})", True, NAVY)]),
    (1, [(data["forcesVerdict"], False)]),
    (0, [("Strongest: ", False), (data["strongestForce"], True, ORANGE)]),
    (0, [("Weakest: ", False), (data["weakestForce"], True, NAVY)]),
    (0, [("Rated high: ", False), (f"{data['forcesHigh']} forces", True, NAVY)]),
], sizes=(22, 18))
anim(s, frame, 2, "wipe-left")
anim(s, right, 3)
s.notes_slide.notes_text_frame.text = (
    f"On average the five forces score {one(data['forcesAverage'])} out of 5, calculated in Excel with AVERAGE. "
    f"{data['forcesVerdict']}. The strongest force is {data['strongestForce'].lower()}; the weakest is the "
    f"{data['weakestForce'].lower()}, because it costs a lot to build factories, brands and distribution. "
    f"{data['forcesHigh']} forces are rated high, so PZ must keep costs low and brands strong to stay profitable.")

# ================================================================== 13. Recommendations
s = new_slide("Title and Content", "Recommendations")
body = placeholder(s, 1)
items = []
for r in data["recommendations"]:
    items += [(0, [(r["title"], True, NAVY)]), (1, [(r["why"], False)])]
write(body.text_frame, items, sizes=(24, 20))
for para in body.text_frame.paragraphs:                  # numbered 1. 2. 3. (Home › Numbering)
    if para.level == 0:
        p_pr = para._p.get_or_add_pPr()
        p_pr.set("marL", "457200")
        p_pr.set("indent", "-457200")
        p_pr.append(etree.fromstring(f'<a:buFont xmlns:a="{A_NS}" typeface="+mj-lt"/>'))
        p_pr.append(etree.fromstring(f'<a:buAutoNum xmlns:a="{A_NS}" type="arabicPeriod"/>'))
    else:
        para._p.get_or_add_pPr().set("marL", "685800")
anim(s, body, 2)
s.notes_slide.notes_text_frame.text = "Based on the analysis, I recommend five actions. " + " ".join(
    f"{i}: {lc(r['title'])}. {r['why']}." for i, r in enumerate(data["recommendations"], start=1))

# ================================================================== 14. Conclusion
s = new_slide("Title and Content", "Conclusion")
body = placeholder(s, 1)
write(body.text_frame, [
    (0, [("PESTLE: ", True, NAVY), (f"{top['factor']} factors matter most (average {one(top['average'])} of 25); "
                                    f"the biggest issue is {lc(top_issue['issue'])}", False)]),
    (0, [("SWOT: ", True, NAVY), (f"strong inside (IFE {two(data['ifeTotal'])}), but a weaker response to the "
                                  f"outside (EFE {two(data['efeTotal'])})", False)]),
    (0, [("Industry: ", True, NAVY), (f"{data['forcesLevel'].lower()} competitive pressure (Five Forces average "
                                      f"{one(data['forcesAverage'])} of 5); {data['strongestForce'].lower()} is the "
                                      "strongest force", False)]),
    (0, [("Strategy: ", True, NAVY), (f"{lc(pos['strategy'])}, through {pos['suggested'].lower()}", False)]),
], sizes=(28,))
anim(s, body, 2)
s.notes_slide.notes_text_frame.text = (
    f"To conclude: PZ Nigeria is strong inside, with well-known brands and a wide distribution network "
    f"(IFE {two(data['ifeTotal'])}), but its environment is hard. {top['factor']} factors, especially "
    f"{lc(top_issue['issue'])}, matter most, and the industry is highly competitive (five forces average "
    f"{one(data['forcesAverage'])} out of 5). The best strategy is to {pos['strategy'].lower()}, through "
    f"{pos['suggested'].lower()}.")

# ================================================================== 15. Thank you
s = new_slide("Title Slide", "Thank You")
anim(s, accent_line(s), 1, "wipe-left")
sub = placeholder(s, 1)
write(sub.text_frame, [(0, [("Questions?", False)])], sizes=(28,))
anim(s, sub, 2)
s.notes_slide.notes_text_frame.text = "Thank you for listening. I am happy to take questions."

# ------------------------------------------------------------------ document properties, save
props = prs.core_properties
now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
props.title = "PZ Nigeria Limited: PESTLE, SWOT and Industry Analysis"
props.subject = "PESTLE, SWOT and industry analysis of PZ Nigeria Limited"
props.author = props.last_modified_by = "Marketing Department"
props.comments = ""
props.keywords = "PZ Nigeria, PESTLE, SWOT, Five Forces"
props.revision = 1
props.created = props.modified = now
buf = io.BytesIO()
prs.save(buf)

# The notes pages python-pptx adds use its own (Office 2007) theme; give them the 2013 one too,
# and record the slide count in the document statistics.
with zipfile.ZipFile(buf) as zin, zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as zout:
    for info in zin.infolist():
        blob = zin.read(info.filename)
        if re.fullmatch(r"ppt/theme/theme\d+\.xml", info.filename):
            blob = THEME_XML.encode()
        elif info.filename == "docProps/app.xml":
            n = len(prs.slides)
            blob = re.sub(rb"<Slides>\d+</Slides>", f"<Slides>{n}</Slides>".encode(), blob)
            blob = re.sub(rb"<Notes>\d+</Notes>", f"<Notes>{n}</Notes>".encode(), blob)
        zout.writestr(info, blob)
with open(re.sub(r"\.pptx$", ".anim.json", out_path), "w") as fh:
    json.dump(plan, fh, indent=2)
print(f"wrote {out_path} ({len(prs.slides)} slides)")
