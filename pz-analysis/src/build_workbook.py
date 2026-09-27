#!/usr/bin/env python3
"""Build PZ_Nigeria_Analysis.xlsx: PESTLE, SWOT and industry analysis of PZ Nigeria Limited.

Ratings and weights are inputs (blue). Every score, average, count, ranking and conclusion is
a live Excel formula (black), using only functions that exist in Excel 2013: SUM, AVERAGE,
AVERAGEIF, COUNTIF, COUNTIFS, COUNTA, SUMIF, SUMPRODUCT, MAX, MIN, RANK, INDEX, MATCH, CHOOSE,
ROUND, TEXT and IF.

Next to the workbook it writes <output>.cells.json, the address of every table, which
export_slide_data.py uses to read the calculated results for the slides.

Usage: python build_workbook.py output.xlsx
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.page import PageMargins

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "PZ_Nigeria_Analysis.xlsx"

# ================================================================ content
COMPANY = [
    ("Company", "PZ Nigeria Limited, part of the UK-based PZ Cussons group"),
    ("Department", "Marketing Department, Ilupeju, Lagos"),
    ("Managing Director", "Mr Oghale Ogueni"),
    ("Chief Financial Officer", "Mr Oladare Oresukan"),
    ("Consumer business (Family Care)", "Imperial Leather, Premier Cool, Joy, Morning Fresh"),
    ("Electrical business", "Air conditioners, fridges and freezers, washing machines"),
    ("Structure", "Managing Director (with the CFO and HR Director) > Category Managers > "
                  "Brand Managers > Assistant Brand Managers"),
    ("How it is organised", "Product and brand structure: each brand has its own manager, and "
                            "shared support teams serve all the brands"),
    ("Industry", "Fast-moving consumer goods (personal care and home care) and consumer electricals"),
]

# factor, issue, effect on PZ Nigeria, type, impact (1-5), likelihood (1-5)
PESTLE = [
    ("Political", "Fuel subsidy removal and FX reforms", "Higher fuel and transport costs; easier access to dollars over time", "Threat", 4, 5),
    ("Political", "Import duties and levies", "Imported raw materials and parts cost more", "Threat", 3, 4),
    ("Political", "2027 elections and policy changes", "Uncertainty can delay investment and change the rules", "Threat", 3, 3),
    ("Economic", "Naira devaluation and FX shortages", "Imported raw materials cost much more; losses on foreign-currency debts", "Threat", 5, 5),
    ("Economic", "High inflation", "Shoppers buy less or switch to cheaper brands", "Threat", 5, 5),
    ("Economic", "High interest rates", "Borrowing is expensive; fewer appliances bought on credit", "Threat", 3, 4),
    ("Social", "Large, young, growing population", "A big and growing market for soaps, detergents and appliances", "Opportunity", 4, 5),
    ("Social", "Price-sensitive shoppers", "Demand for smaller, cheaper packs", "Opportunity", 4, 4),
    ("Social", "Growing hygiene awareness", "More demand for soaps and cleaning products", "Opportunity", 3, 4),
    ("Technological", "Social media and digital marketing", "Cheaper, targeted adverts that reach young buyers", "Opportunity", 3, 5),
    ("Technological", "E-commerce and modern retail", "New ways to sell: online stores and supermarkets", "Opportunity", 3, 4),
    ("Technological", "Poor power supply", "Factories spend more on diesel and solar; demand for energy-saving appliances", "Threat", 4, 4),
    ("Legal", "NAFDAC and SON product rules", "Every product must be registered and meet standards", "Threat", 3, 5),
    ("Legal", "FCCPC and ARCON rules", "Adverts need approval; fair prices and clear labels", "Threat", 2, 4),
    ("Legal", "Tax reforms and ₦70,000 minimum wage", "Higher staff costs and new tax rules to follow", "Threat", 3, 4),
    ("Environmental", "Plastic waste rules", "Must change packaging and pay for recycling (EPR)", "Threat", 3, 4),
    ("Environmental", "Flooding and climate change", "Floods disrupt deliveries and damage stock", "Threat", 3, 3),
    ("Environmental", "Demand for eco-friendly products", "A chance to launch greener packs and appliances", "Opportunity", 2, 3),
]
FACTORS = ["Political", "Economic", "Social", "Technological", "Legal", "Environmental"]

# SWOT points with their IFE / EFE weight and rating.
# IFE rating: strengths 3 (minor) or 4 (major), weaknesses 1 (major) or 2 (minor).
# EFE rating: how well PZ responds today, 1 (poor) to 4 (superior).
STRENGTHS = [
    ("Well-known brands such as Imperial Leather and Joy", 0.15, 4),
    ("More than 100 years in Nigeria; trusted by buyers", 0.10, 4),
    ("Wide distribution network across Nigeria", 0.12, 3),
    ("Each business focuses on its own customers", 0.08, 3),
    ("One manager per brand; shared teams save money", 0.07, 3),
    ("Support from the UK owner (PZ Cussons)", 0.08, 3),
]
WEAKNESSES = [
    ("Waiting for the UK owner's approval can slow work", 0.08, 2),
    ("Too many management levels can cause delays", 0.07, 2),
    ("Brands may fight over the same marketing money", 0.05, 2),
    ("Naira problems force hard decisions on imports", 0.12, 1),
    ("Prices often higher than local low-cost brands", 0.08, 2),
]
OPPORTUNITIES = [
    ("Large, young and growing population", 0.12, 3),
    ("Social media, e-commerce and modern retail", 0.08, 2),
    ("Smaller, affordable packs for shoppers on a budget", 0.10, 3),
    ("Customers left by Unilever and P&G since 2023", 0.10, 3),
    ("Exports to other African countries under AfCFTA", 0.04, 2),
    ("Buying raw materials locally to cut dollar costs", 0.06, 2),
]
THREATS = [
    ("Naira devaluation and foreign-exchange shortages", 0.14, 2),
    ("High inflation reduces what consumers can spend", 0.12, 2),
    ("Strong competition and cheap imports", 0.09, 3),
    ("Fake (counterfeit) products", 0.05, 2),
    ("High energy and transport costs", 0.06, 2),
    ("Policy changes and insecurity", 0.04, 2),
]

OVERVIEW = [
    ("Industry", "Fast-moving consumer goods (FMCG): personal care (soaps, bathing) and home care "
                 "(dishwashing), plus consumer electricals"),
    ("Market", "Over 200 million people, most of them young: one of Africa's largest consumer markets"),
    ("Customers", "Households, reached through distributors, wholesalers, open markets, "
                  "supermarkets and online stores"),
    ("Key trends", "Prices rising with inflation; shoppers moving to smaller packs and cheaper brands; "
                   "multinationals cutting local production; more selling online and on social media"),
    ("Key success factors", "Strong brands, low costs, wide distribution, local raw materials and "
                            "affordable pack sizes"),
]
# force, rating (1-5), why
FORCES = [
    ("Competitive rivalry", 5, "Many local and foreign brands fight on price, promotions and shelf space"),
    ("Bargaining power of buyers", 4, "Shoppers are price-sensitive and switch brands easily; big distributors bargain hard"),
    ("Bargaining power of suppliers", 4, "Key raw materials, packaging and parts are imported and priced in dollars"),
    ("Threat of substitutes", 3, "Cheaper unbranded and local soaps; fans instead of air conditioners"),
    ("Threat of new entrants", 2, "Factories, brands and distribution cost a lot, but cheap imports still get in"),
]
# competitor, key brands, competes with PZ in, position
COMPETITORS = [
    ("Unilever Nigeria", "Vaseline, Close-Up, Pepsodent", "Personal care",
     "Left home care and skin cleansing (soaps) in 2023"),
    ("Procter & Gamble", "Ariel, Pampers, Always", "Home care",
     "Stopped making products in Nigeria; imports only (announced 2023)"),
    ("Reckitt Nigeria", "Dettol, Harpic, Jik", "Antiseptic soap, home care",
     "Strong in hygiene and health products"),
    ("Hayat Kimya", "Molfix, Bingo, Familia", "Home care",
     "Makes products locally; low prices; growing fast"),
    ("Local low-cost brands", "Many local soaps, detergents and dishwashing liquids", "All Family Care brands",
     "Win on price, especially in open markets"),
    ("LG, Samsung, Hisense, Scanfrost", "Air conditioners, fridges, washing machines", "Electrical",
     "Strong brands, plus many cheaper imports"),
]
RECOMMENDATIONS = [
    ("Buy more raw materials in Nigeria", "Cuts dollar costs and the damage from naira devaluation (the biggest PESTLE issue)"),
    ("Sell smaller, affordable packs", "Keeps price-sensitive shoppers buying Imperial Leather, Premier Cool, Joy and Morning Fresh"),
    ("Grow digital marketing and online sales", "Reaches the large young population cheaply through social media and e-commerce"),
    ("Win customers competitors left behind", "Use the wide distribution network to fill shelves Unilever and P&G have left"),
    ("Make decisions faster", "Give local managers more authority and cut approval levels to fix the main weaknesses"),
]

# ================================================================ styling
FONT = "Arial"
INK = "1F3864"        # dark blue: titles and header rows (as in the slides)
TEXT = "000000"
INPUT = "0000FF"      # inputs you can change (financial-model convention)
MUTED = "404040"
GRID = "BFBFBF"
BAND = "F2F2F2"       # light grey band
LIGHT = "DEEBF7"      # light blue, the box colour used on the slides
WHITE = "FFFFFF"

thin = Side(style="thin", color=GRID)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP_RULE = Border(left=thin, right=thin, top=Side(style="medium", color=INK), bottom=thin)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def font(size=10, bold=False, italic=False, color=TEXT):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def al(h="left", wrap=True, indent=1):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=indent if h == "left" else 0)


def put(ws, ref, value, *, f=None, fl=None, fmt=None, a=None, bd=BOX):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fl is not None:
        c.fill = fl
    if fmt is not None:
        c.number_format = fmt
    c.alignment = a or al()
    if bd is not None:
        c.border = bd
    return c


def span(ws, first, last, row, value, **kw):
    """Write value into first..last (merged) on one row, with the border on every cell."""
    if first != last:
        ws.merge_cells(f"{first}{row}:{last}{row}")
    cols = [chr(c) for c in range(ord(first), ord(last) + 1)]
    for col in cols[1:]:
        c = ws[f"{col}{row}"]
        c.border = kw.get("bd", BOX)
        if kw.get("fl") is not None:
            c.fill = kw["fl"]
    return put(ws, f"{first}{row}", value, **kw)


def sheet_title(ws, last_col, title, subtitle, sub_height=32):
    ws.merge_cells(f"A1:{last_col}1")
    put(ws, "A1", title, f=font(16, bold=True, color=WHITE), fl=fill(INK), a=al("center"), bd=None)
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{last_col}2")
    put(ws, "A2", subtitle, f=font(10, italic=True, color=MUTED), a=al("center"), bd=None)
    ws.row_dimensions[2].height = sub_height
    ws.row_dimensions[3].height = 8


def header(ws, row, cells, height=24):
    """cells: [(first_col, last_col, text, align)]"""
    for first, last, text, h in cells:
        span(ws, first, last, row, text, f=font(10, bold=True, color=WHITE), fl=fill(INK), a=al(h))
    ws.row_dimensions[row].height = height


def section(ws, row, last_col, text):
    span(ws, "A", last_col, row, text, f=font(11, bold=True, color=INK), fl=fill(LIGHT), a=al())
    ws.row_dimensions[row].height = 24


def setup(ws, widths, tab=INK, landscape=False):
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
    ws.oddFooter.center.text = "PZ Nigeria Limited - &A - Page &P of &N"
    ws.oddFooter.center.size = 8


# ---------------------------------------------------------------- charts (black, grey, white)
def rich(size=900, color=TEXT, bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT), sz=1200, b=True, solidFill=TEXT)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def grey_bars(ws, anchor, title, values, cats, fmt, vmax, major, width=16.0, height=8.0, horizontal=False):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.title = chart_title(title)
    ch.add_data(values, titles_from_data=False)
    ch.set_categories(cats)
    ch.gapWidth = 60
    ch.legend = None
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill="595959", ln=LineProperties(noFill=True))
    ch.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                                  showLegendKey=False, showPercent=False)
    ch.dataLabels.numFmt = fmt
    ch.dataLabels.position = "outEnd"
    ch.dataLabels.txPr = rich(900, TEXT, bold=True)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.txPr = rich(900, TEXT)
    ch.y_axis.txPr = rich(800, MUTED)
    ch.y_axis.numFmt = fmt
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = vmax
    ch.y_axis.majorUnit = major
    ch.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525)))
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="A6A6A6", w=9525))
    if horizontal:
        ch.x_axis.scaling.orientation = "maxMin"
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="BFBFBF", w=9525))
    ch.width, ch.height = width, height
    ws.add_chart(ch, anchor)


# ================================================================ workbook
wb = Workbook()
wb._fonts = IndexedList([Font(name=FONT, size=10, family=2)])
wb._named_styles["Normal"].font = Font(name=FONT, size=10, family=2)
cells = {}

# ---------------------------------------------------------------- 1. Company Profile
ws = wb.active
ws.title = "Company Profile"
setup(ws, {"A": 34, "B": 96})
sheet_title(ws, "B", "PZ NIGERIA LIMITED: COMPANY PROFILE",
            "The company facts behind the PESTLE, SWOT and industry analysis in this workbook")
header(ws, 4, [("A", "A", "Item", "left"), ("B", "B", "Details", "left")])
for i, (label, text) in enumerate(COMPANY):
    r = 5 + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", label, f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", text, fl=band)
    ws.row_dimensions[r].height = 22
r = 5 + len(COMPANY) + 1
section(ws, r, "B", "What is in this workbook")
contents = [
    ("PESTLE", "Political, Economic, Social, Technological, Legal and Environmental factors, "
               "each scored Impact × Likelihood, with a summary and chart per factor"),
    ("SWOT", "Strengths, Weaknesses, Opportunities and Threats in a 2 × 2 matrix"),
    ("SWOT Scoring", "The SWOT points weighted and rated (IFE and EFE matrices) and the "
                     "strategy they point to"),
    ("Industry Analysis", "Industry overview, Porter's Five Forces and the main competitors"),
    ("Summary", "The key results from every sheet, and the recommendations"),
]
for i, (name, text) in enumerate(contents):
    rr = r + 1 + i
    put(ws, f"A{rr}", name, f=font(10, bold=True))
    put(ws, f"B{rr}", text)
    ws.row_dimensions[rr].height = 30
rr = r + len(contents) + 2
ws.merge_cells(f"A{rr}:B{rr}")
put(ws, f"A{rr}", "Blue figures are inputs (ratings and weights) you can change. Black figures are formulas: "
                  "change any blue figure and every score, summary, chart and conclusion updates.",
    f=font(9, italic=True, color=MUTED), bd=None)
ws.row_dimensions[rr].height = 28

# ---------------------------------------------------------------- 2. PESTLE
ws = wb.create_sheet("PESTLE")
setup(ws, {"A": 16, "B": 36, "C": 54, "D": 13, "E": 11, "F": 12, "G": 14, "H": 11})
sheet_title(ws, "H", "PESTLE ANALYSIS: PZ NIGERIA LIMITED",
            "Impact and Likelihood are rated 1 (low) to 5 (high). Score = Impact × Likelihood (1 to 25). "
            "Priority: High 16 and above, Medium 9 to 15, Low below 9.")
header(ws, 4, [("A", "A", "Factor", "left"), ("B", "B", "Issue", "left"),
               ("C", "C", "Effect on PZ Nigeria", "left"), ("D", "D", "Type", "center"),
               ("E", "E", "Impact (1-5)", "center"), ("F", "F", "Likelihood (1-5)", "center"),
               ("G", "G", "Score", "center"), ("H", "H", "Priority", "center")], height=30)
P_FIRST = 5
P_LAST = P_FIRST + len(PESTLE) - 1
for i, (factor, issue, effect, kind, impact, likely) in enumerate(PESTLE):
    r = P_FIRST + i
    band = fill(BAND) if FACTORS.index(factor) % 2 else fill(WHITE)
    put(ws, f"A{r}", factor, f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", issue, fl=band)
    put(ws, f"C{r}", effect, fl=band)
    put(ws, f"D{r}", kind, fl=band, a=al("center"))
    put(ws, f"E{r}", impact, f=font(10, color=INPUT), fl=band, a=al("center"))
    put(ws, f"F{r}", likely, f=font(10, color=INPUT), fl=band, a=al("center"))
    put(ws, f"G{r}", f"=E{r}*F{r}", f=font(10, bold=True), fl=band, a=al("center"))
    put(ws, f"H{r}", f'=IF(G{r}>=16,"High",IF(G{r}>=9,"Medium","Low"))', fl=band, a=al("center"))
    ws.row_dimensions[r].height = 30
r = P_LAST + 1
span(ws, "A", "F", r, "Average score (all issues)", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
put(ws, f"G{r}", f"=AVERAGE(G{P_FIRST}:G{P_LAST})", f=font(10, bold=True), fl=fill(LIGHT), fmt="0.0",
    a=al("center"), bd=TOP_RULE)
put(ws, f"H{r}", f'=IF(G{r}>=16,"High",IF(G{r}>=9,"Medium","Low"))', f=font(10, bold=True), fl=fill(LIGHT),
    a=al("center"), bd=TOP_RULE)
ws.row_dimensions[r].height = 22
P_AVG = r

# Summary by factor
S_HEAD = P_AVG + 2
section(ws, S_HEAD, "H", "PESTLE summary by factor")
header(ws, S_HEAD + 1, [("A", "A", "Factor", "left"), ("B", "B", "Number of issues (COUNTIF)", "center"),
                        ("C", "C", "Average score (AVERAGEIF)", "center"),
                        ("D", "D", "Highest score", "center"), ("E", "E", "High priority", "center"),
                        ("F", "F", "Threats", "center"), ("G", "G", "Opportunities", "center"),
                        ("H", "H", "Rank", "center")], height=30)
S_FIRST = S_HEAD + 2
S_LAST = S_FIRST + len(FACTORS) - 1
rngA = f"$A${P_FIRST}:$A${P_LAST}"
rngD = f"$D${P_FIRST}:$D${P_LAST}"
rngG = f"$G${P_FIRST}:$G${P_LAST}"
rngH = f"$H${P_FIRST}:$H${P_LAST}"
for i, factor in enumerate(FACTORS):
    r = S_FIRST + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", factor, f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", f"=COUNTIF({rngA},A{r})", fl=band, a=al("center"))
    put(ws, f"C{r}", f"=AVERAGEIF({rngA},A{r},{rngG})", f=font(10, bold=True), fl=band, fmt="0.0", a=al("center"))
    # MAXIFS is not in Excel 2013; SUMPRODUCT evaluates the array without Ctrl+Shift+Enter.
    put(ws, f"D{r}", f"=SUMPRODUCT(MAX(({rngA}=A{r})*{rngG}))", fl=band, a=al("center"))
    put(ws, f"E{r}", f'=COUNTIFS({rngA},A{r},{rngH},"High")', fl=band, a=al("center"))
    put(ws, f"F{r}", f'=COUNTIFS({rngA},A{r},{rngD},"Threat")', fl=band, a=al("center"))
    put(ws, f"G{r}", f'=COUNTIFS({rngA},A{r},{rngD},"Opportunity")', fl=band, a=al("center"))
    put(ws, f"H{r}", f"=RANK(C{r},$C${S_FIRST}:$C${S_LAST})", fl=band, a=al("center"))
    ws.row_dimensions[r].height = 21
r = S_LAST + 1
put(ws, f"A{r}", "All factors", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
for col, formula, fmt in [("B", f"=SUM(B{S_FIRST}:B{S_LAST})", None),
                          ("C", f"=AVERAGE(G{P_FIRST}:G{P_LAST})", "0.0"),
                          ("D", f"=MAX(G{P_FIRST}:G{P_LAST})", None),
                          ("E", f"=SUM(E{S_FIRST}:E{S_LAST})", None),
                          ("F", f"=SUM(F{S_FIRST}:F{S_LAST})", None),
                          ("G", f"=SUM(G{S_FIRST}:G{S_LAST})", None),
                          ("H", "", None)]:
    put(ws, f"{col}{r}", formula or None, f=font(10, bold=True), fl=fill(LIGHT), fmt=fmt, a=al("center"), bd=TOP_RULE)
ws.row_dimensions[r].height = 22
S_TOTAL = r

# Key findings
K_HEAD = S_TOTAL + 2
section(ws, K_HEAD, "H", "Key findings")
header(ws, K_HEAD + 1, [("A", "B", "Finding", "left"), ("C", "F", "Result", "left"),
                        ("G", "H", "Excel functions", "center")])
avg_rng = f"$C${S_FIRST}:$C${S_LAST}"
fac_rng = f"$A${S_FIRST}:$A${S_LAST}"
K_FIRST = K_HEAD + 2
findings = [
    ("Most important factor", f"=INDEX({fac_rng},MATCH(MAX({avg_rng}),{avg_rng},0))&\" (average score \"&TEXT(MAX({avg_rng}),\"0.0\")&\")\"",
     "INDEX, MATCH, MAX"),
    ("Least important factor", f"=INDEX({fac_rng},MATCH(MIN({avg_rng}),{avg_rng},0))&\" (average score \"&TEXT(MIN({avg_rng}),\"0.0\")&\")\"",
     "INDEX, MATCH, MIN"),
    ("Highest-scoring issue", f"=INDEX($B${P_FIRST}:$B${P_LAST},MATCH(MAX({rngG}),{rngG},0))&\" (score \"&MAX({rngG})&\" of 25)\"",
     "INDEX, MATCH, MAX"),
    ("High-priority issues", f"=COUNTIF({rngH},\"High\")&\" of \"&COUNTA({rngA})&\" issues\"",
     "COUNTIF, COUNTA"),
    ("Threats and opportunities", f"=COUNTIF({rngD},\"Threat\")&\" threats, \"&COUNTIF({rngD},\"Opportunity\")&\" opportunities\"",
     "COUNTIF"),
]
for i, (label, formula, funcs) in enumerate(findings):
    r = K_FIRST + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    span(ws, "A", "B", r, label, f=font(10, bold=True), fl=band)
    span(ws, "C", "F", r, formula, fl=band)
    span(ws, "G", "H", r, funcs, f=font(9, color=MUTED), fl=band, a=al("center"))
    ws.row_dimensions[r].height = 21
K_LAST = K_FIRST + len(findings) - 1

CH = K_LAST + 2
section(ws, CH, "H", "Chart: average PESTLE score by factor")
grey_bars(ws, f"A{CH + 1}", "Average PESTLE Score by Factor (out of 25)",
          Reference(ws, min_col=3, min_row=S_FIRST, max_row=S_LAST),
          Reference(ws, min_col=1, min_row=S_FIRST, max_row=S_LAST), "0.0", 25, 5, width=20, height=8)
for rr in range(CH + 1, CH + 18):
    ws.row_dimensions[rr].height = 15
ws.freeze_panes = "A5"
cells["pestle"] = {"first": P_FIRST, "last": P_LAST, "average": P_AVG,
                   "summary_first": S_FIRST, "summary_last": S_LAST, "summary_total": S_TOTAL,
                   "findings_first": K_FIRST, "findings_last": K_LAST}

# ---------------------------------------------------------------- 3. SWOT
ws = wb.create_sheet("SWOT")
setup(ws, {"A": 13, "B": 66, "C": 66})
sheet_title(ws, "C", "SWOT ANALYSIS: PZ NIGERIA LIMITED",
            "Strengths and weaknesses are inside the company; opportunities and threats come from outside it")
put(ws, "A4", "", fl=fill(WHITE), bd=None)
put(ws, "B4", "HELPFUL", f=font(11, bold=True, color=WHITE), fl=fill(INK), a=al("center"))
put(ws, "C4", "HARMFUL", f=font(11, bold=True, color=WHITE), fl=fill(INK), a=al("center"))
ws.row_dimensions[4].height = 24
ROWS = max(len(STRENGTHS), len(WEAKNESSES), len(OPPORTUNITIES), len(THREATS))
blocks = [("INTERNAL", 5, [("B", "STRENGTHS", STRENGTHS), ("C", "WEAKNESSES", WEAKNESSES)]),
          ("EXTERNAL", 5 + ROWS + 1, [("B", "OPPORTUNITIES", OPPORTUNITIES), ("C", "THREATS", THREATS)])]
swot_ranges = {}
for axis, top, quads in blocks:
    ws.merge_cells(f"A{top}:A{top + ROWS}")
    put(ws, f"A{top}", axis, f=font(11, bold=True, color=WHITE), fl=fill(INK),
        a=Alignment(horizontal="center", vertical="center", text_rotation=90))
    for rr in range(top + 1, top + ROWS + 1):
        ws[f"A{rr}"].border = BOX
    ws.row_dimensions[top].height = 24
    for col, name, points in quads:
        put(ws, f"{col}{top}", name, f=font(11, bold=True), fl=fill(LIGHT))
        for k in range(ROWS):
            rr = top + 1 + k
            text = f"{k + 1}. {points[k][0]}" if k < len(points) else None
            put(ws, f"{col}{rr}", text)
            ws.row_dimensions[rr].height = 24
        swot_ranges[name] = (col, top + 1, top + ROWS)
C_HEAD = 5 + 2 * (ROWS + 1) + 1
header(ws, C_HEAD, [("A", "B", "SWOT count", "left"), ("C", "C", "Number of points (COUNTA)", "center")])
for i, name in enumerate(["STRENGTHS", "WEAKNESSES", "OPPORTUNITIES", "THREATS"]):
    col, a, b = swot_ranges[name]
    r = C_HEAD + 1 + i
    span(ws, "A", "B", r, name.title(), f=font(10, bold=True), fl=fill(BAND) if i % 2 else fill(WHITE))
    put(ws, f"C{r}", f"=COUNTA({col}{a}:{col}{b})", fl=fill(BAND) if i % 2 else fill(WHITE), a=al("center"))
r = C_HEAD + 5
span(ws, "A", "B", r, "Total", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
put(ws, f"C{r}", f"=SUM(C{C_HEAD + 1}:C{C_HEAD + 4})", f=font(10, bold=True), fl=fill(LIGHT), a=al("center"), bd=TOP_RULE)
cells["swot"] = {k: {"col": v[0], "first": v[1], "last": v[2]} for k, v in swot_ranges.items()}

# ---------------------------------------------------------------- 4. SWOT Scoring (IFE and EFE)
ws = wb.create_sheet("SWOT Scoring")
setup(ws, {"A": 15, "B": 66, "C": 12, "D": 12, "E": 16})
sheet_title(ws, "E", "SWOT SCORING: IFE AND EFE MATRICES",
            "Weight = importance (each matrix adds up to 1.00). IFE rating: strengths 3-4, weaknesses 1-2. "
            "EFE rating: how well PZ responds, 1 (poor) to 4 (superior). Weighted score = Weight × Rating; "
            "2.50 is average.", sub_height=44)


def factor_matrix(top, title, groups, good_text, bad_text):
    section(ws, top, "E", title)
    header(ws, top + 1, [("A", "A", "Type", "left"), ("B", "B", "Key factor", "left"),
                         ("C", "C", "Weight", "center"), ("D", "D", "Rating", "center"),
                         ("E", "E", "Weighted score", "center")])
    r = top + 2
    first = r
    for kind, points in groups:
        for k, (text, weight, rating) in enumerate(points):
            band = fill(BAND) if kind in ("Weakness", "Threat") else fill(WHITE)
            put(ws, f"A{r}", kind, f=font(10, bold=True), fl=band)
            put(ws, f"B{r}", text, fl=band)
            put(ws, f"C{r}", weight, f=font(10, color=INPUT), fl=band, fmt="0.00", a=al("center"))
            put(ws, f"D{r}", rating, f=font(10, color=INPUT), fl=band, a=al("center"))
            put(ws, f"E{r}", f"=C{r}*D{r}", f=font(10, bold=True), fl=band, fmt="0.00", a=al("center"))
            ws.row_dimensions[r].height = 20
            r += 1
    last = r - 1
    span(ws, "A", "B", r, "Total", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
    put(ws, f"C{r}", f"=SUM(C{first}:C{last})", f=font(10, bold=True), fl=fill(LIGHT), fmt="0.00", a=al("center"), bd=TOP_RULE)
    put(ws, f"D{r}", None, fl=fill(LIGHT), bd=TOP_RULE)
    put(ws, f"E{r}", f"=SUM(E{first}:E{last})", f=font(11, bold=True), fl=fill(LIGHT), fmt="0.00", a=al("center"), bd=TOP_RULE)
    total = r
    span(ws, "A", "B", r + 1, "Weights check", f=font(10, bold=True))
    span(ws, "C", "E", r + 1, f'=IF(ROUND(C{total},2)=1,"OK: weights add up to 1.00","Check: weights must add up to 1.00")')
    span(ws, "A", "B", r + 2, "Result", f=font(10, bold=True))
    span(ws, "C", "E", r + 2, f'=IF(E{total}>=2.5,"{good_text}","{bad_text}")', f=font(10, bold=True))
    ws.row_dimensions[r + 2].height = 30
    return first, last, total


IFE_FIRST, IFE_LAST, IFE_TOTAL = factor_matrix(
    4, "Internal Factor Evaluation (IFE): strengths and weaknesses",
    [("Strength", STRENGTHS), ("Weakness", WEAKNESSES)],
    "Above average: strengths outweigh weaknesses", "Below average: weaknesses outweigh strengths")
EFE_TOP = IFE_TOTAL + 4
EFE_FIRST, EFE_LAST, EFE_TOTAL = factor_matrix(
    EFE_TOP, "External Factor Evaluation (EFE): opportunities and threats",
    [("Opportunity", OPPORTUNITIES), ("Threat", THREATS)],
    "Above average: PZ responds well to opportunities and threats",
    "Below average: PZ must respond better to opportunities and threats")

R_HEAD = EFE_TOTAL + 4
section(ws, R_HEAD, "E", "Strategic position")
header(ws, R_HEAD + 1, [("A", "B", "Measure", "left"), ("C", "D", "Result", "center"),
                        ("E", "E", "Excel functions", "center")])
ife_a, ife_e = f"$A${IFE_FIRST}:$A${IFE_LAST}", f"$E${IFE_FIRST}:$E${IFE_LAST}"
efe_a, efe_e = f"$A${EFE_FIRST}:$A${EFE_LAST}", f"$E${EFE_FIRST}:$E${EFE_LAST}"
R_FIRST = R_HEAD + 2
# IE matrix: IFE across (3+ strong, 2-2.99 average, below 2 weak), EFE down (3+ high, 2-2.99
# medium, below 2 low); cells I to IX, row by row.
cell_no = (f"(IF(E{EFE_TOTAL}>=3,1,IF(E{EFE_TOTAL}>=2,2,3))-1)*3"
           f"+IF(E{IFE_TOTAL}>=3,1,IF(E{IFE_TOTAL}>=2,2,3))")
position = [
    ("Strengths score", f'=SUMIF({ife_a},"Strength",{ife_e})', "0.00", "SUMIF"),
    ("Weaknesses score", f'=SUMIF({ife_a},"Weakness",{ife_e})', "0.00", "SUMIF"),
    ("Opportunities score", f'=SUMIF({efe_a},"Opportunity",{efe_e})', "0.00", "SUMIF"),
    ("Threats score", f'=SUMIF({efe_a},"Threat",{efe_e})', "0.00", "SUMIF"),
    ("IFE total (internal)", f"=E{IFE_TOTAL}", "0.00", "SUM"),
    ("EFE total (external)", f"=E{EFE_TOTAL}", "0.00", "SUM"),
    ("IE matrix cell", f'=CHOOSE({cell_no},"I","II","III","IV","V","VI","VII","VIII","IX")', None, "IF, CHOOSE"),
    ("Strategy", f'=CHOOSE({cell_no},"Grow and build","Grow and build","Hold and maintain","Grow and build",'
                 f'"Hold and maintain","Harvest or divest","Hold and maintain","Harvest or divest","Harvest or divest")',
     None, "IF, CHOOSE"),
    ("Suggested strategies", f'=IF(C{R_FIRST + 7}="Hold and maintain","Market penetration and product development",'
                             f'IF(C{R_FIRST + 7}="Grow and build","Market penetration, market development and product development",'
                             f'"Cut costs or sell weak businesses"))', None, "IF"),
]
for i, (label, formula, fmt, funcs) in enumerate(position):
    r = R_FIRST + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    span(ws, "A", "B", r, label, f=font(10, bold=True), fl=band)
    span(ws, "C", "D", r, formula, f=font(10, bold=True), fl=band, fmt=fmt, a=al("center"))
    put(ws, f"E{r}", funcs, f=font(9, color=MUTED), fl=band, a=al("center"))
    ws.row_dimensions[r].height = 30 if i == len(position) - 1 else 21
R_LAST = R_FIRST + len(position) - 1
CH = R_LAST + 2
section(ws, CH, "E", "Chart: weighted SWOT scores")
grey_bars(ws, f"A{CH + 1}", "Weighted SWOT Scores",
          Reference(ws, min_col=3, min_row=R_FIRST, max_row=R_FIRST + 3),
          Reference(ws, min_col=1, min_row=R_FIRST, max_row=R_FIRST + 3), "0.00", 2.5, 0.5, width=18, height=7.5)
for rr in range(CH + 1, CH + 17):
    ws.row_dimensions[rr].height = 15
cells["scoring"] = {"ife_first": IFE_FIRST, "ife_last": IFE_LAST, "ife_total": IFE_TOTAL,
                    "efe_first": EFE_FIRST, "efe_last": EFE_LAST, "efe_total": EFE_TOTAL,
                    "position_first": R_FIRST, "position_last": R_LAST}

# ---------------------------------------------------------------- 5. Industry Analysis
ws = wb.create_sheet("Industry Analysis")
setup(ws, {"A": 32, "B": 30, "C": 26, "D": 62})
sheet_title(ws, "D", "INDUSTRY ANALYSIS: NIGERIAN FMCG AND CONSUMER ELECTRICALS",
            "Industry overview, Porter's Five Forces (each force rated 1 = weak to 5 = strong) and PZ's main competitors")
section(ws, 4, "D", "Industry overview")
for i, (label, text) in enumerate(OVERVIEW):
    r = 5 + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", label, f=font(10, bold=True), fl=band)
    span(ws, "B", "D", r, text, fl=band)
    ws.row_dimensions[r].height = 30
F_HEAD = 5 + len(OVERVIEW) + 1
section(ws, F_HEAD, "D", "Porter's Five Forces")
header(ws, F_HEAD + 1, [("A", "A", "Force", "left"), ("B", "B", "Strength (1-5)", "center"),
                        ("C", "C", "Level", "center"), ("D", "D", "Why", "left")])
F_FIRST = F_HEAD + 2
for i, (force, rating, why) in enumerate(FORCES):
    r = F_FIRST + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", force, f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", rating, f=font(10, color=INPUT), fl=band, a=al("center"))
    put(ws, f"C{r}", f'=IF(B{r}>=4,"High",IF(B{r}>=3,"Medium","Low"))', f=font(10, bold=True), fl=band, a=al("center"))
    put(ws, f"D{r}", why, fl=band)
    ws.row_dimensions[r].height = 30
F_LAST = F_FIRST + len(FORCES) - 1
r = F_LAST + 1
put(ws, f"A{r}", "Average strength", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
put(ws, f"B{r}", f"=AVERAGE(B{F_FIRST}:B{F_LAST})", f=font(10, bold=True), fl=fill(LIGHT), fmt="0.0", a=al("center"), bd=TOP_RULE)
put(ws, f"C{r}", f'=IF(B{r}>=3.5,"High",IF(B{r}>=2.5,"Medium","Low"))', f=font(10, bold=True), fl=fill(LIGHT), a=al("center"), bd=TOP_RULE)
put(ws, f"D{r}", f'=IF(B{r}>=3.5,"Strong competitive pressure: profits are hard to earn",'
                 f'IF(B{r}>=2.5,"Moderate competitive pressure","Weak competitive pressure: an attractive industry"))',
    f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
F_AVG = r
frng, brng, crng = f"$A${F_FIRST}:$A${F_LAST}", f"$B${F_FIRST}:$B${F_LAST}", f"$C${F_FIRST}:$C${F_LAST}"
extra = [
    ("Strongest force", f"=INDEX({frng},MATCH(MAX({brng}),{brng},0))"),
    ("Weakest force", f"=INDEX({frng},MATCH(MIN({brng}),{brng},0))"),
    ("Forces rated High", f'=COUNTIF({crng},"High")&" of "&COUNTA({frng})'),
]
for i, (label, formula) in enumerate(extra):
    r = F_AVG + 1 + i
    put(ws, f"A{r}", label, f=font(10, bold=True))
    span(ws, "B", "D", r, formula)
    ws.row_dimensions[r].height = 21
F_EXTRA_LAST = F_AVG + len(extra)

M_HEAD = F_EXTRA_LAST + 2
section(ws, M_HEAD, "D", "Main competitors")
header(ws, M_HEAD + 1, [("A", "A", "Competitor", "left"), ("B", "B", "Key brands", "left"),
                        ("C", "C", "Competes with PZ in", "left"), ("D", "D", "Position", "left")])
M_FIRST = M_HEAD + 2
for i, row in enumerate(COMPETITORS):
    r = M_FIRST + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    for col, value in zip("ABCD", row):
        put(ws, f"{col}{r}", value, f=font(10, bold=(col == "A")), fl=band)
    ws.row_dimensions[r].height = 30
M_LAST = M_FIRST + len(COMPETITORS) - 1
r = M_LAST + 1
put(ws, f"A{r}", "Number of competitors listed", f=font(10, bold=True), fl=fill(LIGHT), bd=TOP_RULE)
span(ws, "B", "D", r, f"=COUNTA(A{M_FIRST}:A{M_LAST})", f=font(10, bold=True), fl=fill(LIGHT), a=al("left"), bd=TOP_RULE)
CH = r + 2
section(ws, CH, "D", "Chart: strength of the five forces")
grey_bars(ws, f"A{CH + 1}", "Porter's Five Forces: Strength (1 to 5)",
          Reference(ws, min_col=2, min_row=F_FIRST, max_row=F_LAST),
          Reference(ws, min_col=1, min_row=F_FIRST, max_row=F_LAST), "0", 5, 1, width=20, height=8)
for rr in range(CH + 1, CH + 18):
    ws.row_dimensions[rr].height = 15
cells["industry"] = {"overview_first": 5, "overview_last": 4 + len(OVERVIEW),
                     "forces_first": F_FIRST, "forces_last": F_LAST, "forces_average": F_AVG,
                     "competitors_first": M_FIRST, "competitors_last": M_LAST}

# ---------------------------------------------------------------- 6. Summary
ws = wb.create_sheet("Summary")
setup(ws, {"A": 36, "B": 92})
sheet_title(ws, "B", "SUMMARY AND RECOMMENDATIONS",
            "The key results, linked to the other sheets: they update when any rating or weight changes")
header(ws, 4, [("A", "A", "Measure", "left"), ("B", "B", "Result", "left")])
p, sc, ind = cells["pestle"], cells["scoring"], cells["industry"]
summary = [
    ("PESTLE: most important factor", f"=PESTLE!C{p['findings_first']}"),
    ("PESTLE: highest-scoring issue", f"=PESTLE!C{p['findings_first'] + 2}"),
    ("PESTLE: high-priority issues", f"=PESTLE!C{p['findings_first'] + 3}"),
    ("SWOT: points listed", "=SWOT!C" + str(C_HEAD + 5) + '&" (strengths "&SWOT!C' + str(C_HEAD + 1)
     + '&", weaknesses "&SWOT!C' + str(C_HEAD + 2) + '&", opportunities "&SWOT!C' + str(C_HEAD + 3)
     + '&", threats "&SWOT!C' + str(C_HEAD + 4) + '&")"'),
    ("IFE total (internal strength)", f"=TEXT('SWOT Scoring'!E{sc['ife_total']},\"0.00\")&\" out of 4. \"&'SWOT Scoring'!C{sc['ife_total'] + 2}"),
    ("EFE total (external response)", f"=TEXT('SWOT Scoring'!E{sc['efe_total']},\"0.00\")&\" out of 4. \"&'SWOT Scoring'!C{sc['efe_total'] + 2}"),
    ("Strategy (IE matrix)", f"='SWOT Scoring'!C{sc['position_first'] + 7}&\" (cell \"&'SWOT Scoring'!C{sc['position_first'] + 6}&\"): \"&'SWOT Scoring'!C{sc['position_first'] + 8}"),
    ("Five Forces: average strength", f"=TEXT('Industry Analysis'!B{ind['forces_average']},\"0.0\")&\" out of 5. \"&'Industry Analysis'!D{ind['forces_average']}"),
    ("Five Forces: strongest force", f"='Industry Analysis'!B{ind['forces_average'] + 1}"),
]
for i, (label, formula) in enumerate(summary):
    r = 5 + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", label, f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", formula, fl=band)
    ws.row_dimensions[r].height = 24
REC_HEAD = 5 + len(summary) + 1
section(ws, REC_HEAD, "B", "Recommendations")
header(ws, REC_HEAD + 1, [("A", "A", "Recommendation", "left"), ("B", "B", "Why", "left")])
for i, (rec, why) in enumerate(RECOMMENDATIONS):
    r = REC_HEAD + 2 + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    put(ws, f"A{r}", f"{i + 1}. {rec}", f=font(10, bold=True), fl=band)
    put(ws, f"B{r}", why, fl=band)
    ws.row_dimensions[r].height = 30
r = REC_HEAD + 2 + len(RECOMMENDATIONS) + 1
ws.merge_cells(f"A{r}:B{r}")
put(ws, f"A{r}", "Ratings and weights are the analyst's judgement, based on the company information and "
                 "public news up to 2025. Check the latest figures (exchange rate, inflation) before presenting.",
    f=font(9, italic=True, color=MUTED), bd=None)
ws.row_dimensions[r].height = 28
cells["summary"] = {"first": 5, "last": 4 + len(summary), "rec_first": REC_HEAD + 2,
                    "rec_last": REC_HEAD + 1 + len(RECOMMENDATIONS)}

for sheet in wb.worksheets:
    sheet.sheet_view.selection[0].activeCell = "A1"
    sheet.sheet_view.selection[0].sqref = "A1"
wb.active = 0

wb.properties.title = "PZ Nigeria Limited: PESTLE, SWOT and Industry Analysis"
wb.properties.subject = "PESTLE, SWOT (IFE/EFE) and Porter's Five Forces analysis of PZ Nigeria Limited"
wb.properties.creator = "Marketing Department"
wb.properties.lastModifiedBy = "Marketing Department"
wb.properties.keywords = "PZ Nigeria, PESTLE, SWOT, Five Forces, industry analysis"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
Path(str(OUT) + ".cells.json").write_text(json.dumps(cells, indent=2))
print(f"wrote {OUT}")
