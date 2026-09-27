#!/usr/bin/env python3
"""Read the finished workbook's calculated values and write them to JSON for the slides.

Every rating, score and conclusion in the deck comes from here, so the slides always agree
with the Excel file.

Usage: python export_slide_data.py PZ_Nigeria_Analysis.xlsx cells.json slide_data.json
"""
import json
import re
import sys

from openpyxl import load_workbook


def rows(ws, first, last, cols):
    return [[ws[f"{c}{r}"].value for c in cols] for r in range(first, last + 1)]


def main(xlsx, cells_path, out):
    wb = load_workbook(xlsx, data_only=True)
    cells = json.load(open(cells_path))
    data = {}

    ws = wb["Company Profile"]
    data["company"] = {label: text for label, text in rows(ws, 5, 13, "AB")}

    ws, p = wb["PESTLE"], cells["pestle"]
    keys = ["factor", "issue", "effect", "type", "impact", "likelihood", "score", "priority"]
    data["pestle"] = [dict(zip(keys, r)) for r in rows(ws, p["first"], p["last"], "ABCDEFGH")]
    keys = ["factor", "count", "average", "highest", "high", "threats", "opportunities", "rank"]
    data["pestleSummary"] = [dict(zip(keys, r)) for r in rows(ws, p["summary_first"], p["summary_last"], "ABCDEFGH")]
    data["pestleAverage"] = ws[f"G{p['average']}"].value
    data["pestleFindings"] = [ws[f"C{r}"].value for r in range(p["findings_first"], p["findings_last"] + 1)]

    ws = wb["SWOT"]
    data["swot"] = {}
    for name, rng in cells["swot"].items():
        pts = [ws[f"{rng['col']}{r}"].value for r in range(rng["first"], rng["last"] + 1)]
        data["swot"][name.lower()] = [re.sub(r"^\d+\.\s*", "", t) for t in pts if t]

    ws, s = wb["SWOT Scoring"], cells["scoring"]
    keys = ["type", "factor", "weight", "rating", "weighted"]
    data["ife"] = [dict(zip(keys, r)) for r in rows(ws, s["ife_first"], s["ife_last"], "ABCDE")]
    data["efe"] = [dict(zip(keys, r)) for r in rows(ws, s["efe_first"], s["efe_last"], "ABCDE")]
    for name in ("ife", "efe"):
        total = s[f"{name}_total"]
        data[f"{name}Total"] = ws[f"E{total}"].value
        data[f"{name}Result"] = ws[f"C{total + 2}"].value
    pos = [ws[f"C{r}"].value for r in range(s["position_first"], s["position_last"] + 1)]
    data["position"] = dict(zip(["strengths", "weaknesses", "opportunities", "threats", "ife", "efe",
                                 "ieCell", "strategy", "suggested"], pos))

    ws, ind = wb["Industry Analysis"], cells["industry"]
    data["overview"] = {label: text for label, text, *_ in rows(ws, ind["overview_first"], ind["overview_last"], "AB")}
    keys = ["force", "rating", "level", "why"]
    data["forces"] = [dict(zip(keys, r)) for r in rows(ws, ind["forces_first"], ind["forces_last"], "ABCD")]
    a = ind["forces_average"]
    data["forcesAverage"] = ws[f"B{a}"].value
    data["forcesLevel"] = ws[f"C{a}"].value
    data["forcesVerdict"] = ws[f"D{a}"].value
    data["strongestForce"] = ws[f"B{a + 1}"].value
    data["weakestForce"] = ws[f"B{a + 2}"].value
    data["forcesHigh"] = ws[f"B{a + 3}"].value
    keys = ["name", "brands", "competes", "position"]
    data["competitors"] = [dict(zip(keys, r)) for r in rows(ws, ind["competitors_first"], ind["competitors_last"], "ABCD")]

    ws, sm = wb["Summary"], cells["summary"]
    data["recommendations"] = [
        {"title": re.sub(r"^\d+\.\s*", "", t), "why": w}
        for t, w in rows(ws, sm["rec_first"], sm["rec_last"], "AB")
    ]

    missing = [k for k, v in data.items() if v is None]
    if missing:
        sys.exit(f"workbook has no cached value for: {missing}")
    with open(out, "w") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
