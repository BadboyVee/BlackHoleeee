#!/usr/bin/env python3
"""Read the finished workbook's calculated values and write them to JSON for the slides.

Every number in the deck comes from here, so the slides always agree with the Excel file.

Usage: python export_slide_data.py Marketing_Department_Budget.xlsx slide_data.json
"""
import json
import sys

from openpyxl import load_workbook

UNIT_META = {
    "Advertising": {"title": "Advertising (Ads)", "color": "2A78D6", "icon": "FaBullhorn",
                    "covers": "Paid media campaigns on TV, radio, print, outdoor and online channels"},
    "Marketing": {"title": "Marketing", "color": "EB6834", "icon": "FaBullseye",
                  "covers": "Market research, sales promotions, product launches and trade shows"},
    "Public Relations": {"title": "Public Relations (PR)", "color": "1BAF7A", "icon": "FaHandshake",
                         "covers": "Media relations, corporate events, sponsorships and community relations"},
    "e-Business": {"title": "e-Business", "color": "EDA100", "icon": "FaShoppingCart",
                   "covers": "Company website, e-commerce platform, search and social media marketing"},
}


def main(xlsx, out):
    ws = load_workbook(xlsx, data_only=True)["Marketing Budget"]
    header = [ws.cell(4, c).value for c in range(1, 8)]
    assert header[:1] == ["Budget Item"] and header[5:] == ["Total", "Average per Unit"], header

    items = []
    r = 5
    while ws.cell(r, 1).value != "Total Budget":
        items.append({
            "name": ws.cell(r, 1).value,
            "amounts": [ws.cell(r, c).value for c in range(2, 6)],
            "total": ws.cell(r, 6).value,
            "average": ws.cell(r, 7).value,
        })
        r += 1
    total_row, avg_row, share_row = r, r + 1, r + 2
    assert ws.cell(avg_row, 1).value == "Average per Budget Item"
    assert ws.cell(share_row, 1).value == "Share of Total Budget"

    dept_total = ws.cell(total_row, 6).value
    avg_per_unit = ws.cell(total_row, 7).value
    unit_totals = [ws.cell(total_row, c).value for c in range(2, 6)]
    ranks = {t: i + 1 for i, t in enumerate(sorted(unit_totals, reverse=True))}

    units = []
    for j, name in enumerate(header[1:5]):
        unit_items = sorted(
            ({"name": it["name"], "amount": it["amounts"][j]} for it in items),
            key=lambda x: x["amount"], reverse=True)
        total = unit_totals[j]
        units.append({
            "name": name,
            **UNIT_META[name],
            "total": total,
            "share": ws.cell(share_row, 2 + j).value,
            "averagePerItem": ws.cell(avg_row, 2 + j).value,
            "rank": ranks[total],
            "vsAverage": total - avg_per_unit,
            "items": unit_items,
            "column": "BCDE"[j],
        })

    largest_item = max(items, key=lambda it: it["total"])
    data = {
        "year": "2026",
        "departmentTotal": dept_total,
        "averagePerUnit": avg_per_unit,
        "averagePerItem": ws.cell(avg_row, 6).value,
        "itemCount": len(items),
        "units": units,
        "items": [{k: it[k] for k in ("name", "total", "average")} for it in items],
        "largestItem": {"name": largest_item["name"], "total": largest_item["total"],
                        "share": largest_item["total"] / dept_total},
        "tableRange": f"A4:G{share_row}",
        "totalRow": total_row,
    }
    with open(out, "w") as fh:
        json.dump(data, fh, indent=2)
    print(f"wrote {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
