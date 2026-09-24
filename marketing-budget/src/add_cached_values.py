#!/usr/bin/env python3
"""Copy formula results from a recalculated workbook into the openpyxl-built one.

openpyxl saves formulas without results, so phone and e-mail previewers show empty cells
until Excel recalculates. This writes each formula's computed value (taken from a copy that
LibreOffice recalculated) into the original file, leaving its formatting and charts
untouched. Excel still recalculates on open (fullCalcOnLoad stays on).

Usage: python add_cached_values.py built.xlsx recalculated.xlsx output.xlsx [--ole-range A4:G15]

--ole-range sets the workbook's oleSize, the block Excel shows when the workbook is
embedded in PowerPoint as a worksheet object.
"""
import argparse
import os
import shutil
import tempfile
import zipfile

from lxml import etree
from openpyxl import load_workbook

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
Q = f"{{{NS}}}"


def sheet_parts(zf):
    """Map sheet name -> part name inside the package."""
    wb = etree.fromstring(zf.read("xl/workbook.xml"))
    rels = etree.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    targets = {r.get("Id"): r.get("Target") for r in rels}
    parts = {}
    for sheet in wb.iter(f"{Q}sheet"):
        target = targets[sheet.get(f"{{{rel_ns}}}id")].lstrip("/")
        parts[sheet.get("name")] = target if target.startswith("xl/") else f"xl/{target}"
    return parts


def fmt_number(value):
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return str(int(value))
    return repr(float(value))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("built")
    ap.add_argument("recalculated")
    ap.add_argument("output")
    ap.add_argument("--ole-range")
    args = ap.parse_args()

    values = load_workbook(args.recalculated, data_only=True)
    filled = 0
    with zipfile.ZipFile(args.built) as src:
        parts = sheet_parts(src)
        by_part = {part: name for name, part in parts.items()}
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
            out_path = tmp.name
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED) as dst:
            for info in src.infolist():
                data = src.read(info.filename)
                if info.filename in by_part:
                    ws = values[by_part[info.filename]]
                    root = etree.fromstring(data)
                    for c in root.iter(f"{Q}c"):
                        f = c.find(f"{Q}f")
                        if f is None:
                            continue
                        result = ws[c.get("r")].value
                        v = c.find(f"{Q}v")
                        if v is None:
                            v = etree.SubElement(c, f"{Q}v")
                        if result is None:
                            c.remove(v)
                            continue
                        if isinstance(result, str):
                            c.set("t", "str")
                            v.text = result
                        elif isinstance(result, bool):
                            c.set("t", "b")
                            v.text = fmt_number(result)
                        else:
                            c.attrib.pop("t", None)
                            v.text = fmt_number(result)
                        filled += 1
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                elif info.filename == "xl/workbook.xml" and args.ole_range:
                    root = etree.fromstring(data)
                    for old in root.findall(f"{Q}oleSize"):
                        root.remove(old)
                    ole = etree.Element(f"{Q}oleSize", ref=args.ole_range)
                    calc = root.find(f"{Q}calcPr")
                    # CT_Workbook order: ... definedNames, calcPr, oleSize, customWorkbookViews ...
                    anchor = calc if calc is not None else root.find(f"{Q}definedNames")
                    if anchor is None:
                        anchor = root.find(f"{Q}sheets")
                    anchor.addnext(ole)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                dst.writestr(info, data)
    shutil.move(out_path, args.output)
    os.chmod(args.output, 0o644)
    print(f"cached {filled} formula results -> {args.output}")


if __name__ == "__main__":
    main()
