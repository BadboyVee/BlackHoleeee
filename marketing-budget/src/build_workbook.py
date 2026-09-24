#!/usr/bin/env python3
"""Build Marketing_Department_Budget.xlsx: the marketing department budget by unit.

Every budget figure is an input (blue). Every total, average, share and analysis figure is a
live Excel formula (black), using only functions that exist in Excel 2013: SUM, AVERAGE, MAX,
MIN, INDEX, MATCH, COUNTA and IF.

Usage: python build_workbook.py [output.xlsx]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.legend import Legend
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
)
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.page import PageMargins

OUT = (
    Path(sys.argv[1])
    if len(sys.argv) > 1
    else Path(__file__).resolve().parent.parent / "Marketing_Department_Budget.xlsx"
)

# ---------------------------------------------------------------- data (US$, illustrative)
UNITS = ["Advertising", "Marketing", "Public Relations", "e-Business"]
UNIT_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100"]  # same identity colours as the slides
ITEMS = [
    # budget item                          Advertising Marketing  PR     e-Business
    ("Staff Salaries & Benefits",          [66000,      68000,     44000, 56000]),
    ("Media & Advertising Space",          [150000,     20000,     8000,  28000]),
    ("Events, Sponsorships & Promotions",  [12000,      70000,     38000, 8000]),
    ("Digital Platforms & Technology",     [10000,      12000,     6000,  60000]),
    ("Market Research & Analytics",        [8000,       35000,     5000,  12000]),
    ("Printing & Promotional Materials",   [28000,      25000,     12000, 2000]),
    ("Training & Travel",                  [6000,       18000,     9000,  6000]),
    ("Agency & Consultancy Fees",          [40000,      12000,     8000,  8000]),
]

# ---------------------------------------------------------------- styling
FONT = "Arial"
INK = "1B2436"        # headings, formulas
INPUT = "0000FF"      # hardcoded inputs (financial-model convention)
MUTED = "5B6475"      # notes
GRID = "C9D0DC"       # cell borders
BAND = "F3F5F9"       # alternate-row shading
TOTAL_FILL = "DCE3EE"
AVG_FILL = "EEF1F6"
WHITE = "FFFFFF"

CUR = '$#,##0_);($#,##0);"-"_)'
PCT = '0.0%_);(0.0%);"-"_)'

thin = Side(style="thin", color=GRID)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
TOTAL_BOX = Border(left=thin, right=thin, top=Side(style="medium", color=INK), bottom=thin)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def font(size=10, bold=False, italic=False, color=INK):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def style(cell, *, f=None, fl=None, fmt=None, al=None, bd=None):
    if f is not None:
        cell.font = f
    if fl is not None:
        cell.fill = fl
    if fmt is not None:
        cell.number_format = fmt
    if al is not None:
        cell.alignment = al
    if bd is not None:
        cell.border = bd


LEFT = Alignment(horizontal="left", vertical="center", indent=1)
LEFT_WRAP = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---------------------------------------------------------------- workbook
wb = Workbook()
# Arial 10 as the workbook default, so anything typed later matches the sheet.
wb._fonts = IndexedList([Font(name=FONT, size=10, family=2)])
wb._named_styles["Normal"].font = Font(name=FONT, size=10, family=2)

ws = wb.active
ws.title = "Marketing Budget"
ws.sheet_properties.tabColor = INK
ws.sheet_view.showGridLines = False
ws.sheet_view.zoomScale = 100

widths = {"A": 38, "B": 18, "C": 18, "D": 18, "E": 18, "F": 18, "G": 18}
for col, w in widths.items():
    ws.column_dimensions[col].width = w

FIRST, LAST = 5, 5 + len(ITEMS) - 1          # item rows 5..12
TOTAL, AVG, SHARE = LAST + 1, LAST + 2, LAST + 3  # 13, 14, 15
UNIT_COLS = ["B", "C", "D", "E"]

# Title block
ws.merge_cells("A1:G1")
ws["A1"] = "MARKETING DEPARTMENT BUDGET 2026"
style(ws["A1"], f=font(16, bold=True, color=WHITE), fl=fill(INK), al=Alignment(horizontal="center", vertical="center"))
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:G2")
ws["A2"] = ("Annual budget by unit: Advertising (Ads), Marketing, Public Relations (PR) "
            "and e-Business. All amounts in US$.")
style(ws["A2"], f=font(10, italic=True, color=MUTED), al=Alignment(horizontal="center", vertical="center"))
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 8

# Header row
headers = ["Budget Item", *UNITS, "Total", "Average per Unit"]
for i, text in enumerate(headers):
    c = ws.cell(row=4, column=i + 1, value=text)
    style(c, f=font(10, bold=True, color=WHITE), fl=fill(INK), bd=BOX,
          al=LEFT if i == 0 else CENTER)
ws.row_dimensions[4].height = 30

# Budget items: inputs in blue, row total (SUM) and row average (AVERAGE) as formulas
for r_off, (item, amounts) in enumerate(ITEMS):
    r = FIRST + r_off
    band = fill(BAND) if r_off % 2 else fill(WHITE)
    c = ws.cell(row=r, column=1, value=item)
    style(c, f=font(10), fl=band, bd=BOX, al=LEFT)
    for j, amount in enumerate(amounts):
        c = ws.cell(row=r, column=2 + j, value=amount)
        style(c, f=font(10, color=INPUT), fl=band, bd=BOX, fmt=CUR, al=RIGHT)
    c = ws.cell(row=r, column=6, value=f"=SUM(B{r}:E{r})")
    style(c, f=font(10, bold=True), fl=band, bd=BOX, fmt=CUR, al=RIGHT)
    c = ws.cell(row=r, column=7, value=f"=AVERAGE(B{r}:E{r})")
    style(c, f=font(10), fl=band, bd=BOX, fmt=CUR, al=RIGHT)
    ws.row_dimensions[r].height = 21

# Total row: SUM down each unit column
ws.cell(row=TOTAL, column=1, value="Total Budget")
for col in UNIT_COLS:
    ws[f"{col}{TOTAL}"] = f"=SUM({col}{FIRST}:{col}{LAST})"
ws[f"F{TOTAL}"] = f"=SUM(B{TOTAL}:E{TOTAL})"
ws[f"G{TOTAL}"] = f"=AVERAGE(B{TOTAL}:E{TOTAL})"
for col in "ABCDEFG":
    c = ws[f"{col}{TOTAL}"]
    style(c, f=font(10.5, bold=True), fl=fill(TOTAL_FILL), bd=TOTAL_BOX,
          fmt=None if col == "A" else CUR, al=LEFT if col == "A" else RIGHT)
ws.row_dimensions[TOTAL].height = 24

# Average row: AVERAGE down each column
ws.cell(row=AVG, column=1, value="Average per Budget Item")
for col in [*UNIT_COLS, "F", "G"]:
    ws[f"{col}{AVG}"] = f"=AVERAGE({col}{FIRST}:{col}{LAST})"
for col in "ABCDEFG":
    c = ws[f"{col}{AVG}"]
    style(c, f=font(10, bold=True, italic=True), fl=fill(AVG_FILL), bd=BOX,
          fmt=None if col == "A" else CUR, al=LEFT if col == "A" else RIGHT)
ws.row_dimensions[AVG].height = 22

# Share row: each unit's slice of the department total (guarded against a zero total)
ws.cell(row=SHARE, column=1, value="Share of Total Budget")
for col in UNIT_COLS:
    ws[f"{col}{SHARE}"] = f"=IF($F${TOTAL}=0,0,{col}{TOTAL}/$F${TOTAL})"
ws[f"F{SHARE}"] = f"=SUM(B{SHARE}:E{SHARE})"
ws[f"G{SHARE}"] = f"=AVERAGE(B{SHARE}:E{SHARE})"
for col in "ABCDEFG":
    c = ws[f"{col}{SHARE}"]
    style(c, f=font(10, italic=True), fl=fill(WHITE), bd=BOX,
          fmt=None if col == "A" else PCT, al=LEFT if col == "A" else RIGHT)
ws.row_dimensions[SHARE].height = 21

# Notes and colour legend
NOTES = SHARE + 2  # 17
notes = [
    ("Notes", font(10, bold=True)),
    ("• Blue figures are the budget inputs. Change any of them and every total, average, share, "
     "analysis figure and chart updates automatically.", font(9, color=MUTED)),
    ("• Black figures are formulas: SUM gives the totals and AVERAGE gives the averages. "
     "Click any black figure to see its formula in the formula bar.", font(9, color=MUTED)),
    ("• The budget figures are illustrative sample data prepared for this assignment.",
     font(9, color=MUTED)),
]
for i, (text, f) in enumerate(notes):
    r = NOTES + i
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    ws.cell(row=r, column=1, value=text)
    style(ws.cell(row=r, column=1), f=f, al=Alignment(horizontal="left", vertical="center", wrap_text=True))
    ws.row_dimensions[r].height = 17

# Budget analysis block
AN = NOTES + len(notes) + 1  # header row 22
an_headers = [("A", "A", "Budget Analysis"), ("B", "B", "Amount"),
              ("C", "E", "Unit / Budget Item"), ("F", "G", "Excel Function Used")]
for start, end, text in an_headers:
    if start != end:
        ws.merge_cells(f"{start}{AN}:{end}{AN}")
    c = ws[f"{start}{AN}"]
    c.value = text
    style(c, f=font(10, bold=True, color=WHITE), fl=fill(INK), bd=BOX,
          al=LEFT if start == "A" else CENTER)
    for col in "ABCDEFG"["ABCDEFG".index(start):"ABCDEFG".index(end) + 1]:
        style(ws[f"{col}{AN}"], fl=fill(INK), bd=BOX)
ws.row_dimensions[AN].height = 24

units_rng = f"$B${TOTAL}:$E${TOTAL}"
analysis = [
    ("Total department budget", f"=SUM(B{TOTAL}:E{TOTAL})",
     "All four units combined", "SUM"),
    ("Average budget per unit", f"=AVERAGE(B{TOTAL}:E{TOTAL})",
     '="Total budget ÷ "&COUNTA($B$4:$E$4)&" units"', "AVERAGE, COUNTA"),
    ("Highest unit budget", f"=MAX(B{TOTAL}:E{TOTAL})",
     f"=INDEX($B$4:$E$4,MATCH(B{AN + 3},{units_rng},0))", "MAX, INDEX / MATCH"),
    ("Lowest unit budget", f"=MIN(B{TOTAL}:E{TOTAL})",
     f"=INDEX($B$4:$E$4,MATCH(B{AN + 4},{units_rng},0))", "MIN, INDEX / MATCH"),
    ("Largest budget item (all units)", f"=MAX(F{FIRST}:F{LAST})",
     f"=INDEX($A${FIRST}:$A${LAST},MATCH(B{AN + 5},$F${FIRST}:$F${LAST},0))", "MAX, INDEX / MATCH"),
    ("Average cost per budget item", f"=AVERAGE(F{FIRST}:F{LAST})",
     f'="Total budget ÷ "&COUNTA($A${FIRST}:$A${LAST})&" budget items"', "AVERAGE, COUNTA"),
]
for i, (label, amount, detail, func) in enumerate(analysis):
    r = AN + 1 + i
    band = fill(BAND) if i % 2 else fill(WHITE)
    ws.merge_cells(f"C{r}:E{r}")
    ws.merge_cells(f"F{r}:G{r}")
    ws[f"A{r}"] = label
    ws[f"B{r}"] = amount
    ws[f"C{r}"] = detail
    ws[f"F{r}"] = func
    for col in "ABCDEFG":
        style(ws[f"{col}{r}"], fl=band, bd=BOX)
    style(ws[f"A{r}"], f=font(10), al=LEFT)
    style(ws[f"B{r}"], f=font(10, bold=True), fmt=CUR, al=RIGHT)
    style(ws[f"C{r}"], f=font(10), al=LEFT)
    style(ws[f"F{r}"], f=font(9, color=MUTED), al=Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[r].height = 21
AN_LAST = AN + len(analysis)

# Charts section
CH = AN_LAST + 2
ws.merge_cells(f"A{CH}:G{CH}")
ws[f"A{CH}"] = "Budget Charts"
style(ws[f"A{CH}"], f=font(10, bold=True, color=WHITE), fl=fill(INK), al=LEFT)
ws.row_dimensions[CH].height = 24


def rich(size=900, color=INK, bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def title(text):
    cp = CharacterProperties(latin=DrawingFont(typeface=FONT), sz=1200, b=True, solidFill=INK)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def no_line():
    return GraphicalProperties(ln=LineProperties(noFill=True))


values = Reference(ws, min_col=2, max_col=5, min_row=TOTAL, max_row=TOTAL)
names = Reference(ws, min_col=2, max_col=5, min_row=4, max_row=4)

# Column chart: total budget by unit
bar = BarChart()
bar.type = "col"
bar.grouping = "clustered"
bar.title = title("Total Budget by Unit")
bar.add_data(values, from_rows=True, titles_from_data=False)
bar.set_categories(names)
bar.gapWidth = 70
bar.legend = None
series = bar.series[0]
series.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
for idx, color in enumerate(UNIT_COLORS):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(noFill=True))
    series.dPt.append(pt)
bar.dataLabels = DataLabelList(showVal=True, showSerName=False, showCatName=False,
                               showLegendKey=False, showPercent=False)
bar.dataLabels.numFmt = "$#,##0"
bar.dataLabels.position = "outEnd"
bar.dataLabels.txPr = rich(900, INK, bold=True)
bar.x_axis.delete = False
bar.y_axis.delete = False
bar.x_axis.txPr = rich(900, INK)
bar.y_axis.txPr = rich(800, MUTED)
bar.y_axis.numFmt = "$#,##0"
bar.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E1E4EA", w=9525)))
bar.y_axis.spPr = no_line()
bar.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill="C3C8D2", w=9525))
bar.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D5DAE3", w=9525))
bar.width, bar.height = 13.6, 8.2
ws.add_chart(bar, f"A{CH + 2}")

# Doughnut chart: share of the department budget
dn = DoughnutChart(holeSize=52)
dn.title = title("Share of Department Budget")
dn.add_data(values, from_rows=True, titles_from_data=False)
dn.set_categories(names)
series = dn.series[0]
for idx, color in enumerate(UNIT_COLORS):
    pt = DataPoint(idx=idx)
    pt.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=WHITE, w=19050))
    series.dPt.append(pt)
dn.dataLabels = DataLabelList(showPercent=True, showVal=False, showSerName=False,
                              showCatName=False, showLegendKey=False, showLeaderLines=False)
dn.dataLabels.numFmt = "0.0%"
dn.dataLabels.txPr = rich(900, INK, bold=True)
dn.legend = Legend()
dn.legend.position = "r"
dn.legend.txPr = rich(900, INK)
dn.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D5DAE3", w=9525))
dn.width, dn.height = 13.0, 8.2
ws.add_chart(dn, f"D{CH + 2}")

CHART_ROWS = 17
for r in range(CH + 1, CH + 2 + CHART_ROWS):
    ws.row_dimensions[r].height = 15

# Print setup: one landscape page, centred
LAST_ROW = CH + 1 + CHART_ROWS
ws.print_area = f"A1:G{LAST_ROW}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5, header=0.3, footer=0.3)
ws.oddFooter.center.text = "Marketing Department Budget 2026 - Page &P of &N"
ws.oddFooter.center.size = 8
ws.freeze_panes = None
ws.sheet_view.selection[0].activeCell = "A1"
ws.sheet_view.selection[0].sqref = "A1"

wb.properties.title = "Marketing Department Budget 2026"
wb.properties.subject = "Budget by unit: Advertising, Marketing, Public Relations, e-Business"
wb.properties.creator = "Marketing Department"
wb.properties.lastModifiedBy = "Marketing Department"
wb.properties.keywords = "budget, marketing, SUM, AVERAGE"

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"wrote {OUT}")
