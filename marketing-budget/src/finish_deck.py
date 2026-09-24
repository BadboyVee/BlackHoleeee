#!/usr/bin/env python3
"""Final pass over the deck.

- Keeps one paragraph-settings block (<a:pPr>) per paragraph. pptxgenjs writes one before
  every text run, which the schema forbids and PowerPoint 2013 may reject.
- Adds a Fade transition to every slide.

Usage: python finish_deck.py deck.pptx out.pptx
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn


def single_ppr(root):
    """Keep each paragraph's first <a:pPr> (it holds the real settings) and drop the rest."""
    fixed = 0
    for para in root.iter(qn("a:p")):
        pprs = para.findall(qn("a:pPr"))
        if not pprs:
            continue
        keep = pprs[0]
        for extra in pprs[1:]:
            para.remove(extra)
            fixed += 1
        if para.index(keep) != 0:
            para.remove(keep)
            para.insert(0, keep)
    return fixed


def add_fade(sld):
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    trans = etree.SubElement(sld, qn("p:transition"), spd="med")
    etree.SubElement(trans, qn("p:fade"))
    sld.remove(trans)
    # CT_Slide order: cSld, clrMapOvr, transition, timing, extLst
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(trans)


def main(deck, out):
    prs = Presentation(deck)
    dropped = 0
    for slide in prs.slides:
        dropped += single_ppr(slide._element)
        if slide.has_notes_slide:
            dropped += single_ppr(slide.notes_slide._element)
        add_fade(slide._element)
    prs.save(out)
    print(f"fade on {len(prs.slides)} slides, {dropped} duplicate paragraph blocks removed -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:3])
