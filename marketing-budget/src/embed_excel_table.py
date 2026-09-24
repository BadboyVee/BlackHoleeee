#!/usr/bin/env python3
"""Swap the table picture on the Excel slide for a real embedded Excel worksheet object.

PowerPoint's Paste Special > "Microsoft Excel Worksheet Object" stores the workbook inside
the deck and shows a picture of the pasted cells; double-clicking the picture opens Excel.
This does the same thing: it embeds the workbook (whose oleSize names the table block)
where the picture called "Excel Budget Table" sits, and uses that picture as the display.

Usage: python embed_excel_table.py deck.pptx workbook.xlsx preview.png table_size.json out.pptx
"""
import json
import sys

from pptx import Presentation
from pptx.enum.shapes import PROG_ID
from pptx.oxml.ns import qn
from pptx.util import Emu

NAME = "Excel Budget Table"
EMU_PER_POINT = 12700


def main(deck, workbook, preview, size_json, out):
    prs = Presentation(deck)
    with open(size_json) as fh:
        natural_w, natural_h = json.load(fh)["points"]

    matches = [(slide, shape) for slide in prs.slides for shape in slide.shapes if shape.name == NAME]
    if len(matches) != 1:
        sys.exit(f"expected exactly one shape named {NAME!r}, found {len(matches)}")
    slide, pic = matches[0]
    alt_text = pic._element.nvPicPr.cNvPr.get("descr", "")

    frame = slide.shapes.add_ole_object(
        workbook, PROG_ID.XLSX, pic.left, pic.top, pic.width, pic.height,
        icon_file=preview,
        icon_width=Emu(round(natural_w * EMU_PER_POINT)),
        icon_height=Emu(round(natural_h * EMU_PER_POINT)),
    )
    gf = frame._element

    # Display the cells themselves rather than an icon, as Paste Special does.
    ole = gf.find(qn("a:graphic")).find(qn("a:graphicData")).find(qn("p:oleObj"))
    ole.attrib.pop("showAsIcon", None)
    ole.set("name", "Worksheet")

    cnvpr = gf.find(qn("p:nvGraphicFramePr")).find(qn("p:cNvPr"))
    cnvpr.set("name", NAME)
    cnvpr.set("descr", alt_text)
    locks = gf.find(qn("p:nvGraphicFramePr")).find(qn("p:cNvGraphicFramePr")).find(qn("a:graphicFrameLocks"))
    locks.set("noChangeAspect", "1")

    # The picture inside p:oleObj needs an id of its own; no two shapes on a slide may share one.
    inner = ole.find(qn("p:pic")).find(qn("p:nvPicPr")).find(qn("p:cNvPr"))
    used = [int(e.get("id")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    inner.set("id", str(max(used) + 1))
    inner.set("name", NAME + " picture")

    # Take the picture's place in the z-order, then remove the picture.
    pic._element.addprevious(gf)
    pic._element.getparent().remove(pic._element)

    prs.save(out)
    print(f"embedded {workbook} on slide {prs.slides.index(slide) + 1} -> {out}")


if __name__ == "__main__":
    main(*sys.argv[1:6])
