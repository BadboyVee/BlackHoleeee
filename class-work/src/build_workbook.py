#!/usr/bin/env python3
"""Build Promotion_Query.xlsx: five students, five subjects, and a query result that promotes
each student to LAW if their English score is 70 or more, and to SCIENCE otherwise:
    =IF(C5>=70,"LAW","SCIENCE")
Excel 2013 style: Office 2013 theme, Calibri 11. The names and scores are sample data.

Usage: python build_workbook.py out.xlsx
"""
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList

from office2013 import THEME_XML

SUBJECTS = ["English", "Mathematics", "Biology", "Chemistry", "Physics"]
STUDENTS = [
    ("Adebayo Tolulope", [78, 64, 59, 66, 61]),
    ("Okafor Chinedu", [65, 88, 81, 79, 84]),
    ("Ibrahim Aisha", [70, 58, 62, 55, 49]),
    ("Eze Ngozi", [55, 72, 77, 69, 74]),
    ("Bello Musa", [82, 60, 54, 58, 63]),
]

font = lambda **kw: Font(name="Calibri", size=kw.pop("size", 11), family=2, scheme="minor", **kw)
thin = Side(style="thin", color="000000")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
grey = PatternFill("solid", fgColor="D9D9D9")
centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center")

wb = Workbook()
wb.loaded_theme = THEME_XML.encode()
wb._fonts = IndexedList([font()])
wb._named_styles["Normal"].font = font()
wb.properties.creator = wb.properties.lastModifiedBy = ""
wb.properties.title = "Promotion Query"
ws = wb.active
ws.title = "Class Work"

last = chr(ord("B") + len(SUBJECTS) + 1)                     # H: the Promoted To column
ws.merge_cells(f"A1:{last}1")
ws["A1"] = "QUERY RESULT: STUDENTS PROMOTED TO LAW OR SCIENCE"
ws["A1"].font = font(size=14, bold=True)
ws["A1"].alignment = centre
ws.merge_cells(f"A2:{last}2")
ws["A2"] = "Rule: a student whose English score is 70 or more is promoted to LAW; others go to SCIENCE."
ws["A2"].font = font(italic=True)
ws["A2"].alignment = centre

heads = ["S/N", "Name of Student"] + SUBJECTS + ["Promoted To"]
for c, text in enumerate(heads, start=1):
    cell = ws.cell(4, c, text)
    cell.font, cell.fill, cell.border, cell.alignment = font(bold=True), grey, box, centre
ws.row_dimensions[4].height = 30

for i, (name, scores) in enumerate(STUDENTS):
    r = 5 + i
    values = [i + 1, name] + scores + [f'=IF(C{r}>=70,"LAW","SCIENCE")']
    for c, value in enumerate(values, start=1):
        cell = ws.cell(r, c, value)
        cell.border, cell.font = box, font(bold=c == len(values))
        cell.alignment = left if c == 2 else centre

r = 5 + len(STUDENTS) + 1
for k, (label, word) in enumerate((("Number promoted to LAW", "LAW"), ("Number promoted to SCIENCE", "SCIENCE"))):
    ws.merge_cells(f"A{r + k}:G{r + k}")
    ws[f"A{r + k}"] = label
    ws[f"H{r + k}"] = f'=COUNTIF(H5:H{4 + len(STUDENTS)},"{word}")'
    for c in range(1, 9):
        ws.cell(r + k, c).border = box
        ws.cell(r + k, c).font = font(bold=True)
    ws[f"A{r + k}"].alignment = left
    ws[f"H{r + k}"].alignment = centre

for col, w in {"A": 6, "B": 22, "C": 11, "D": 13, "E": 11, "F": 11, "G": 11, "H": 15}.items():
    ws.column_dimensions[col].width = w
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
wb.save(sys.argv[1])
print(f"wrote {sys.argv[1]}")
