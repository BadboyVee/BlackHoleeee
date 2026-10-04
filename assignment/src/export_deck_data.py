#!/usr/bin/env python3
"""Read the figures the slides use out of the finished Relative_Interest_Pie_Chart.xlsx.

Writes deck_data.json for build_deck.js: each programme's average score, rank and share of the
total as Excel worked them out, plus its colour, so the slides always match the workbook.

Usage: python export_deck_data.py Relative_Interest_Pie_Chart.xlsx deck_data.json
"""
import json
import sys

from openpyxl import load_workbook

from assignment_data import PROGRAMMES, SLICES, SOURCE

workbook, out = sys.argv[1:3]
ws = load_workbook(workbook, data_only=True)["Relative Interest"]
first = 5                                      # the table's first programme row
rows = []
for i, colour in enumerate(SLICES):
    name, score, rank, share = (ws.cell(first + i, c).value for c in range(1, 5))
    if name != PROGRAMMES[i][0] or None in (score, rank, share):
        sys.exit(f"row {first + i} of {workbook} is {name!r}: run the workbook through LibreOffice first")
    rows.append({"name": name, "score": score, "rank": rank, "share": share, "colour": colour})
total = ws.cell(first + len(SLICES), 2).value
with open(out, "w") as fh:
    json.dump({"source": SOURCE, "total": total, "programmes": rows}, fh, indent=2)
print(f"wrote {out}")
