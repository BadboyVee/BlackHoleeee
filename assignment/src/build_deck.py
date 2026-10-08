#!/usr/bin/env python3
"""Build the relative-interest pie chart deck in one of three colour palettes, made the way it
would be in PowerPoint 2013.

Usage: python build_deck.py deck_data.json template.pptx wood.jpg STYLE out.pptx
  STYLE is marketing, pz or mixed: the palettes in deck_kit.py (the Marketing Budget deck's,
  the PZ Nigeria deck's wood design, and both mixed).

All three have the same five simple slides: 1. title, 2. the scores, 3. the pie chart,
4. observations, 5. thank you. deck_data.json comes from export_deck_data.py, which reads the
scores, ranks and shares out of the finished Relative_Interest_Pie_Chart.xlsx, so the slides
match the workbook. check_fit.py confirms every text box fits its text, and build.sh then gives
each slide its transition (finish_deck.py).
"""
import json
import sys

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from deck_kit import Deck, into_placeholder, place, placeholder, plot_layout, rgb

data_path, template_path, wood_path, style_name, out_path = sys.argv[1:6]
deck = Deck(style_name, template_path, wood_path)
S, prs, W, H, LAYOUT = deck.S, deck.prs, deck.W, deck.H, deck.layout
write, textbox, rect, content, title_card = deck.write, deck.textbox, deck.rect, deck.content, deck.title_card
data = json.load(open(data_path))
P = data["programmes"]                                  # pie order: largest share first


def pct(share):
    return f"{int(share * 100 + 0.5)}%"                 # rounded as Excel shows 0%


# ================================================================== 1. Title
s = prs.slides.add_slide(LAYOUT["Title Slide"])
title_card(s, "Relative Interest", ["Digital Skills Programmes on Google Trends",
                                     "Market Research Lab, Task 1  ·  Nigeria, past 12 months"])
s.notes_slide.notes_text_frame.text = (
    "My presentation is on the relative interest in five digital skills programmes in Nigeria, using Google Trends.")

# ================================================================== 2. The scores
s = prs.slides.add_slide(LAYOUT["Title and Content"])
top = content(s, "Relative-Interest Observations", "The average Google Trends score of each programme")
ph = placeholder(s, 1)
rows = [("Programme", "Average score", "Rank")] + [(p["name"], str(p["score"]), str(p["rank"])) for p in P]
widths = [Inches(5.9), Inches(2.6), Inches(2.1)]
ROW_H = Inches(0.6)
table_w = sum(widths, Emu(0))
# In the content placeholder's slot, as the Insert Table icon puts it (then moved to the middle).
deck.table(s, ph, (W - table_w) // 2, top + Inches(0.1), rows, widths, ROW_H, 20)
textbox(s, Inches(0.8), top + Inches(0.1) + ROW_H * len(rows) + Inches(0.3), W - Inches(1.6), Inches(0.45),
        ["Google Trends gives each programme a score from 0 to 100, where 100 means the highest interest."],
        size=16, colour=S["muted"], align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = (
    "These are the average scores I got from Google Trends for the past 12 months. Cybersecurity had the "
    "highest score, and Generative AI & Prompt Engineering and Web Development had the lowest.")

# ================================================================== 3. The pie chart
s = prs.slides.add_slide(LAYOUT["Title and Content"])
top = content(s, "Relative Interest by Programme", "Each programme's share of the total interest")
ph = placeholder(s, 1)
chart_data = CategoryChartData(number_format="0")
chart_data.categories = [p["name"] for p in P]
chart_data.add_series("Average score", [p["score"] for p in P])
CHART = (Inches(1.0), top - Inches(0.05), Inches(6.6), Inches(4.65))
gframe = s.shapes.add_chart(XL_CHART_TYPE.PIE, *CHART, chart_data)
chart = gframe.chart
chart.font.name = S["body_font"]
chart.font.size = Pt(16)
chart.font.color.rgb = rgb(S["text"])
chart.has_title = True
write(chart.chart_title.text_frame, ["Share of Total Interest"], size=16, colour=S["text"])
chart.has_legend = False                                # the key beside the chart names the slices
plot = chart.plots[0]
plot.vary_by_categories = True
series = plot.series[0]
for i, colour in enumerate(S["pie"]):
    point = series.points[i].format
    point.fill.solid()
    point.fill.fore_color.rgb = rgb(colour)
    point.line.color.rgb = rgb("FFFFFF")                # white lines between the slices
    point.line.width = Pt(1.5)
plot.has_data_labels = True
labels = plot.data_labels
labels.show_percentage = True
labels.show_value = False
labels.show_category_name = False
labels.position = XL_LABEL_POSITION.OUTSIDE_END
labels.font.size = Pt(16)
labels.font.bold = True
labels.font.name = S["body_font"]
labels.font.color.rgb = rgb(S["text"])
# A big pie under the chart title, in the content placeholder's slot (Insert Chart icon).
plot_layout(chart, 0.235, 0.18, 0.53, 0.74)
into_placeholder(ph, gframe)
# The key: each programme's shade, name, score and share.
for i, (p, colour) in enumerate(zip(P, S["pie"])):
    y = top + Inches(0.25) + Inches(0.86) * i
    rect(s, Inches(7.85), y + Inches(0.06), Inches(0.28), Inches(0.28), fill=colour, line=S["key_line"], width=0.75)
    key = textbox(s, Inches(8.25), y - Inches(0.04), Inches(4.5), Inches(0.82),
                  [[(p["name"], True)], f"Score {p['score']}  ·  {pct(p['share'])}"], size=16, colour=S["text"])
    key.text_frame.paragraphs[1].runs[0].font.size = Pt(15)
textbox(s, Inches(1.0), H - Inches(0.8), W - Inches(2.0), Inches(0.3),
        ["Source: Google Trends, Nigeria, past 12 months (28 September 2025 to 28 September 2026)."],
        size=12, colour=S["muted"], align=PP_ALIGN.CENTER)
s.notes_slide.notes_text_frame.text = (
    f"This pie chart shows the share of interest for each programme. {P[0]['name']} has the biggest share at "
    f"{pct(P[0]['share'])}, then {P[1]['name']} at {pct(P[1]['share'])} and {P[2]['name']} at "
    f"{pct(P[2]['share'])}. {P[3]['name']} and {P[4]['name']} have {pct(P[4]['share'])} each.")

# ================================================================== 4. Observations
lowest = [p for p in P if p["score"] == P[-1]["score"]]
assert len(P) == 5 and len(lowest) == 2, "the sentences below expect a tie for last place"
s = prs.slides.add_slide(LAYOUT["Title and Content"])
top = content(s, "Observations", "What the pie chart shows")
box = rect(s, Inches(1.0), top + Inches(0.05), W - Inches(2.0), Inches(4.45), fill=S["box"])
body = placeholder(s, 1)
place(body, Inches(1.35), top + Inches(0.3), W - Inches(2.7), Inches(4.0))
box._element.getparent().remove(box._element)
body._element.addprevious(box._element)                 # the box behind the text
name = lambda p: (p["name"], True)
write(body.text_frame, [
    [name(P[0]), (f" had the highest interest, with {pct(P[0]['share'])} of the total.", False)],
    [name(P[1]), (f" came second with {pct(P[1]['share'])}.", False)],
    [name(P[2]), (f" came third with {pct(P[2]['share'])}.", False)],
    [name(lowest[0]), (" and ", False), name(lowest[1]),
     (f" had the lowest interest, with {pct(lowest[0]['share'])} each.", False)],
    [(f"This means {P[0]['name']} and {P[1]['name']} should be promoted first.", False)],
], size=24, colour=S["text"], after=10)
body.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE     # the sentences centred in their box
for para in body.text_frame.paragraphs:                 # programme names in the style's colour
    for run in para.runs:
        if run.font.bold:
            run.font.color.rgb = rgb(S["names"])
s.notes_slide.notes_text_frame.text = (
    f"From the chart, {P[0]['name']} is the most popular programme, followed by {P[1]['name']}. The organisation "
    "should promote these two first.")

# ================================================================== 5. Thank you
s = prs.slides.add_slide(LAYOUT["Title Slide"])
title_card(s, "Thank You", ["Any questions?", "Market Research Lab, Task 1"])
s.notes_slide.notes_text_frame.text = "Thank you for listening."

deck.save(out_path, title="Relative Interest in Digital Skills Programmes",
          subject="Pie chart of the Google Trends relative-interest observations",
          keywords="Google Trends, relative interest, pie chart")
