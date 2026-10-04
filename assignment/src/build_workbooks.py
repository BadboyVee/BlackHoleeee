#!/usr/bin/env python3
"""Build the two assignment workbooks in Excel 2013's own style.

1. Student_Test_Results.xlsx: the test results of 25 students in 7 subjects. Each subject has a
   score (typed in, out of 100) and a grade worked out with the assignment's nested IF formula,
       =IF(C5>=80,"A1",IF(C5>=70,"B2",IF(C5>=66,"B3",IF(C5>=60,"C4",IF(C5>=56,"C5",
         IF(C5>=50,"C6",IF(C5>=46,"D7",IF(C5>=40,"E8","F9"))))))))
   pointed at that subject's score cell. Each student also gets a total (SUM), an average
   (AVERAGE), an overall grade (the same IF formula on the average) and a position (RANK).
   Under the table: each subject's highest, lowest and average score (MAX, MIN, AVERAGE) and
   how many students got each grade (COUNTIF).
2. Relative_Interest_Pie_Chart.xlsx: the relative-interest observations from Market Research Lab
   Task 1 (Google Trends), with each programme's rank (RANK) and share of the total, and a pie
   chart of them, ready to copy into PowerPoint.

Both use the Office 2013 theme, Calibri 11 and Excel 2013's default chart style, with dark blue
("Blue, Accent 1, Darker 50%") header rows and orange ("Orange, Accent 2, Darker 25%") for F9
grades and the largest slice. Only functions that exist in Excel 2013 are used. The student
names and scores are sample data: type over them and every grade, total, average, position and
count updates.

Usage: python build_workbooks.py results.xlsx interest.xlsx
"""
import re
import shutil
import sys
import tempfile
import zipfile

from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.drawing.colors import ColorChoice, SchemeColor
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    CharacterProperties,
    Font as DrawingFont,
    Paragraph,
    ParagraphProperties,
    RegularTextRun,
)
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins

from assignment_data import GRADES, PROGRAMMES, SLICES, SOURCE, STUDENTS, SUBJECTS, grade_formula
from office2013 import THEME_XML

results_path, interest_path = sys.argv[1:3]

# ================================================================ styling (as in the PZ workbook)
NAVY = Color(theme=4, tint=-0.499984740745262)      # "Blue, Accent 1, Darker 50%"    1F4E79
ORANGE = Color(theme=5, tint=-0.249977111117893)    # "Orange, Accent 2, Darker 25%"  C55A11
WHITE = Color(theme=0)                              # "White, Background 1"
MUTED = "595959"                                    # Excel 2013's chart text grey
thin = Side(style="thin", color=Color(theme=0, tint=-0.249977111117893))   # White, Background 1, Darker 25%
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)                  # Home > Borders > All Borders


def font(size=11, bold=False, italic=False, color="000000"):
    return Font(name="Calibri", size=size, bold=bold, italic=italic, color=color, family=2, scheme="minor")


def al(h="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical="center", wrap_text=wrap, indent=indent)


def put(ws, ref, value, *, f=None, fill=None, fmt=None, a=None, bd=BOX):
    c = ws[ref]
    c.value = value
    c.font = f or font()
    if fill is not None:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt is not None:
        c.number_format = fmt
    c.alignment = a or al()
    if bd is not None:
        c.border = bd
    return c


def span(ws, first, last, value, **kw):
    """Write value into the merged range first:last, with the border (and fill) on every cell."""
    if first != last:
        ws.merge_cells(f"{first}:{last}")
        for row in ws[f"{first}:{last}"]:
            for c in row:
                c.border = kw.get("bd", BOX)
                if kw.get("fill") is not None:
                    c.fill = PatternFill("solid", fgColor=kw["fill"])
    return put(ws, first, value, **kw)


def head(ws, first, last, text, wrap=True):
    """A header cell (or merged block): dark blue, white bold text."""
    return span(ws, first, last, text, f=font(bold=True, color=WHITE), fill=NAVY, a=al(wrap=wrap))


def title(ws, last_col, text, subtitle):
    ws.merge_cells(f"A1:{last_col}1")
    put(ws, "A1", text, f=font(14, bold=True, color=NAVY), a=al(), bd=None)
    ws.row_dimensions[1].height = 21
    ws.merge_cells(f"A2:{last_col}2")
    put(ws, "A2", subtitle, f=font(italic=True), a=al(), bd=None)


def page(ws, landscape):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75, header=0.3, footer=0.3)


def new_workbook(subject):
    wb = Workbook()
    wb.loaded_theme = THEME_XML.encode()      # the Office 2013 theme (Calibri, 2013 colours)
    wb._fonts = IndexedList([font()])
    wb._named_styles["Normal"].font = font()
    wb.properties.creator = wb.properties.lastModifiedBy = ""
    wb.properties.subject = subject
    return wb


# ================================================================ 1. test results
wb = new_workbook("Test results of 25 students in 7 subjects, with the grade of each subject")
wb.properties.title = "Test Results"
ws = wb.active
ws.title = "Test Results"
page(ws, landscape=True)

FIRST, LAST = 5, 4 + len(STUDENTS)                       # student rows 5..29
SCORE = [get_column_letter(3 + 2 * j) for j in range(len(SUBJECTS))]   # C, E, G, ... O
GRADE = [get_column_letter(4 + 2 * j) for j in range(len(SUBJECTS))]   # D, F, H, ... P
TOTAL, AVERAGE, OVERALL, POSITION = (get_column_letter(3 + 2 * len(SUBJECTS) + k) for k in range(4))  # Q R S T

widths = {"A": 6, "B": 24, TOTAL: 9, AVERAGE: 9, OVERALL: 9, POSITION: 9}
widths.update({c: 8 for c in SCORE + GRADE})
for col, w in widths.items():
    ws.column_dimensions[col].width = w

title(ws, POSITION, "TEST RESULTS OF 25 STUDENTS IN 7 SUBJECTS",
      "Each subject is scored out of 100. The grades, totals, averages and positions are worked out "
      "with formulas (IF, SUM, AVERAGE and RANK).")

# Two header rows: the subject over its Score and Grade columns.
head(ws, "A3", "A4", "S/N")
head(ws, "B3", "B4", "Name of Student")
for subject, s, g in zip(SUBJECTS, SCORE, GRADE):
    head(ws, f"{s}3", f"{g}3", subject)
    head(ws, f"{s}4", f"{s}4", "Score")
    head(ws, f"{g}4", f"{g}4", "Grade")
for col, text in ((TOTAL, "Total"), (AVERAGE, "Average"), (OVERALL, "Overall Grade"), (POSITION, "Position")):
    head(ws, f"{col}3", f"{col}4", text)
ws.row_dimensions[3].height = 30
ws.row_dimensions[4].height = 18

scores_of = lambda r: ",".join(f"{s}{r}" for s in SCORE)          # C5,E5,G5,I5,K5,M5,O5
for i, (name, scores) in enumerate(STUDENTS):
    r = FIRST + i
    put(ws, f"A{r}", i + 1)
    put(ws, f"B{r}", name, a=al("left", indent=1))
    for s, g, score in zip(SCORE, GRADE, scores):
        put(ws, f"{s}{r}", score)
        put(ws, f"{g}{r}", grade_formula(f"{s}{r}"))
    put(ws, f"{TOTAL}{r}", f"=SUM({scores_of(r)})", f=font(bold=True))
    put(ws, f"{AVERAGE}{r}", f"=AVERAGE({scores_of(r)})", fmt="0.0")
    put(ws, f"{OVERALL}{r}", grade_formula(f"{AVERAGE}{r}"), f=font(bold=True))
    put(ws, f"{POSITION}{r}", f"=RANK({TOTAL}{r},${TOTAL}${FIRST}:${TOTAL}${LAST})")

# Each subject's highest, lowest and average score.
for k, (label, func, fmt) in enumerate((("Highest score", "MAX", "0"), ("Lowest score", "MIN", "0"),
                                        ("Class average", "AVERAGE", "0.0"))):
    r = LAST + 1 + k
    span(ws, f"A{r}", f"B{r}", label, f=font(bold=True), a=al("left", indent=1))
    for col in SCORE + [TOTAL, AVERAGE]:
        put(ws, f"{col}{r}", f"={func}({col}{FIRST}:{col}{LAST})", f=font(bold=True),
            fmt="0.0" if col == AVERAGE else fmt)
    for col in GRADE + [OVERALL, POSITION]:
        put(ws, f"{col}{r}", None)
STATS_LAST = LAST + 3

# How many students got each grade, per subject and overall.
G_TITLE = STATS_LAST + 2
ws.merge_cells(f"A{G_TITLE}:{POSITION}{G_TITLE}")
put(ws, f"A{G_TITLE}", "GRADE SUMMARY: NUMBER OF STUDENTS WITH EACH GRADE (COUNTIF)",
    f=font(12, bold=True, color=NAVY), a=al("left"), bd=None)
ws.row_dimensions[G_TITLE].height = 18
G_HEAD = G_TITLE + 1
head(ws, f"A{G_HEAD}", f"A{G_HEAD}", "Grade")
head(ws, f"B{G_HEAD}", f"B{G_HEAD}", "Score (remark)")
for subject, s, g in zip(SUBJECTS, SCORE, GRADE):
    head(ws, f"{s}{G_HEAD}", f"{g}{G_HEAD}", subject)
head(ws, f"{TOTAL}{G_HEAD}", f"{AVERAGE}{G_HEAD}", "All subjects")
head(ws, f"{OVERALL}{G_HEAD}", f"{POSITION}{G_HEAD}", "Overall grade")
ws.row_dimensions[G_HEAD].height = 30
G_FIRST = G_HEAD + 1
G_LAST = G_FIRST + len(GRADES) - 1
for k, (grade, low, high, remark) in enumerate(GRADES):
    r = G_FIRST + k
    put(ws, f"A{r}", grade, f=font(bold=True))
    put(ws, f"B{r}", f"{low} to {high} ({remark})", a=al("left", indent=1))
    for s, g in zip(SCORE, GRADE):
        span(ws, f"{s}{r}", f"{g}{r}", f"=COUNTIF({g}${FIRST}:{g}${LAST},$A{r})")
    span(ws, f"{TOTAL}{r}", f"{AVERAGE}{r}", f"=SUM({SCORE[0]}{r}:{GRADE[-1]}{r})", f=font(bold=True))
    span(ws, f"{OVERALL}{r}", f"{POSITION}{r}", f"=COUNTIF(${OVERALL}${FIRST}:${OVERALL}${LAST},$A{r})",
         f=font(bold=True))
r = G_LAST + 1
span(ws, f"A{r}", f"B{r}", "Number of students", f=font(bold=True), a=al("left", indent=1))
for first, last in list(zip(SCORE, GRADE)) + [(TOTAL, AVERAGE), (OVERALL, POSITION)]:
    span(ws, f"{first}{r}", f"{last}{r}", f"=SUM({first}{G_FIRST}:{first}{G_LAST})", f=font(bold=True))

# F9 (fail) grades in orange, following the formulas (Home > Conditional Formatting).
grade_cells = " ".join(f"{g}{FIRST}:{g}{LAST}" for g in GRADE + [OVERALL])
ws.conditional_formatting.add(grade_cells, CellIsRule(operator="equal", formula=['"F9"'],
                                                      font=Font(bold=True, color=ORANGE)))
# Scores must be whole numbers from 0 to 100 (Data > Data Validation).
check = DataValidation(type="whole", operator="between", formula1="0", formula2="100", allow_blank=True,
                       showErrorMessage=True, errorTitle="Score", error="Type a whole number from 0 to 100.")
for s in SCORE:
    check.add(f"{s}{FIRST}:{s}{LAST}")
ws.add_data_validation(check)
ws.freeze_panes = f"C{FIRST}"
wb.save(results_path)
print(f"wrote {results_path}")

# ================================================================ 2. relative interest and pie chart
wb = new_workbook("Relative-interest observations from Google Trends, with a pie chart")
wb.properties.title = "Relative-Interest Observations"
ws = wb.active
ws.title = "Relative Interest"
page(ws, landscape=False)
for col, w in {"A": 36, "B": 15, "C": 10, "D": 15}.items():
    ws.column_dimensions[col].width = w
title(ws, "D", "RELATIVE-INTEREST OBSERVATIONS",
      "Average interest in each programme, from 0 to 100 (read off the chart, so approximate).")
ws.merge_cells("A3:D3")
put(ws, "A3", f"Source: {SOURCE}.", f=font(9, italic=True, color=MUTED), a=al(), bd=None)

P_FIRST, P_LAST = 5, 4 + len(PROGRAMMES)
for col, text in zip("ABCD", ("Programme", "Average score", "Rank", "Share of total")):
    head(ws, f"{col}4", f"{col}4", text)
ws.row_dimensions[4].height = 20
P_TOTAL = P_LAST + 1
for i, (name, score) in enumerate(PROGRAMMES):
    r = P_FIRST + i
    put(ws, f"A{r}", name, a=al("left", indent=1))
    put(ws, f"B{r}", score)
    put(ws, f"C{r}", f"=RANK(B{r},$B${P_FIRST}:$B${P_LAST})")
    put(ws, f"D{r}", f"=B{r}/$B${P_TOTAL}", fmt="0%")
put(ws, f"A{P_TOTAL}", "Total", f=font(bold=True), a=al("left", indent=1))
put(ws, f"B{P_TOTAL}", f"=SUM(B{P_FIRST}:B{P_LAST})", f=font(bold=True))
put(ws, f"C{P_TOTAL}", None)
put(ws, f"D{P_TOTAL}", f"=SUM(D{P_FIRST}:D{P_LAST})", f=font(bold=True), fmt="0%")


def rich(size=900, color=MUTED, bold=False):
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=size, b=bold, solidFill=color)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def chart_title(text):
    cp = CharacterProperties(latin=DrawingFont(typeface="+mn-lt"), sz=1400, b=False, solidFill=MUTED)
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])
    return Title(tx=Text(rich=RichText(p=[para])), overlay=False)


def slice_fill(accent, brightness):
    """A slice in a theme colour with PowerPoint's Lighter/Darker setting, edged in white as in
    Excel 2013's default pie style."""
    if brightness < 0:
        colour = SchemeColor(val=accent, lumMod=round((1 + brightness) * 100000))
    elif brightness > 0:
        colour = SchemeColor(val=accent, lumMod=round((1 - brightness) * 100000),
                             lumOff=round(brightness * 100000))
    else:
        colour = SchemeColor(val=accent)
    return GraphicalProperties(solidFill=ColorChoice(schemeClr=colour),
                               ln=LineProperties(solidFill=ColorChoice(schemeClr=SchemeColor(val="bg1")), w=19050))


pie = PieChart()
pie.title = chart_title("Relative Interest by Programme")
pie.add_data(Reference(ws, min_col=2, min_row=4, max_row=P_LAST), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=1, min_row=P_FIRST, max_row=P_LAST))
pie.firstSliceAng = 0
series = pie.series[0]
for i, (accent, brightness) in enumerate(SLICES):
    series.dPt.append(DataPoint(idx=i, spPr=slice_fill(accent, brightness)))
# Each slice shows its score and its share of the total, outside the pie.
pie.dataLabels = DataLabelList(showVal=True, showPercent=True, showCatName=False, showSerName=False,
                               showLegendKey=False, showLeaderLines=True, separator="\n")
pie.dataLabels.position = "outEnd"
pie.dataLabels.txPr = rich(1000, "404040", bold=True)
pie.legend.position = "r"
pie.legend.txPr = rich(1000)
# Leave room for the labels around the pie and the legend on the right.
pie.layout = Layout(manualLayout=ManualLayout(layoutTarget="inner", xMode="edge", yMode="edge",
                                              x=0.07, y=0.2, w=0.42, h=0.7))
pie.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill="D9D9D9", w=9525))
pie.width, pie.height = 16.5, 9.5
ws.add_chart(pie, f"A{P_TOTAL + 2}")
wb.save(interest_path)


def fix_separators(path):
    """openpyxl writes a data label separator as <c:separator val="..."/> after
    <c:showLeaderLines>; the chart schema wants the text inside the element, before it."""
    c = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    with zipfile.ZipFile(path) as zin, tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if re.fullmatch(r"xl/charts/chart\d+\.xml", info.filename):
                    root = etree.fromstring(data)
                    for sep in list(root.iter(f"{c}separator")):
                        if "val" in sep.attrib:
                            sep.text = sep.attrib.pop("val")
                        leader = sep.getparent().find(f"{c}showLeaderLines")
                        if leader is not None:
                            leader.addprevious(sep)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zout.writestr(info, data)
    shutil.move(tmp.name, path)


fix_separators(interest_path)
print(f"wrote {interest_path}")
